from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


@dataclass(frozen=True)
class ApplicabilityIssue:
    severity: str
    issue_code: str
    message: str
    source_reference: str | None = None


def _resolve_path(project_root: Path, value: str) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = project_root / path
    return path.resolve()


def _normalize_key(series: Any, size: Any) -> str:
    return f"{str(series).strip().upper()}|{str(size).strip().upper()}"


def _normalize_marker(value: Any, *, uppercase: bool = True) -> str | None:
    if value in (None, ""):
        return None

    if isinstance(value, bool):
        text = str(value)
    elif isinstance(value, (int, float)):
        # Pump Options should be categorical. Numeric values are retained for review.
        text = str(value)
    else:
        text = str(value).strip()

    if not text:
        return None

    return text.upper() if uppercase else text


def _canonical_label_map(field_reconciliation: list[dict[str, Any]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in field_reconciliation:
        canonical = item["canonical_field_code"]
        for label in item.get("labels", []):
            result[str(label).strip().upper()] = canonical
    return result


def _domain_map(domains: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for item in domains:
        result[str(item["field_label"]).strip().upper()] = item
    return result


def _option_values(domain: dict[str, Any]) -> list[dict[str, Any]]:
    headers = domain.get("headers", [])
    primary = domain["field_label"]
    values: list[dict[str, Any]] = []

    for row_index, row in enumerate(domain.get("rows", []), start=1):
        value = row.get(primary)
        source_header = primary

        if value in (None, ""):
            # Preserve alignment for multi-column option tables such as Seal Type / OLD JC Style.
            for header in headers:
                candidate = row.get(header)
                if candidate not in (None, ""):
                    value = candidate
                    source_header = header
                    break

        values.append(
            {
                "domain_row_index": row_index,
                "option_value": value,
                "option_source_header": source_header,
            }
        )

    return values


def _authoritative_model(
    model_row: dict[str, Any],
    *,
    require_base_identifier: bool,
) -> tuple[str | None, str | None]:
    model = model_row.get("authoritative_model_identifier")
    base = model_row.get("authoritative_base_identifier")

    if not model and model_row.get("classification") == "LOGIC_ONLY_SUPPLEMENTAL_READY":
        model = model_row.get("logic_model_identifier")

    if not base and model and str(model).upper().startswith("A"):
        base = "D" + str(model)[1:]

    if require_base_identifier and not base:
        return model, None

    return model, base


def compile_model_option_applicability(
    *,
    project_root: Path,
    profile_path: Path,
) -> dict[str, Any]:
    project_root = project_root.resolve()
    profile = json.loads(profile_path.read_text(encoding="utf-8"))

    reconciliation_path = _resolve_path(project_root, profile["reconciliation_json"])
    workbook_path = _resolve_path(project_root, profile["workbook"])

    if not reconciliation_path.exists():
        raise FileNotFoundError(reconciliation_path)
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)

    reconciliation = json.loads(
        reconciliation_path.read_text(encoding="utf-8")
    )

    if reconciliation.get("summary", {}).get("error_count", 0):
        raise ValueError(
            "M023.1 reconciliation contains errors. "
            "Applicability compilation is blocked until reconciliation errors are zero."
        )

    worksheet_name = profile["worksheet"]
    accepted_classes = set(profile["accepted_model_classifications"])
    blocked_classes = set(profile.get("blocked_model_classifications", []))
    marker_semantics = {
        str(key).upper(): value
        for key, value in profile["marker_semantics"].items()
    }
    unknown_policy = profile["unknown_marker_policy"]
    uppercase_marker = bool(
        profile.get("normalization", {}).get("uppercase_marker", True)
    )
    require_base = bool(
        profile.get("safety", {}).get("require_reconciled_base_identifier", True)
    )

    issues: list[ApplicabilityIssue] = []
    label_to_canonical = _canonical_label_map(
        reconciliation["field_reconciliation"]
    )
    domains_by_label = _domain_map(reconciliation["logic_option_domains"])

    # Index each reconciliation row by BOTH its effective/authoritative
    # model key and its original Pump Options key.  M023.1 may reconcile
    # a Logic workbook series/size alias (for example PH3170 -> PH2170).
    # Pump Options still contains the original Logic key, so applicability
    # compilation must be able to resolve that original row without
    # weakening the authoritative Rev2 identity.
    model_reconciliation: dict[str, dict[str, Any]] = {}

    for item in reconciliation["model_reconciliation"]:
        effective_key = item["model_key"]
        model_reconciliation[effective_key] = item

        original_logic_key = item.get("logic_original_model_key")
        if original_logic_key:
            existing = model_reconciliation.get(original_logic_key)
            if existing is not None and existing is not item:
                raise ValueError(
                    "M023.1 reconciliation produced an ambiguous original "
                    f"Pump Options model key: {original_logic_key}"
                )
            model_reconciliation[original_logic_key] = item

    groups = reconciliation["pump_options_matrix"]["groups"]
    model_rows = reconciliation["pump_options_matrix"]["models"]

    # Cache model reconciliation decisions before touching Excel.
    accepted_models: dict[str, dict[str, Any]] = {}
    blocked_models: dict[str, dict[str, Any]] = {}

    for model in model_rows:
        key = model["model_key"]
        resolved = model_reconciliation.get(key)

        if not resolved:
            issues.append(
                ApplicabilityIssue(
                    "Error",
                    "MODEL_RECONCILIATION_MISSING",
                    f"Pump Options model {key} has no M023.1 reconciliation record.",
                    key,
                )
            )
            continue

        classification = resolved["classification"]

        if classification in accepted_classes:
            authoritative_model, base_identifier = _authoritative_model(
                resolved,
                require_base_identifier=require_base,
            )
            if require_base and not base_identifier:
                issues.append(
                    ApplicabilityIssue(
                        "Error",
                        "BASE_IDENTIFIER_MISSING",
                        f"Accepted model {key} has no reconciled Dean base identifier.",
                        key,
                    )
                )
                continue

            accepted_models[key] = {
                **model,
                "reconciliation_classification": classification,
                "authoritative_model_identifier": authoritative_model,
                "base_identifier": base_identifier,
            }
        elif classification in blocked_classes:
            blocked_models[key] = {
                **model,
                "reconciliation_classification": classification,
            }
        else:
            issues.append(
                ApplicabilityIssue(
                    "Error",
                    "MODEL_CLASSIFICATION_UNHANDLED",
                    f"Model {key} has unhandled reconciliation classification {classification!r}.",
                    key,
                )
            )

    # Build field/group alignment from normalized Config Options rows.
    compiled_groups: list[dict[str, Any]] = []

    for group in groups:
        field_label = str(group["field_label"]).strip()
        label_key = field_label.upper()
        canonical = label_to_canonical.get(
            label_key,
            group.get("canonical_field_code"),
        )
        domain = domains_by_label.get(label_key)

        if domain is None:
            issues.append(
                ApplicabilityIssue(
                    "Warning",
                    "MATRIX_GROUP_WITHOUT_OPTION_DOMAIN",
                    f"Pump Options group {field_label!r} has no matching Config Options domain. "
                    "The group is blocked from applicability publication.",
                    group["group_column"],
                )
            )
            continue

        options = _option_values(domain)
        start_col = column_index_from_string(group["group_column"])
        end_col = column_index_from_string(group["span_end_column"])
        span_width = end_col - start_col + 1

        if len(options) > span_width:
            issues.append(
                ApplicabilityIssue(
                    "Error",
                    "OPTION_DOMAIN_EXCEEDS_MATRIX_SPAN",
                    f"{field_label!r} has {len(options)} Config Options rows but only "
                    f"{span_width} columns in its Pump Options matrix span.",
                    f"{group['group_column']}:{group['span_end_column']}",
                )
            )
            continue

        aligned_options = []
        for offset, option in enumerate(options):
            aligned_options.append(
                {
                    **option,
                    "column_index": start_col + offset,
                }
            )

        compiled_groups.append(
            {
                "field_label": field_label,
                "canonical_field_code": canonical,
                "source_group_column": group["group_column"],
                "source_span_end_column": group["span_end_column"],
                "domain_table": domain["table_name"],
                "domain_row_count": len(options),
                "aligned_options": aligned_options,
            }
        )

    candidates: list[dict[str, Any]] = []
    reviews: list[dict[str, Any]] = []
    marker_counts: Counter[str] = Counter()
    status_counts: Counter[str] = Counter()
    field_counts: Counter[str] = Counter()
    allowed_counts: Counter[str] = Counter()

    # Only the small PumpConfiguration_Logic workbook is opened here.
    wb = load_workbook(
        workbook_path,
        read_only=False,
        data_only=True,
        keep_vba=False,
        keep_links=False,
    )

    try:
        ws = wb[worksheet_name]

        for model_key, model in accepted_models.items():
            row = int(model["source_row"])

            # Guard against workbook drift since M023.1.
            live_model = ws.cell(row, 1).value
            live_series = ws.cell(row, 2).value
            live_size = ws.cell(row, 3).value

            live_key = _normalize_key(live_series, live_size)

            if live_key != model_key:
                issues.append(
                    ApplicabilityIssue(
                        "Error",
                        "PUMP_OPTIONS_ROW_DRIFT",
                        f"Cached M023.1 row {row} expected {model_key}, but the current workbook "
                        f"contains {live_key}. Rerun the full M023.1 compiler.",
                        f"{worksheet_name}!A{row}:C{row}",
                    )
                )
                continue

            for group in compiled_groups:
                for option in group["aligned_options"]:
                    option_value = option["option_value"]

                    if option_value in (None, ""):
                        reviews.append(
                            {
                                "review_type": "EMPTY_OPTION_VALUE",
                                "model_key": model_key,
                                "series": model["series"],
                                "size": model["size"],
                                "field_code": group["canonical_field_code"],
                                "field_label": group["field_label"],
                                "option_value": None,
                                "marker_raw": None,
                                "source_row": row,
                                "source_column_index": option["column_index"],
                                "source_domain_table": group["domain_table"],
                                "publication_status": "BLOCKED_REVIEW",
                            }
                        )
                        continue

                    raw_value = ws.cell(row, option["column_index"]).value
                    marker = _normalize_marker(
                        raw_value,
                        uppercase=uppercase_marker,
                    )

                    # Blank is an explicit non-applicable state and does not become a candidate.
                    if marker is None:
                        continue

                    marker_counts[marker] += 1

                    semantics = marker_semantics.get(marker)
                    known_marker = semantics is not None

                    if semantics is None:
                        semantics = unknown_policy

                    status = semantics["status"]
                    is_allowed = bool(semantics["is_allowed"])
                    is_standard = bool(semantics["is_standard"])
                    publication_status = semantics["publication_status"]

                    record = {
                        "family_code": profile["family_code"].upper(),
                        "applicability_version": profile["version"],
                        "model_key": model_key,
                        "logic_model_identifier": live_model,
                        "authoritative_model_identifier": model[
                            "authoritative_model_identifier"
                        ],
                        "base_identifier": model["base_identifier"],
                        "series": str(live_series).strip(),
                        "size": str(live_size).strip(),
                        "model_reconciliation_classification": model[
                            "reconciliation_classification"
                        ],
                        "field_code": group["canonical_field_code"],
                        "field_label": group["field_label"],
                        "option_value": option_value,
                        "option_source_header": option["option_source_header"],
                        "domain_table": group["domain_table"],
                        "domain_row_index": option["domain_row_index"],
                        "marker_raw": raw_value,
                        "marker_normalized": marker,
                        "applicability_status": status,
                        "is_allowed": is_allowed,
                        "is_standard": is_standard,
                        "marker_is_known": known_marker,
                        "publication_status": publication_status,
                        "source_workbook": workbook_path.name,
                        "source_worksheet": worksheet_name,
                        "source_row": row,
                        "source_column_index": option["column_index"],
                    }

                    candidates.append(record)
                    status_counts[status] += 1
                    field_counts[group["canonical_field_code"]] += 1

                    if is_allowed:
                        allowed_counts[group["canonical_field_code"]] += 1

                    if publication_status != "READY":
                        reviews.append(
                            {
                                "review_type": (
                                    "KNOWN_BLOCKED_MARKER"
                                    if known_marker
                                    else "UNKNOWN_MARKER"
                                ),
                                **record,
                            }
                        )

        # Record blocked models separately; they can never leak into candidates.
        for model_key, model in blocked_models.items():
            reviews.append(
                {
                    "review_type": "BLOCKED_MODEL",
                    "model_key": model_key,
                    "series": model["series"],
                    "size": model["size"],
                    "logic_model_identifier": model["model_identifier"],
                    "model_reconciliation_classification": model[
                        "reconciliation_classification"
                    ],
                    "publication_status": "BLOCKED_REVIEW",
                    "source_row": model["source_row"],
                }
            )
    finally:
        wb.close()

    # Safety validations.
    standard_by_model_field: Counter[tuple[str, str]] = Counter()
    allowed_by_model_field: Counter[tuple[str, str]] = Counter()

    for row in candidates:
        key = (row["model_key"], row["field_code"])
        if row["is_standard"]:
            standard_by_model_field[key] += 1
        if row["is_allowed"]:
            allowed_by_model_field[key] += 1

    for key, count in standard_by_model_field.items():
        if count > 1:
            issues.append(
                ApplicabilityIssue(
                    "Warning",
                    "MULTIPLE_STANDARD_OPTIONS",
                    f"{key[0]} / {key[1]} has {count} STD options. "
                    "The compiler preserves them but blocks default-selection assumptions.",
                    f"{key[0]}|{key[1]}",
                )
            )

    zero_allowed_pairs = []
    for model_key in accepted_models:
        for group in compiled_groups:
            pair = (model_key, group["canonical_field_code"])
            if allowed_by_model_field[pair] == 0:
                zero_allowed_pairs.append(pair)

    # Do not call these errors because some fields legitimately do not apply to every model.
    for pair in zero_allowed_pairs[:100]:
        issues.append(
            ApplicabilityIssue(
                "Info",
                "NO_ALLOWED_OPTION_FOR_MODEL_FIELD",
                f"{pair[0]} / {pair[1]} has no STD/X option in Pump Options.",
                f"{pair[0]}|{pair[1]}",
            )
        )

    error_count = sum(1 for issue in issues if issue.severity == "Error")
    warning_count = sum(1 for issue in issues if issue.severity == "Warning")
    info_count = sum(1 for issue in issues if issue.severity == "Info")

    ready_candidates = [
        row for row in candidates if row["publication_status"] == "READY"
    ]

    summary = {
        "family_code": profile["family_code"].upper(),
        "applicability_version": profile["version"],
        "source_reconciliation_version": reconciliation.get(
            "reconciliation_version"
        ),
        "workbook_reads": 1,
        "rev2_workbook_reads": 0,
        "pump_options_model_rows": len(model_rows),
        "accepted_model_count": len(accepted_models),
        "blocked_model_count": len(blocked_models),
        "matrix_group_count": len(groups),
        "compiled_group_count": len(compiled_groups),
        "candidate_count": len(candidates),
        "ready_candidate_count": len(ready_candidates),
        "review_record_count": len(reviews),
        "allowed_candidate_count": sum(1 for row in candidates if row["is_allowed"]),
        "standard_candidate_count": sum(1 for row in candidates if row["is_standard"]),
        "marker_counts": dict(sorted(marker_counts.items())),
        "status_counts": dict(sorted(status_counts.items())),
        "candidate_counts_by_field": dict(sorted(field_counts.items())),
        "allowed_counts_by_field": dict(sorted(allowed_counts.items())),
        "model_classification_counts": dict(
            sorted(
                Counter(
                    model["reconciliation_classification"]
                    for model in accepted_models.values()
                ).items()
            )
        ),
        "zero_allowed_model_field_pair_count": len(zero_allowed_pairs),
        "issue_count": len(issues),
        "error_count": error_count,
        "warning_count": warning_count,
        "info_count": info_count,
        "publication_gate": "PASS" if error_count == 0 else "FAIL",
        "marker_policy": {
            key: value for key, value in marker_semantics.items()
        },
    }

    return {
        "summary": summary,
        "compiled_groups": [
            {
                key: value
                for key, value in group.items()
                if key != "aligned_options"
            }
            | {
                "aligned_options": [
                    {
                        key: value
                        for key, value in option.items()
                        if key != "column_index"
                    }
                    for option in group["aligned_options"]
                ]
            }
            for group in compiled_groups
        ],
        "candidates": candidates,
        "ready_candidates": ready_candidates,
        "reviews": reviews,
        "issues": [asdict(issue) for issue in issues],
    }


def _flatten_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    if not rows:
        path.write_text("", encoding="utf-8")
        return

    fieldnames: list[str] = []
    seen: set[str] = set()

    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                fieldnames.append(key)

    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    key: _flatten_value(row.get(key))
                    for key in fieldnames
                }
            )


def save_applicability_outputs(
    report: dict[str, Any],
    export_dir: Path,
) -> dict[str, Path]:
    export_dir.mkdir(parents=True, exist_ok=True)

    summary_path = export_dir / "m023_dean_applicability_summary.json"
    candidates_path = export_dir / "m023_dean_applicability_candidates.csv"
    ready_path = export_dir / "m023_dean_applicability_ready.csv"
    reviews_path = export_dir / "m023_dean_applicability_review.csv"
    issues_path = export_dir / "m023_dean_applicability_issues.csv"
    groups_path = export_dir / "m023_dean_applicability_groups.json"
    marker_path = export_dir / "m023_dean_applicability_marker_summary.csv"

    summary_path.write_text(
        json.dumps(report["summary"], indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    groups_path.write_text(
        json.dumps(report["compiled_groups"], indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    _write_csv(candidates_path, report["candidates"])
    _write_csv(ready_path, report["ready_candidates"])
    _write_csv(reviews_path, report["reviews"])
    _write_csv(issues_path, report["issues"])

    marker_rows = []
    semantics = report["summary"]["marker_policy"]
    for marker, count in report["summary"]["marker_counts"].items():
        rule = semantics.get(marker)
        marker_rows.append(
            {
                "marker": marker,
                "count": count,
                "known": rule is not None,
                "status": rule["status"] if rule else "REVIEW",
                "is_allowed": rule["is_allowed"] if rule else False,
                "is_standard": rule["is_standard"] if rule else False,
                "publication_status": (
                    rule["publication_status"] if rule else "BLOCKED_REVIEW"
                ),
            }
        )
    _write_csv(marker_path, marker_rows)

    return {
        "summary": summary_path,
        "candidates": candidates_path,
        "ready": ready_path,
        "reviews": reviews_path,
        "issues": issues_path,
        "groups": groups_path,
        "markers": marker_path,
    }
