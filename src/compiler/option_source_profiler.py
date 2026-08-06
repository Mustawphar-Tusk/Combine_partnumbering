from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


HEX_HEADER_PATTERNS = (
    "hex",
    "code",
    "smart number",
    "smartnumber",
    "part number code",
)
DESCRIPTION_HEADER_PATTERNS = (
    "description",
    "option",
    "selection",
    "attribute",
    "value",
    "material",
    "series",
    "size",
)
INVALID_VALUES = {"ERR", "TBD__", "____"}


@dataclass(frozen=True)
class NamedRangeCandidate:
    family_code: str
    workbook_role: str
    workbook_name: str
    name: str
    scope: str
    refers_to: str
    resolved_sheet: str | None
    resolved_range: str | None
    row_count: int | None
    column_count: int | None


@dataclass(frozen=True)
class ValidationCandidate:
    family_code: str
    workbook_role: str
    workbook_name: str
    worksheet_name: str
    target_range: str
    validation_type: str | None
    formula1: str | None
    formula2: str | None
    allow_blank: bool | None
    source_kind: str
    resolved_source: str | None


@dataclass(frozen=True)
class OptionTableCandidate:
    family_code: str
    workbook_role: str
    workbook_name: str
    worksheet_name: str
    header_row: int
    first_data_row: int
    last_data_row: int
    description_column: str
    code_column: str | None
    hex_column: str | None
    confidence: str
    reason: str


@dataclass(frozen=True)
class OptionSourceIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    worksheet_name: str | None = None
    source_reference: str | None = None


@dataclass(frozen=True)
class OptionSourceProfileReport:
    named_range_count: int
    validation_count: int
    option_table_count: int
    issue_count: int
    named_ranges: tuple[NamedRangeCandidate, ...]
    validations: tuple[ValidationCandidate, ...]
    option_tables: tuple[OptionTableCandidate, ...]
    issues: tuple[OptionSourceIssue, ...]


def _resolve_workbook(
    project_root: Path,
    discovery: dict[str, Any],
    family_code: str,
    workbook_role: str,
) -> Path:
    matches = [
        row
        for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == workbook_role
        and row.get("discovery_status") == "discovered"
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected exactly one workbook for {family_code}/{workbook_role}; "
            f"found {len(matches)}."
        )
    return project_root / matches[0]["relative_path"]


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def _column_letter(cell) -> str:
    return re.sub(r"\d+", "", cell.coordinate)


def _matches_any(value: str, patterns: tuple[str, ...]) -> bool:
    return any(pattern in value for pattern in patterns)


def _defined_names(workbook, family_code: str, role: str, name: str):
    candidates: list[NamedRangeCandidate] = []

    for defined_name in workbook.defined_names.values():
        scope = (
            "workbook"
            if defined_name.localSheetId is None
            else f"worksheet:{defined_name.localSheetId}"
        )
        refers_to = defined_name.attr_text or ""
        resolved_sheet = None
        resolved_range = None
        row_count = None
        column_count = None

        try:
            destinations = list(defined_name.destinations)
        except Exception:
            destinations = []

        if len(destinations) == 1:
            resolved_sheet, resolved_range = destinations[0]
            try:
                min_col, min_row, max_col, max_row = range_boundaries(
                    resolved_range.replace("$", "")
                )
                row_count = max_row - min_row + 1
                column_count = max_col - min_col + 1
            except Exception:
                pass

        candidates.append(
            NamedRangeCandidate(
                family_code=family_code,
                workbook_role=role,
                workbook_name=name,
                name=defined_name.name,
                scope=scope,
                refers_to=refers_to,
                resolved_sheet=resolved_sheet,
                resolved_range=resolved_range,
                row_count=row_count,
                column_count=column_count,
            )
        )

    return candidates


def _validation_source_kind(formula: str | None) -> tuple[str, str | None]:
    if not formula:
        return "none", None

    stripped = formula.strip()

    if stripped.startswith('"') and stripped.endswith('"'):
        return "inline_list", stripped

    if stripped.startswith("="):
        stripped = stripped[1:]

    if "!" in stripped:
        return "range_reference", stripped

    if re.fullmatch(r"[A-Za-z_\\][A-Za-z0-9_.\\]*", stripped):
        return "defined_name", stripped

    if stripped.upper().startswith(("INDIRECT(", "OFFSET(", "INDEX(")):
        return "formula_driven", stripped

    return "formula_or_expression", stripped


def _validations(workbook, family_code: str, role: str, workbook_name: str):
    rows: list[ValidationCandidate] = []

    for worksheet in workbook.worksheets:
        collection = worksheet.data_validations
        if collection is None:
            continue

        for validation in collection.dataValidation:
            source_kind, resolved = _validation_source_kind(validation.formula1)

            rows.append(
                ValidationCandidate(
                    family_code=family_code,
                    workbook_role=role,
                    workbook_name=workbook_name,
                    worksheet_name=worksheet.title,
                    target_range=str(validation.sqref),
                    validation_type=validation.type,
                    formula1=validation.formula1,
                    formula2=validation.formula2,
                    allow_blank=validation.allow_blank,
                    source_kind=source_kind,
                    resolved_source=resolved,
                )
            )

    return rows


def _candidate_tables(
    workbook,
    family_code: str,
    role: str,
    workbook_name: str,
    scan_rows: int = 30,
    scan_columns: int = 100,
):
    results: list[OptionTableCandidate] = []
    seen: set[tuple[str, int, str]] = set()

    for worksheet in workbook.worksheets:
        max_row = min(worksheet.max_row or 0, scan_rows)
        max_column = min(worksheet.max_column or 0, scan_columns)

        for row in range(1, max_row + 1):
            headers = []
            for column in range(1, max_column + 1):
                cell = worksheet.cell(row=row, column=column)
                header = _normalize_header(cell.value)
                if header:
                    headers.append((column, cell, header))

            if not headers:
                continue

            description_cells = [
                (column, cell, header)
                for column, cell, header in headers
                if _matches_any(header, DESCRIPTION_HEADER_PATTERNS)
            ]
            code_cells = [
                (column, cell, header)
                for column, cell, header in headers
                if "code" in header and "description" not in header
            ]
            hex_cells = [
                (column, cell, header)
                for column, cell, header in headers
                if _matches_any(header, HEX_HEADER_PATTERNS)
            ]

            for _, description_cell, description_header in description_cells:
                key = (
                    worksheet.title,
                    row,
                    description_cell.coordinate,
                )
                if key in seen:
                    continue
                seen.add(key)

                code_cell = code_cells[0][1] if code_cells else None
                hex_cell = hex_cells[0][1] if hex_cells else None

                nonblank_rows = []
                for data_row in range(row + 1, min(worksheet.max_row, row + 500) + 1):
                    value = worksheet.cell(
                        row=data_row,
                        column=description_cell.column,
                    ).value
                    if value is not None and str(value).strip():
                        if str(value).strip().upper() not in INVALID_VALUES:
                            nonblank_rows.append(data_row)

                if not nonblank_rows:
                    continue

                confidence = "high" if hex_cell or code_cell else "medium"
                reason_parts = [
                    f"Description-like header '{description_header}'."
                ]
                if code_cell:
                    reason_parts.append(
                        f"Code-like header at {code_cell.coordinate}."
                    )
                if hex_cell:
                    reason_parts.append(
                        f"Hex/code-like header at {hex_cell.coordinate}."
                    )

                results.append(
                    OptionTableCandidate(
                        family_code=family_code,
                        workbook_role=role,
                        workbook_name=workbook_name,
                        worksheet_name=worksheet.title,
                        header_row=row,
                        first_data_row=min(nonblank_rows),
                        last_data_row=max(nonblank_rows),
                        description_column=_column_letter(description_cell),
                        code_column=(
                            _column_letter(code_cell) if code_cell else None
                        ),
                        hex_column=(
                            _column_letter(hex_cell) if hex_cell else None
                        ),
                        confidence=confidence,
                        reason=" ".join(reason_parts),
                    )
                )

    return results


def profile_option_sources(
    project_root: Path,
    discovery_path: Path,
    profile_path: Path,
) -> OptionSourceProfileReport:
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))
    configuration = json.loads(profile_path.read_text(encoding="utf-8"))

    named_ranges: list[NamedRangeCandidate] = []
    validations: list[ValidationCandidate] = []
    option_tables: list[OptionTableCandidate] = []
    issues: list[OptionSourceIssue] = []

    for item in configuration["workbooks"]:
        family_code = item["family_code"].upper()
        role = item["workbook_role"]

        path = _resolve_workbook(
            project_root,
            discovery,
            family_code,
            role,
        )

        workbook = load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_vba=path.suffix.casefold() in {".xlsm", ".xltm"},
            keep_links=True,
        )

        try:
            named_ranges.extend(
                _defined_names(
                    workbook,
                    family_code,
                    role,
                    path.name,
                )
            )
            validations.extend(
                _validations(
                    workbook,
                    family_code,
                    role,
                    path.name,
                )
            )
            option_tables.extend(
                _candidate_tables(
                    workbook,
                    family_code,
                    role,
                    path.name,
                    scan_rows=int(
                        configuration.get("table_scan_rows", 30)
                    ),
                    scan_columns=int(
                        configuration.get("table_scan_columns", 100)
                    ),
                )
            )
        except Exception as exc:
            issues.append(
                OptionSourceIssue(
                    family_code=family_code,
                    severity="Error",
                    issue_code="WORKBOOK_PROFILE_FAILED",
                    message=f"{path.name}: {exc}",
                )
            )
        finally:
            workbook.close()

    return OptionSourceProfileReport(
        named_range_count=len(named_ranges),
        validation_count=len(validations),
        option_table_count=len(option_tables),
        issue_count=len(issues),
        named_ranges=tuple(named_ranges),
        validations=tuple(validations),
        option_tables=tuple(option_tables),
        issues=tuple(issues),
    )


def save_json(report: OptionSourceProfileReport, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(report: OptionSourceProfileReport, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    _write_csv(
        directory / "option_source_named_ranges.csv",
        [asdict(row) for row in report.named_ranges],
    )
    _write_csv(
        directory / "option_source_validations.csv",
        [asdict(row) for row in report.validations],
    )
    _write_csv(
        directory / "option_source_table_candidates.csv",
        [asdict(row) for row in report.option_tables],
    )
    _write_csv(
        directory / "option_source_profile_issues.csv",
        [asdict(row) for row in report.issues],
    )


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
