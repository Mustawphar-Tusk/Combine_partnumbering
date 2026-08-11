from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


@dataclass(frozen=True)
class PriceCandidateCondition:
    sequence_no: int
    field_code: str
    comparison_operator: str
    comparison_value: str | None
    source_value: str | None = None


@dataclass(frozen=True)
class PriceCandidate:
    family_code: str
    component_code: str
    series_code: str
    source_series_code: str
    size_value: str
    source_size_value: str
    option_field_code: str
    option_value: str
    source_option_value: str
    amount: float | None
    pricing_status: str
    source_price_value: Any
    currency_code: str
    workbook_name: str
    worksheet_name: str
    table_name: str
    source_cell: str
    conditions: tuple[
        PriceCandidateCondition,
        ...
    ] = ()


@dataclass(frozen=True)
class PricingIssue:
    family_code: str
    component_code: str
    series_code: str | None
    severity: str
    issue_code: str
    message: str
    source_reference: str | None = None


@dataclass(frozen=True)
class PricingCompilationReport:
    candidate_count: int
    issue_count: int
    candidates: tuple[PriceCandidate, ...]
    issues: tuple[PricingIssue, ...]


def _text(value: Any) -> str | None:
    if value is None:
        return None

    value = str(value).strip()
    return value or None


def _transform_row_value(
    value: str,
    transform_code: str | None,
) -> str:

    if not transform_code:
        return value.strip().lower()

    if transform_code == "STRIP_PAREN_SUFFIX":
        value = value.split(" (", 1)[0]
        return value.strip().lower()

    raise ValueError(
        f"Unsupported pricing row transform: "
        f"{transform_code}"
    )


def compile_base_pump_pricing(
    workbook_path: Path,
    profile_path: Path,
) -> PricingCompilationReport:

    profile = json.loads(
        profile_path.read_text(
            encoding="utf-8-sig",
        )
    )

    family_code = (
        profile["family_code"]
        .strip()
        .upper()
    )

    component_code = profile[
        "component_code"
    ]

    worksheet_name = profile[
        "worksheet_name"
    ]

    simple_series = set(
        profile["simple_matrix_series"]
    )

    ignored_columns = {
        str(value).strip()
        for value in profile[
            "ignored_columns"
        ]
    }

    row_key_default = profile.get(
        "row_key",
        "SIZE",
    )

    row_key_by_series = profile.get(
        "row_key_by_series",
        {},
    )

    row_value_transform_by_series = (
        profile.get(
            "row_value_transform_by_series",
            {},
        )
    )

    material_equivalences = (
        profile
        .get(
            "value_equivalences",
            {},
        )
        .get(
            "PUMP_MATERIAL",
            {},
        )
    )

    series_equivalences = profile.get(
        "series_value_equivalences",
        {},
    )

    price_value_status_map = {
        str(key).strip().upper(): value
        for key, value in profile.get(
            "price_value_status_map",
            {},
        ).items()
    }

    candidates: list[PriceCandidate] = []
    issues: list[PricingIssue] = []

    workbook = load_workbook(
        workbook_path,
        data_only=True,
        read_only=False,
        keep_vba=True,
    )

    try:
        worksheet = workbook[
            worksheet_name
        ]

        for (
            series_code,
            table_name,
        ) in profile[
            "series_tables"
        ].items():

            if series_code not in simple_series:
                continue

            if table_name not in worksheet.tables:
                issues.append(
                    PricingIssue(
                        family_code=family_code,
                        component_code=component_code,
                        series_code=series_code,
                        severity="Error",
                        issue_code="PRICE_TABLE_NOT_FOUND",
                        message=(
                            f"Pricing table "
                            f"{table_name} "
                            f"was not found."
                        ),
                    )
                )
                continue

            table = worksheet.tables[
                table_name
            ]

            (
                min_col,
                min_row,
                max_col,
                max_row,
            ) = range_boundaries(
                table.ref
            )

            headers = {
                col: _text(
                    worksheet.cell(
                        min_row,
                        col,
                    ).value
                )
                for col in range(
                    min_col,
                    max_col + 1,
                )
            }

            row_key = row_key_by_series.get(
                series_code,
                row_key_default,
            )

            row_key_column = next(
                (
                    col
                    for col, header
                    in headers.items()
                    if header == row_key
                ),
                None,
            )

            if row_key_column is None:
                issues.append(
                    PricingIssue(
                        family_code=family_code,
                        component_code=component_code,
                        series_code=series_code,
                        severity="Error",
                        issue_code=(
                            "ROW_KEY_COLUMN_NOT_FOUND"
                        ),
                        message=(
                            f"{table_name} has no "
                            f"{row_key!r} column."
                        ),
                        source_reference=(
                            f"{worksheet_name}!"
                            f"{table.ref}"
                        ),
                    )
                )
                continue

            transform_code = (
                row_value_transform_by_series
                .get(series_code)
            )

            for row in range(
                min_row + 1,
                max_row + 1,
            ):
                source_size = _text(
                    worksheet.cell(
                        row,
                        row_key_column,
                    ).value
                )

                if not source_size:
                    continue

                try:
                    canonical_size = (
                        _transform_row_value(
                            source_size,
                            transform_code,
                        )
                    )
                except ValueError as exc:
                    issues.append(
                        PricingIssue(
                            family_code=family_code,
                            component_code=component_code,
                            series_code=series_code,
                            severity="Error",
                            issue_code=(
                                "ROW_VALUE_TRANSFORM_FAILED"
                            ),
                            message=str(exc),
                            source_reference=(
                                f"{worksheet_name}!"
                                f"{worksheet.cell(row, row_key_column).coordinate}"
                            ),
                        )
                    )
                    continue

                for (
                    col,
                    header,
                ) in headers.items():

                    if not header:
                        continue

                    if header in ignored_columns:
                        continue

                    raw_value = worksheet.cell(
                        row,
                        col,
                    ).value

                    source_material = header

                    canonical_material = (
                        material_equivalences.get(
                            source_material,
                            source_material,
                        )
                    )

                    amount: float | None
                    pricing_status: str

                    if isinstance(
                        raw_value,
                        (int, float),
                    ):
                        amount = float(raw_value)
                        pricing_status = "found"

                    else:
                        source_text = _text(raw_value)

                        if not source_text:
                            continue

                        mapped_status = (
                            price_value_status_map.get(
                                source_text.upper()
                            )
                        )

                        if not mapped_status:
                            continue

                        amount = None
                        pricing_status = mapped_status

                    source_cell = worksheet.cell(
                        row,
                        col,
                    ).coordinate

                    candidates.append(
                        PriceCandidate(
                            family_code=family_code,
                            component_code=component_code,
                            series_code=(
                                series_equivalences.get(
                                    series_code,
                                    series_code,
                                )
                            ),
                            source_series_code=series_code,
                            size_value=canonical_size,
                            source_size_value=source_size,
                            option_field_code=(
                                "PUMP_MATERIAL"
                            ),
                            option_value=(
                                canonical_material
                            ),
                            source_option_value=(
                                source_material
                            ),
                            amount=amount,
                            pricing_status=(
                                pricing_status
                            ),
                            source_price_value=(
                                raw_value
                            ),
                            currency_code="USD",
                            workbook_name=(
                                workbook_path.name
                            ),
                            worksheet_name=(
                                worksheet_name
                            ),
                            table_name=table_name,
                            source_cell=source_cell,
                            conditions=(
                                PriceCandidateCondition(
                                    sequence_no=1,
                                    field_code="SIZE",
                                    comparison_operator="EQ",
                                    comparison_value=canonical_size,
                                    source_value=source_size,
                                ),
                                PriceCandidateCondition(
                                    sequence_no=2,
                                    field_code="PUMP_MATERIAL",
                                    comparison_operator="EQ",
                                    comparison_value=canonical_material,
                                    source_value=source_material,
                                ),
                            ),
                        )
                    )

        return PricingCompilationReport(
            candidate_count=len(candidates),
            issue_count=len(issues),
            candidates=tuple(candidates),
            issues=tuple(issues),
        )

    finally:
        workbook.close()


def save_report(
    report: PricingCompilationReport,
    output_path: Path,
) -> None:

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_path.write_text(
        json.dumps(
            asdict(report),
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
