from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


ALPHABET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
INVALID_VALUES = {"ERR", "TBD__", "____"}


@dataclass(frozen=True)
class SegmentCombinationCandidate:
    family_code: str
    workbook_role: str
    workbook_name: str
    worksheet_name: str
    segment_code: str
    segment_name: str
    source_row: int
    source_id: int
    segment_value: str
    expected_width: int
    combination_key: str
    selections_json: str
    source_cells_json: str
    source_profile: str


@dataclass(frozen=True)
class SegmentCombinationIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    worksheet_name: str | None = None
    source_reference: str | None = None


@dataclass(frozen=True)
class SegmentCombinationReport:
    combination_count: int
    segment_count: int
    issue_count: int
    combinations: tuple[SegmentCombinationCandidate, ...]
    issues: tuple[SegmentCombinationIssue, ...]


def to_base36(value: int) -> str:
    if value < 0:
        raise ValueError("Base-36 input cannot be negative.")
    if value == 0:
        return "0"

    digits: list[str] = []
    while value:
        value, remainder = divmod(value, 36)
        digits.append(ALPHABET[remainder])

    return "".join(reversed(digits))


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_workbook(
    project_root: Path,
    discovery: dict[str, Any],
    family_code: str,
    workbook_role: str,
) -> Path:
    matches = [
        row
        for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == workbook_role
        and row.get("discovery_status") == "discovered"
    ]

    if len(matches) != 1:
        raise ValueError(
            f"Expected one workbook for {family_code}/{workbook_role}; "
            f"found {len(matches)}."
        )

    return project_root / matches[0]["relative_path"]


def _column_value(
    row_cells: tuple,
    column_letter: str,
    minimum_column: int,
) -> Any:
    absolute_column = column_index_from_string(column_letter)
    relative_index = absolute_column - minimum_column

    if relative_index < 0 or relative_index >= len(row_cells):
        return None

    return row_cells[relative_index].value


def _column_coordinate(
    row_cells: tuple,
    column_letter: str,
    minimum_column: int,
    row_number: int,
) -> str:
    absolute_column = column_index_from_string(column_letter)
    relative_index = absolute_column - minimum_column

    if 0 <= relative_index < len(row_cells):
        return row_cells[relative_index].coordinate

    return f"{column_letter}{row_number}"


def compile_segment_combinations(
    project_root: Path,
    discovery_path: Path,
    profile_path: Path,
) -> SegmentCombinationReport:
    discovery = json.loads(
        discovery_path.read_text(encoding="utf-8")
    )
    profile = json.loads(
        profile_path.read_text(encoding="utf-8")
    )

    family_code = profile["family_code"].upper()
    workbook_role = profile["workbook_role"]

    workbook_path = _resolve_workbook(
        project_root,
        discovery,
        family_code,
        workbook_role,
    )

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=False,
        keep_links=False,
    )

    combinations: list[SegmentCombinationCandidate] = []
    issues: list[SegmentCombinationIssue] = []

    seen_ids: set[tuple[str, int]] = set()
    seen_keys: set[tuple[str, str]] = set()

    try:
        for segment in profile["segments"]:
            worksheet_name = segment["worksheet_name"]
            segment_code = segment["segment_code"]
            segment_name = segment["segment_name"]
            expected_width = int(segment["expected_width"])
            id_column = segment["id_column"]
            key_column = segment.get("combination_key_column")
            row_from = int(segment["row_from"])
            row_to = int(segment.get("row_to", 100000))
            selection_columns = segment["selection_columns"]

            max_consecutive_blank_ids = int(
                segment.get("max_consecutive_blank_ids", 100)
            )

            if worksheet_name not in workbook.sheetnames:
                issues.append(
                    SegmentCombinationIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="WORKSHEET_NOT_FOUND",
                        message=f"Worksheet '{worksheet_name}' was not found.",
                        worksheet_name=worksheet_name,
                    )
                )
                continue

            worksheet = workbook[worksheet_name]

            required_columns = [
                column_index_from_string(id_column),
                *(
                    [column_index_from_string(key_column)]
                    if key_column
                    else []
                ),
                *[
                    column_index_from_string(column)
                    for column in selection_columns.values()
                ],
            ]

            minimum_column = min(required_columns)
            maximum_column = max(required_columns)
            blank_id_count = 0

            for row_number, row_cells in enumerate(
                worksheet.iter_rows(
                    min_row=row_from,
                    max_row=row_to,
                    min_col=minimum_column,
                    max_col=maximum_column,
                    values_only=False,
                ),
                start=row_from,
            ):
                raw_id = _column_value(
                    row_cells,
                    id_column,
                    minimum_column,
                )

                if raw_id is None or str(raw_id).strip() == "":
                    blank_id_count += 1

                    if blank_id_count >= max_consecutive_blank_ids:
                        break

                    continue

                blank_id_count = 0

                try:
                    source_id = int(raw_id)
                except (TypeError, ValueError):
                    issues.append(
                        SegmentCombinationIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="INVALID_SOURCE_ID",
                            message=f"Source ID '{raw_id}' is not an integer.",
                            worksheet_name=worksheet_name,
                            source_reference=f"{id_column}{row_number}",
                        )
                    )
                    continue

                selections: dict[str, str] = {}
                source_cells: dict[str, str] = {}
                invalid = False

                for field_code, column in selection_columns.items():
                    raw_value = _column_value(
                        row_cells,
                        column,
                        minimum_column,
                    )
                    value = _text(raw_value)

                    if value is None:
                        invalid = True
                        issues.append(
                            SegmentCombinationIssue(
                                family_code=family_code,
                                severity="Warning",
                                issue_code="MISSING_SELECTION_VALUE",
                                message=(
                                    f"Missing value for {field_code} "
                                    f"in {segment_code} row {row_number}."
                                ),
                                worksheet_name=worksheet_name,
                                source_reference=f"{column}{row_number}",
                            )
                        )
                        break

                    if value.upper() in INVALID_VALUES:
                        invalid = True
                        break

                    selections[field_code] = value
                    source_cells[field_code] = _column_coordinate(
                        row_cells,
                        column,
                        minimum_column,
                        row_number,
                    )

                if invalid:
                    continue

                combination_key = None

                if key_column:
                    raw_key = _column_value(
                        row_cells,
                        key_column,
                        minimum_column,
                    )

                    if isinstance(raw_key, str) and raw_key.startswith("="):
                        raw_key = None

                    combination_key = _text(raw_key)

                if not combination_key:
                    combination_key = "|".join(
                        selections[field_code]
                        for field_code in selection_columns
                    )

                id_key = (segment_code, source_id)
                combination_key_id = (segment_code, combination_key)

                if id_key in seen_ids:
                    issues.append(
                        SegmentCombinationIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="DUPLICATE_SEGMENT_ID",
                            message=(
                                f"Duplicate source ID {source_id} "
                                f"for segment {segment_code}."
                            ),
                            worksheet_name=worksheet_name,
                            source_reference=f"{id_column}{row_number}",
                        )
                    )
                    continue

                if combination_key_id in seen_keys:
                    issues.append(
                        SegmentCombinationIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="DUPLICATE_COMBINATION_KEY",
                            message=(
                                f"Duplicate combination key for segment "
                                f"{segment_code}."
                            ),
                            worksheet_name=worksheet_name,
                            source_reference=(
                                f"{key_column}{row_number}"
                                if key_column
                                else str(row_number)
                            ),
                        )
                    )
                    continue

                segment_value = to_base36(source_id).zfill(expected_width)

                if len(segment_value) > expected_width:
                    issues.append(
                        SegmentCombinationIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="SEGMENT_WIDTH_EXCEEDED",
                            message=(
                                f"Segment value '{segment_value}' "
                                f"exceeds width {expected_width}."
                            ),
                            worksheet_name=worksheet_name,
                            source_reference=f"{id_column}{row_number}",
                        )
                    )
                    continue

                seen_ids.add(id_key)
                seen_keys.add(combination_key_id)

                combinations.append(
                    SegmentCombinationCandidate(
                        family_code=family_code,
                        workbook_role=workbook_role,
                        workbook_name=workbook_path.name,
                        worksheet_name=worksheet_name,
                        segment_code=segment_code,
                        segment_name=segment_name,
                        source_row=row_number,
                        source_id=source_id,
                        segment_value=segment_value,
                        expected_width=expected_width,
                        combination_key=combination_key,
                        selections_json=json.dumps(
                            selections,
                            ensure_ascii=False,
                            separators=(",", ":"),
                        ),
                        source_cells_json=json.dumps(
                            source_cells,
                            separators=(",", ":"),
                        ),
                        source_profile=str(profile_path),
                    )
                )

    finally:
        workbook.close()

    return SegmentCombinationReport(
        combination_count=len(combinations),
        segment_count=len(
            {item.segment_code for item in combinations}
        ),
        issue_count=len(issues),
        combinations=tuple(combinations),
        issues=tuple(issues),
    )


def save_json(
    report: SegmentCombinationReport,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(
    report: SegmentCombinationReport,
    directory: Path,
) -> None:
    directory.mkdir(parents=True, exist_ok=True)

    _write_csv(
        directory / "segment_combination_candidates.csv",
        [asdict(row) for row in report.combinations],
    )
    _write_csv(
        directory / "segment_combination_issues.csv",
        [asdict(row) for row in report.issues],
    )


def _write_csv(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(rows[0]),
        )
        writer.writeheader()
        writer.writerows(rows)