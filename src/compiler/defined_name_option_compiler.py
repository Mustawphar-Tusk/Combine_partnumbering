from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


INVALID_VALUES = {"ERR", "TBD__", "____"}


@dataclass(frozen=True)
class DefinedNameOptionCandidate:
    family_code: str
    workbook_role: str
    workbook_name: str
    input_worksheet: str
    input_range: str
    field_code: str
    field_name: str
    defined_name: str
    source_worksheet: str
    source_range: str
    option_code: str
    option_description: str
    source_cell: str
    display_order: int
    confidence: str
    source_profile: str


@dataclass(frozen=True)
class DefinedNameOptionIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    defined_name: str | None = None
    source_reference: str | None = None


@dataclass(frozen=True)
class DefinedNameOptionReport:
    option_count: int
    field_count: int
    issue_count: int
    options: tuple[DefinedNameOptionCandidate, ...]
    issues: tuple[DefinedNameOptionIssue, ...]


def normalize_code(value: str) -> str:
    code = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_").upper()
    return code or "UNNAMED"


def _resolve_workbook(
    project_root: Path,
    discovery: dict[str, Any],
    family_code: str,
    workbook_role: str,
) -> Path:
    matches = [
        row
        for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == workbook_role
        and row.get("discovery_status") == "discovered"
    ]

    if len(matches) != 1:
        raise ValueError(
            f"Expected one workbook for {family_code}/{workbook_role}; "
            f"found {len(matches)}."
        )

    return project_root / matches[0]["relative_path"]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_defined_name(workbook, name: str):
    matches = [
        item
        for item in workbook.defined_names.values()
        if item.name == name
    ]

    resolved = []
    unresolved = []

    for item in matches:
        try:
            destinations = list(item.destinations)
        except Exception:
            destinations = []

        if len(destinations) == 1:
            resolved.append((item, destinations[0]))
        else:
            unresolved.append(item)

    return resolved, unresolved


def compile_defined_name_options(
    project_root: Path,
    discovery_path: Path,
    profile_path: Path,
) -> DefinedNameOptionReport:
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))
    profile = json.loads(profile_path.read_text(encoding="utf-8"))

    family_code = profile["family_code"].upper()
    workbook_role = profile["workbook_role"]
    input_worksheet = profile["input_worksheet"]
    source_profile = str(profile_path)

    workbook_path = _resolve_workbook(
        project_root,
        discovery,
        family_code,
        workbook_role,
    )

    workbook = load_workbook(
        workbook_path,
        read_only=False,
        data_only=True,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        keep_links=True,
    )

    options: list[DefinedNameOptionCandidate] = []
    issues: list[DefinedNameOptionIssue] = []
    seen_options: set[tuple[str, str]] = set()
    seen_fields: set[str] = set()

    try:
        for mapping in profile["mappings"]:
            defined_name = mapping["defined_name"]
            field_code = mapping.get(
                "field_code",
                normalize_code(defined_name),
            )
            field_name = mapping.get(
                "field_name",
                defined_name.replace("_", " "),
            )
            input_range = mapping["input_range"]

            resolved, unresolved = _resolve_defined_name(
                workbook,
                defined_name,
            )

            if unresolved:
                issues.append(
                    DefinedNameOptionIssue(
                        family_code=family_code,
                        severity="Warning",
                        issue_code="UNRESOLVED_DEFINED_NAME_RECORD",
                        message=(
                            f"Defined name '{defined_name}' has "
                            f"{len(unresolved)} unresolved record(s)."
                        ),
                        defined_name=defined_name,
                    )
                )

            unique_destinations = {
                (
                    destination[0],
                    destination[1].replace("$", ""),
                )
                for _, destination in resolved
            }

            if not unique_destinations:
                issues.append(
                    DefinedNameOptionIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="DEFINED_NAME_NOT_RESOLVED",
                        message=(
                            f"Defined name '{defined_name}' could not "
                            "be resolved to a worksheet range."
                        ),
                        defined_name=defined_name,
                    )
                )
                continue

            if len(unique_destinations) > 1:
                issues.append(
                    DefinedNameOptionIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="DEFINED_NAME_CONFLICT",
                        message=(
                            f"Defined name '{defined_name}' resolves to "
                            f"multiple ranges: {sorted(unique_destinations)}."
                        ),
                        defined_name=defined_name,
                    )
                )
                continue

            source_worksheet, source_range = next(
                iter(unique_destinations)
            )

            if source_worksheet not in workbook.sheetnames:
                issues.append(
                    DefinedNameOptionIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="SOURCE_WORKSHEET_NOT_FOUND",
                        message=(
                            f"Source worksheet '{source_worksheet}' "
                            f"for '{defined_name}' was not found."
                        ),
                        defined_name=defined_name,
                        source_reference=source_range,
                    )
                )
                continue

            min_col, min_row, max_col, max_row = range_boundaries(
                source_range
            )

            if min_col != max_col:
                issues.append(
                    DefinedNameOptionIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="MULTI_COLUMN_OPTION_RANGE",
                        message=(
                            f"Defined name '{defined_name}' resolves to "
                            f"a multi-column range '{source_range}'."
                        ),
                        defined_name=defined_name,
                        source_reference=(
                            f"{source_worksheet}!{source_range}"
                        ),
                    )
                )
                continue

            worksheet = workbook[source_worksheet]
            display_order = 0

            for row in range(min_row, max_row + 1):
                cell = worksheet.cell(row=row, column=min_col)
                description = _text(cell.value)

                if not description:
                    continue
                if description.upper() in INVALID_VALUES:
                    continue

                option_code = normalize_code(description)
                dedupe_key = (field_code, option_code)

                if dedupe_key in seen_options:
                    continue

                seen_options.add(dedupe_key)
                seen_fields.add(field_code)
                display_order += 10

                options.append(
                    DefinedNameOptionCandidate(
                        family_code=family_code,
                        workbook_role=workbook_role,
                        workbook_name=workbook_path.name,
                        input_worksheet=input_worksheet,
                        input_range=input_range,
                        field_code=field_code,
                        field_name=field_name,
                        defined_name=defined_name,
                        source_worksheet=source_worksheet,
                        source_range=source_range,
                        option_code=option_code,
                        option_description=description,
                        source_cell=cell.coordinate,
                        display_order=display_order,
                        confidence="high",
                        source_profile=source_profile,
                    )
                )

    finally:
        workbook.close()

    return DefinedNameOptionReport(
        option_count=len(options),
        field_count=len(seen_fields),
        issue_count=len(issues),
        options=tuple(options),
        issues=tuple(issues),
    )


def save_json(
    report: DefinedNameOptionReport,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(
    report: DefinedNameOptionReport,
    directory: Path,
) -> None:
    directory.mkdir(parents=True, exist_ok=True)

    _write_csv(
        directory / "fybroc_defined_name_option_candidates.csv",
        [asdict(row) for row in report.options],
    )
    _write_csv(
        directory / "fybroc_defined_name_option_issues.csv",
        [asdict(row) for row in report.issues],
    )


def _write_csv(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(rows[0]),
        )
        writer.writeheader()
        writer.writerows(rows)
