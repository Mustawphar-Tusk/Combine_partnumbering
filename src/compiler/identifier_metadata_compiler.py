from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class PumpModelCandidate:
    family_code: str
    model_identifier: str
    series_code: str
    size_code: str
    base_identifier: str
    source_worksheet: str
    source_row: int
    source_profile: str


@dataclass(frozen=True)
class IdentifierFormatCandidate:
    family_code: str
    identifier_type: str
    format_code: str
    version_number: int
    template: str
    segment_separator: str
    base_identifier_source: str
    source_profile: str


@dataclass(frozen=True)
class IdentifierSegmentCandidate:
    family_code: str
    segment_code: str
    segment_name: str
    assembly_order: int
    expected_width: int | None
    source_field_codes_json: str
    source_profile: str


@dataclass(frozen=True)
class IdentifierMetadataIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    source_reference: str | None = None


@dataclass(frozen=True)
class IdentifierMetadataReport:
    model_count: int
    format_count: int
    segment_count: int
    issue_count: int
    models: tuple[PumpModelCandidate, ...]
    formats: tuple[IdentifierFormatCandidate, ...]
    segments: tuple[IdentifierSegmentCandidate, ...]
    issues: tuple[IdentifierMetadataIssue, ...]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _resolve_workbook(project_root: Path, discovery: dict, family_code: str, role: str) -> Path:
    matches = [
        row for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == role
        and row.get("discovery_status") == "discovered"
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one workbook for {family_code}/{role}; found {len(matches)}."
        )
    return project_root / matches[0]["relative_path"]


def _base_identifier(model_identifier: str, format_profile: dict) -> str:
    prefix = format_profile["base_prefix"]
    remove_prefix = format_profile.get("remove_model_prefix")
    core = model_identifier

    if remove_prefix and core.upper().startswith(remove_prefix.upper()):
        core = core[len(remove_prefix):]

    return f"{prefix}{core}"


def compile_identifier_metadata(
    project_root: Path,
    discovery_path: Path,
    profile_directory: Path,
) -> IdentifierMetadataReport:
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))

    models: list[PumpModelCandidate] = []
    formats: list[IdentifierFormatCandidate] = []
    segments: list[IdentifierSegmentCandidate] = []
    issues: list[IdentifierMetadataIssue] = []

    seen_model_keys: set[tuple[str, str, str]] = set()

    for profile_path in sorted(profile_directory.glob("*.json")):
        profile = json.loads(profile_path.read_text(encoding="utf-8"))
        family_code = profile["family_code"].upper()

        for fmt in profile["identifier_formats"]:
            formats.append(
                IdentifierFormatCandidate(
                    family_code=family_code,
                    identifier_type=fmt["identifier_type"],
                    format_code=fmt["format_code"],
                    version_number=int(fmt.get("version_number", 1)),
                    template=fmt["template"],
                    segment_separator=fmt.get("segment_separator", "-"),
                    base_identifier_source=fmt["base_identifier_source"],
                    source_profile=str(profile_path),
                )
            )

        for segment in profile["segments"]:
            segments.append(
                IdentifierSegmentCandidate(
                    family_code=family_code,
                    segment_code=segment["segment_code"],
                    segment_name=segment["segment_name"],
                    assembly_order=int(segment["assembly_order"]),
                    expected_width=segment.get("expected_width"),
                    source_field_codes_json=json.dumps(
                        segment.get("source_field_codes", []),
                        separators=(",", ":"),
                    ),
                    source_profile=str(profile_path),
                )
            )

        model_source = profile["model_reference"]
        workbook_path = _resolve_workbook(
            project_root,
            discovery,
            family_code,
            model_source["workbook_role"],
        )

        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
            keep_vba=workbook_path.suffix.lower() == ".xlsm",
            keep_links=True,
        )

        try:
            sheet_name = model_source["worksheet_name"]
            if sheet_name not in workbook.sheetnames:
                issues.append(
                    IdentifierMetadataIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="WORKSHEET_NOT_FOUND",
                        message=f"Worksheet '{sheet_name}' was not found.",
                    )
                )
                continue

            worksheet = workbook[sheet_name]
            columns = model_source["columns"]
            row_from = int(model_source["row_from"])
            row_to = int(model_source.get("row_to", worksheet.max_row or row_from))

            for row in range(row_from, row_to + 1):
                series = _text(worksheet[f'{columns["series"]}{row}'].value)
                size = _text(worksheet[f'{columns["size"]}{row}'].value)

                if not series or not size:
                    continue

                source_kind = model_source["model_identifier_source"]
                if source_kind == "column":
                    model_identifier = _text(
                        worksheet[f'{columns["model_identifier"]}{row}'].value
                    )
                elif source_kind == "series":
                    model_identifier = series
                else:
                    raise ValueError(
                        f"Unsupported model_identifier_source '{source_kind}'."
                    )

                if not model_identifier:
                    issues.append(
                        IdentifierMetadataIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="MODEL_IDENTIFIER_MISSING",
                            message=f"Missing model identifier for {series}/{size}.",
                            source_reference=f"{sheet_name}!{row}",
                        )
                    )
                    continue

                key = (family_code, series, size)
                if key in seen_model_keys:
                    issues.append(
                        IdentifierMetadataIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="DUPLICATE_MODEL_REFERENCE",
                            message=f"Duplicate model mapping for {series}/{size}.",
                            source_reference=f"{sheet_name}!{row}",
                        )
                    )
                    continue

                seen_model_keys.add(key)

                models.append(
                    PumpModelCandidate(
                        family_code=family_code,
                        model_identifier=model_identifier,
                        series_code=series,
                        size_code=size,
                        base_identifier=_base_identifier(
                            model_identifier,
                            profile["base_identifier"],
                        ),
                        source_worksheet=sheet_name,
                        source_row=row,
                        source_profile=str(profile_path),
                    )
                )
        finally:
            workbook.close()

    return IdentifierMetadataReport(
        model_count=len(models),
        format_count=len(formats),
        segment_count=len(segments),
        issue_count=len(issues),
        models=tuple(models),
        formats=tuple(formats),
        segments=tuple(segments),
        issues=tuple(issues),
    )


def save_json(report: IdentifierMetadataReport, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(report: IdentifierMetadataReport, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    _write_csv(directory / "identifier_model_candidates.csv", [asdict(x) for x in report.models])
    _write_csv(directory / "identifier_format_candidates.csv", [asdict(x) for x in report.formats])
    _write_csv(directory / "identifier_segment_candidates.csv", [asdict(x) for x in report.segments])
    _write_csv(directory / "identifier_metadata_issues.csv", [asdict(x) for x in report.issues])


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
