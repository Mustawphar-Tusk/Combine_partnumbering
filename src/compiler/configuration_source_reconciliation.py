from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


_HEADER_REF_RE = re.compile(r"\[\[#Headers\],\[(?P<field>.+?)\]\]", re.IGNORECASE)


@dataclass(frozen=True)
class ReconciliationIssue:
    severity: str
    issue_code: str
    message: str
    source_reference: str | None = None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    value = str(value).strip()
    return value or None


def canonical_field_code(label: str) -> str:
    value = label.strip().upper().replace("&", " AND ")
    value = re.sub(r"[^A-Z0-9]+", "_", value)
    return value.strip("_")


def _formula_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    text = getattr(value, "text", None)
    if isinstance(text, str):
        return text
    return str(value)


def _header_reference(value: Any) -> str | None:
    text = _formula_text(value)
    match = _HEADER_REF_RE.search(text)
    if not match:
        return None
    field = match.group("field").strip()
    return field or None


def _load_workbook(path: Path, *, data_only: bool = False):
    return load_workbook(
        path,
        read_only=False,
        data_only=data_only,
        keep_vba=path.suffix.lower() == ".xlsm",
        keep_links=False,
    )


def _resolve_path(project_root: Path, value: str) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = project_root / path
    return path.resolve()


def _transform_base_identifier(
    model_identifier: str,
    *,
    from_prefix: str,
    to_prefix: str,
) -> str | None:
    model = model_identifier.strip()
    if not model.upper().startswith(from_prefix.upper()):
        return None
    return f"{to_prefix}{model[len(from_prefix):]}"


def _model_key(series: str, size: str) -> str:
    return f"{series.strip().upper()}|{size.strip().upper()}"


def _section_for_sequence(sections: list[dict[str, Any]], sequence: int) -> dict[str, Any] | None:
    for section in sections:
        if int(section["sequence_from"]) <= sequence <= int(section["sequence_to"]):
            return section
    return None


def _read_rev2_fields(
    workbook_path: Path,
    compiler_profile_path: Path,
) -> list[dict[str, Any]]:
    profile = json.loads(compiler_profile_path.read_text(encoding="utf-8"))
    worksheet_name = profile["worksheet_name"]
    strategy = profile["strategy"]
    sections = strategy.get("sections", [])

    wb = _load_workbook(workbook_path, data_only=False)
    try:
        ws = wb[worksheet_name]
        records: list[dict[str, Any]] = []
        seen_sequences: set[int] = set()

        for pair in strategy["column_pairs"]:
            seq_col = column_index_from_string(pair["sequence_column"])
            label_col = column_index_from_string(pair["label_column"])
            seq_min = int(pair["sequence_from"])
            seq_max = int(pair["sequence_to"])

            for row in range(int(pair["row_from"]), int(pair["row_to"]) + 1):
                raw_seq = ws.cell(row, seq_col).value
                try:
                    sequence = int(raw_seq)
                except (TypeError, ValueError):
                    continue

                if sequence < seq_min or sequence > seq_max or sequence in seen_sequences:
                    continue

                label = _text(ws.cell(row, label_col).value)
                if not label:
                    continue

                seen_sequences.add(sequence)
                section = _section_for_sequence(sections, sequence)
                canonical = canonical_field_code(label)

                records.append(
                    {
                        "sequence": sequence,
                        "label": label,
                        "canonical_field_code": canonical,
                        "rev2_field_code": f"{canonical}_{sequence:03d}",
                        "section_code": section["section_code"] if section else None,
                        "section_name": section["section_name"] if section else None,
                        "source_workbook": workbook_path.name,
                        "source_worksheet": worksheet_name,
                        "source_cell": f"{pair['label_column']}{row}",
                    }
                )

        return sorted(records, key=lambda item: item["sequence"])
    finally:
        wb.close()



def _validate_rev2_sequences(
    compiler_profile_path: Path,
    rev2_fields: list[dict[str, Any]],
    expected_sequence_gaps: list[int],
    issues: list[ReconciliationIssue],
) -> dict[str, Any]:
    profile = json.loads(compiler_profile_path.read_text(encoding="utf-8"))
    strategy = profile["strategy"]

    minimum = int(
        strategy.get(
            "minimum_sequence",
            min(int(pair["sequence_from"]) for pair in strategy["column_pairs"]),
        )
    )
    maximum = int(
        strategy.get(
            "maximum_sequence",
            max(int(pair["sequence_to"]) for pair in strategy["column_pairs"]),
        )
    )

    found = {int(item["sequence"]) for item in rev2_fields}
    expected = set(range(minimum, maximum + 1))
    missing = sorted(expected - found)
    configured_gaps = sorted({int(value) for value in expected_sequence_gaps})
    unexpected_missing = sorted(set(missing) - set(configured_gaps))
    configured_gaps_now_present = sorted(set(configured_gaps) & found)

    for sequence in unexpected_missing:
        issues.append(
            ReconciliationIssue(
                "Error",
                "UNEXPECTED_REV2_SEQUENCE_GAP",
                f"Rev2 configuration sequence {sequence} is missing and is not declared as an intentional source gap.",
                f"sequence:{sequence}",
            )
        )

    for sequence in configured_gaps_now_present:
        issues.append(
            ReconciliationIssue(
                "Warning",
                "EXPECTED_REV2_GAP_NOW_PRESENT",
                f"Sequence {sequence} is configured as an intentional gap, but the current workbook now contains it. Review the reconciliation profile before publication.",
                f"sequence:{sequence}",
            )
        )

    return {
        "minimum_sequence": minimum,
        "maximum_sequence": maximum,
        "found_count": len(found),
        "missing_sequences": missing,
        "expected_sequence_gaps": configured_gaps,
        "unexpected_missing_sequences": unexpected_missing,
        "configured_gaps_now_present": configured_gaps_now_present,
    }


def _read_rev2_model_references(
    workbook_path: Path,
    identifier_profile_path: Path,
    base_rule: dict[str, Any],
    issues: list[ReconciliationIssue],
) -> list[dict[str, Any]]:
    profile = json.loads(identifier_profile_path.read_text(encoding="utf-8"))
    source = profile["model_reference"]
    worksheet_name = source["worksheet_name"]
    row_from = int(source.get("row_from", 2))
    columns = source["columns"]

    wb = _load_workbook(workbook_path, data_only=True)
    try:
        ws = wb[worksheet_name]
        model_col = column_index_from_string(columns["model_identifier"])
        series_col = column_index_from_string(columns["series"])
        size_col = column_index_from_string(columns["size"])
        records: list[dict[str, Any]] = []
        seen: dict[str, dict[str, Any]] = {}

        for row in range(row_from, ws.max_row + 1):
            model = _text(ws.cell(row, model_col).value)
            series = _text(ws.cell(row, series_col).value)
            size = _text(ws.cell(row, size_col).value)
            if not (model and series and size):
                continue

            base = _transform_base_identifier(
                model,
                from_prefix=base_rule["from_prefix"],
                to_prefix=base_rule["to_prefix"],
            )
            if base is None:
                issues.append(
                    ReconciliationIssue(
                        "Warning",
                        "REV2_MODEL_IDENTIFIER_NON_A_PREFIX",
                        f"Authoritative model identifier {model!r} does not match the configured A-number prefix.",
                        f"{worksheet_name}!{get_column_letter(model_col)}{row}",
                    )
                )

            key = _model_key(series, size)
            record = {
                "model_key": key,
                "model_identifier": model,
                "base_identifier": base,
                "series": series,
                "size": size,
                "source": "REV2_MODEL_REFERENCE",
                "source_workbook": workbook_path.name,
                "source_worksheet": worksheet_name,
                "source_row": row,
            }

            if key in seen and seen[key]["model_identifier"] != model:
                issues.append(
                    ReconciliationIssue(
                        "Error",
                        "DUPLICATE_REV2_MODEL_KEY",
                        f"Series/size {series!r}/{size!r} maps to both "
                        f"{seen[key]['model_identifier']!r} and {model!r} in the authoritative Rev2 model reference.",
                        f"{worksheet_name}!{row}",
                    )
                )
            else:
                seen[key] = record
            records.append(record)

        return records
    finally:
        wb.close()


def _read_config_option_domains(
    workbook_path: Path,
    sheet_name: str,
    issues: list[ReconciliationIssue],
) -> list[dict[str, Any]]:
    wb = _load_workbook(workbook_path, data_only=True)
    try:
        ws = wb[sheet_name]
        result: list[dict[str, Any]] = []

        for table in ws.tables.values():
            min_col, min_row, max_col, max_row = _range_boundaries(table.ref)
            headers = [
                _text(ws.cell(min_row, col).value) or f"COLUMN_{col}"
                for col in range(min_col, max_col + 1)
            ]
            primary_label = headers[0]
            rows: list[dict[str, Any]] = []

            for row in range(min_row + 1, max_row + 1):
                values = [ws.cell(row, col).value for col in range(min_col, max_col + 1)]
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append({headers[i]: values[i] for i in range(len(headers))})

            if not rows:
                issues.append(
                    ReconciliationIssue(
                        "Warning",
                        "EMPTY_LOGIC_OPTION_DOMAIN",
                        f"Config Options table {table.name} ({primary_label}) contains no data rows.",
                        f"{sheet_name}!{table.ref}",
                    )
                )

            result.append(
                {
                    "table_name": table.name,
                    "table_range": table.ref,
                    "field_label": primary_label,
                    "canonical_field_code": canonical_field_code(primary_label),
                    "headers": headers,
                    "row_count": len(rows),
                    "rows": rows,
                    "source_workbook": workbook_path.name,
                    "source_worksheet": sheet_name,
                }
            )

        return sorted(result, key=lambda item: item["canonical_field_code"])
    finally:
        wb.close()


def _range_boundaries(ref: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries
    return range_boundaries(ref)


def _read_codependencies(
    workbook_path: Path,
    sheet_name: str,
) -> list[dict[str, Any]]:
    wb = _load_workbook(workbook_path, data_only=True)
    try:
        ws = wb[sheet_name]
        result: list[dict[str, Any]] = []

        for table in ws.tables.values():
            min_col, min_row, max_col, max_row = _range_boundaries(table.ref)
            headers = [
                _text(ws.cell(min_row, col).value) or f"COLUMN_{col}"
                for col in range(min_col, max_col + 1)
            ]

            rows: list[dict[str, Any]] = []
            for row in range(min_row + 1, max_row + 1):
                values = [ws.cell(row, col).value for col in range(min_col, max_col + 1)]
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append({headers[i]: values[i] for i in range(len(headers))})

            result.append(
                {
                    "table_name": table.name,
                    "table_range": table.ref,
                    "field_labels": headers,
                    "canonical_field_codes": [canonical_field_code(x) for x in headers],
                    "row_count": len(rows),
                    "rows": rows,
                    "source_workbook": workbook_path.name,
                    "source_worksheet": sheet_name,
                }
            )

        return sorted(result, key=lambda item: item["table_name"])
    finally:
        wb.close()


def _find_model_header_row(ws) -> int:
    for row in range(1, min(ws.max_row, 25) + 1):
        a = (_text(ws.cell(row, 1).value) or "").upper()
        b = (_text(ws.cell(row, 2).value) or "").upper()
        c = (_text(ws.cell(row, 3).value) or "").upper()
        if a in {"A NUMBER", "A#", "A# NUMBER"} and b == "SERIES" and c == "SIZE":
            return row
    raise ValueError(f"Could not locate A Number / Series / Size header row on worksheet {ws.title!r}.")


def _matrix_groups(ws, header_row: int) -> list[dict[str, Any]]:
    starts: list[tuple[int, str]] = []

    for col in range(4, ws.max_column + 1):
        field = None
        for row in range(1, header_row + 1):
            field = _header_reference(ws.cell(row, col).value)
            if field:
                break
        if field:
            starts.append((col, field))

    groups: list[dict[str, Any]] = []

    for index, (start_col, field_label) in enumerate(starts):
        next_start = starts[index + 1][0] if index + 1 < len(starts) else ws.max_column + 1
        option_columns: list[dict[str, Any]] = []

        for col in range(start_col + 1, next_start):
            label = None
            for row in range(1, header_row + 1):
                value = ws.cell(row, col).value
                text = _text(value)
                if not text:
                    continue
                if _header_reference(value):
                    continue
                if text.startswith("="):
                    continue
                if text.upper() in {"STD", "X", "C.F.", "CF", "O"}:
                    continue
                label = text
                break
            if label:
                option_columns.append({"column": get_column_letter(col), "label": label, "column_index": col})

        groups.append(
            {
                "field_label": field_label,
                "canonical_field_code": canonical_field_code(field_label),
                "group_column": get_column_letter(start_col),
                "group_column_index": start_col,
                "span_end_column": get_column_letter(next_start - 1),
                "option_columns": option_columns,
            }
        )

    return groups


def _read_matrix(
    workbook_path: Path,
    sheet_name: str,
    base_rule: dict[str, Any],
    issues: list[ReconciliationIssue],
    *,
    collect_models: bool,
) -> dict[str, Any]:
    wb = _load_workbook(workbook_path, data_only=False)
    try:
        ws = wb[sheet_name]
        header_row = _find_model_header_row(ws)
        groups = _matrix_groups(ws, header_row)

        model_rows: list[dict[str, Any]] = []
        data_rows: list[int] = []

        for row in range(header_row + 1, ws.max_row + 1):
            model = _text(ws.cell(row, 1).value)
            series = _text(ws.cell(row, 2).value)
            size = _text(ws.cell(row, 3).value)
            if not (model and series and size):
                continue

            data_rows.append(row)
            if collect_models:
                base = _transform_base_identifier(
                    model,
                    from_prefix=base_rule["from_prefix"],
                    to_prefix=base_rule["to_prefix"],
                )
                if base is None:
                    issues.append(
                        ReconciliationIssue(
                            "Warning",
                            "LOGIC_MODEL_IDENTIFIER_NON_A_PREFIX",
                            f"Logic matrix model identifier {model!r} does not match the configured A-number prefix.",
                            f"{sheet_name}!A{row}",
                        )
                    )
                model_rows.append(
                    {
                        "model_key": _model_key(series, size),
                        "model_identifier": model,
                        "base_identifier": base,
                        "series": series,
                        "size": size,
                        "source": "LOGIC_PUMP_OPTIONS",
                        "source_workbook": workbook_path.name,
                        "source_worksheet": sheet_name,
                        "source_row": row,
                    }
                )

        for group in groups:
            for option in group["option_columns"]:
                counter: Counter[str] = Counter()
                numeric_count = 0
                other_count = 0
                for row in data_rows:
                    value = ws.cell(row, option["column_index"]).value
                    if value in (None, ""):
                        continue
                    if isinstance(value, (int, float)):
                        numeric_count += 1
                    else:
                        text = str(value).strip()
                        if text:
                            counter[text] += 1
                        else:
                            other_count += 1
                option["value_counts"] = dict(counter)
                option["numeric_count"] = numeric_count
                option["other_count"] = other_count
                option.pop("column_index", None)
            group.pop("group_column_index", None)

        return {
            "sheet_name": sheet_name,
            "dimensions": {"rows": ws.max_row, "columns": ws.max_column},
            "model_header_row": header_row,
            "model_row_count": len(model_rows) if collect_models else len(data_rows),
            "models": model_rows,
            "groups": groups,
        }
    finally:
        wb.close()


def _canonical_with_alias(label: str, aliases: dict[str, str]) -> str:
    direct = aliases.get(label)
    if direct:
        return canonical_field_code(direct)
    canonical = canonical_field_code(label)
    for source, target in aliases.items():
        if canonical_field_code(source) == canonical:
            return canonical_field_code(target)
    return canonical


def _reconcile_fields(
    rev2_fields: list[dict[str, Any]],
    domains: list[dict[str, Any]],
    dependencies: list[dict[str, Any]],
    pump_matrix: dict[str, Any],
    price_matrix: dict[str, Any],
    aliases: dict[str, str],
    issues: list[ReconciliationIssue],
) -> list[dict[str, Any]]:
    registry: dict[str, dict[str, Any]] = {}

    def ensure(label: str) -> dict[str, Any]:
        canonical = _canonical_with_alias(label, aliases)
        if canonical not in registry:
            registry[canonical] = {
                "canonical_field_code": canonical,
                "labels": set(),
                "sources": set(),
                "roles": set(),
                "rev2_sequences": set(),
                "logic_option_tables": set(),
                "dependency_tables": set(),
                "pump_matrix_groups": set(),
                "price_matrix_groups": set(),
            }
        registry[canonical]["labels"].add(label)
        return registry[canonical]

    for item in rev2_fields:
        rec = ensure(item["label"])
        rec["sources"].add("DEAN_REV2")
        rec["roles"].add("ORDERED_CONFIGURATION")
        rec["rev2_sequences"].add(item["sequence"])

    for item in domains:
        rec = ensure(item["field_label"])
        rec["sources"].add("PUMPCONFIGURATION_LOGIC")
        rec["roles"].add("OPTION_DOMAIN")
        rec["logic_option_tables"].add(item["table_name"])

    for dep in dependencies:
        for label in dep["field_labels"]:
            rec = ensure(label)
            rec["sources"].add("PUMPCONFIGURATION_LOGIC")
            rec["roles"].add("DEPENDENCY")
            rec["dependency_tables"].add(dep["table_name"])

    for group in pump_matrix["groups"]:
        rec = ensure(group["field_label"])
        rec["sources"].add("PUMPCONFIGURATION_LOGIC")
        rec["roles"].add("MODEL_APPLICABILITY")
        rec["pump_matrix_groups"].add(group["group_column"])

    for group in price_matrix["groups"]:
        rec = ensure(group["field_label"])
        rec["sources"].add("PUMPCONFIGURATION_LOGIC")
        rec["roles"].add("PRICING_LINEAGE")
        rec["price_matrix_groups"].add(group["group_column"])

    result: list[dict[str, Any]] = []
    for canonical, rec in sorted(registry.items()):
        has_rev2 = "DEAN_REV2" in rec["sources"]
        has_logic = "PUMPCONFIGURATION_LOGIC" in rec["sources"]
        roles = rec["roles"]

        if has_rev2 and has_logic:
            classification = "SHARED"
        elif has_rev2:
            classification = "REV2_ONLY"
        elif "OPTION_DOMAIN" in roles or "MODEL_APPLICABILITY" in roles:
            classification = "LOGIC_ONLY_SUPPLEMENTAL"
        elif "DEPENDENCY" in roles:
            classification = "LOGIC_DEPENDENCY_ONLY"
        else:
            classification = "LOGIC_PRICING_ONLY"

        if "DEPENDENCY" in roles and not (
            "OPTION_DOMAIN" in roles or "MODEL_APPLICABILITY" in roles or has_rev2
        ):
            issues.append(
                ReconciliationIssue(
                    "Warning",
                    "DEPENDENCY_FIELD_WITHOUT_DOMAIN",
                    f"{canonical} participates in codependencies but has no Rev2 field, Config Options domain, or Pump Options applicability group.",
                    canonical,
                )
            )

        result.append(
            {
                "canonical_field_code": canonical,
                "labels": sorted(rec["labels"]),
                "classification": classification,
                "sources": sorted(rec["sources"]),
                "roles": sorted(roles),
                "rev2_sequences": sorted(rec["rev2_sequences"]),
                "logic_option_tables": sorted(rec["logic_option_tables"]),
                "dependency_tables": sorted(rec["dependency_tables"]),
                "pump_matrix_groups": sorted(rec["pump_matrix_groups"]),
                "price_matrix_groups": sorted(rec["price_matrix_groups"]),
                "identifier_participation": "UNASSIGNED",
                "publication_status": "CANDIDATE",
            }
        )

    return result


def _normalize_model_identifier(
    model_identifier: str,
    aliases: dict[str, str],
) -> tuple[str, str | None]:
    model = model_identifier.strip()
    lookup = {str(key).strip().upper(): str(value).strip() for key, value in aliases.items()}
    normalized = lookup.get(model.upper())
    if normalized:
        return normalized, normalized
    return model, None


def _logic_key_alias(
    item: dict[str, Any],
    rules: list[dict[str, Any]],
) -> dict[str, Any] | None:
    for rule in rules:
        model_match = str(rule.get("logic_model_identifier", "")).strip().upper()
        series_match = str(rule.get("logic_series", "")).strip().upper()
        size_match = str(rule.get("logic_size", "")).strip().upper()

        if model_match and item["model_identifier"].strip().upper() != model_match:
            continue
        if series_match and item["series"].strip().upper() != series_match:
            continue
        if size_match and item["size"].strip().upper() != size_match:
            continue
        return rule
    return None


def _approved_rev2_override(
    *,
    key: str,
    rev2: dict[str, Any],
    logic: dict[str, Any],
    rules: list[dict[str, Any]],
) -> dict[str, Any] | None:
    for rule in rules:
        rule_key = _model_key(str(rule["series"]), str(rule["size"]))
        if rule_key != key:
            continue
        if str(rule["rev2_model_identifier"]).strip().upper() != rev2["model_identifier"].strip().upper():
            continue
        if str(rule["logic_model_identifier"]).strip().upper() != logic["original_model_identifier"].strip().upper():
            continue
        return rule
    return None


def _reconcile_models(
    rev2_models: list[dict[str, Any]],
    logic_models: list[dict[str, Any]],
    issues: list[ReconciliationIssue],
    reconciliation_rules: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    rules = reconciliation_rules or {}
    identifier_aliases = rules.get("logic_identifier_aliases", {})
    key_alias_rules = rules.get("logic_key_aliases", [])
    approved_overrides = rules.get("approved_rev2_overrides", [])
    allow_logic_only_a_prefix = bool(rules.get("allow_logic_only_a_prefix_models", False))
    a_prefix = str(rules.get("a_number_prefix", "A"))

    rev2_by_key: dict[str, dict[str, Any]] = {}
    logic_by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for item in rev2_models:
        rev2_by_key[item["model_key"]] = item

    for item in logic_models:
        prepared = dict(item)
        prepared["original_model_key"] = item["model_key"]
        prepared["original_model_identifier"] = item["model_identifier"]
        prepared["original_series"] = item["series"]
        prepared["original_size"] = item["size"]

        normalized_model, model_alias = _normalize_model_identifier(
            item["model_identifier"], identifier_aliases
        )
        prepared["normalized_model_identifier"] = normalized_model
        prepared["identifier_alias_applied"] = model_alias

        key_alias = _logic_key_alias(item, key_alias_rules)
        if key_alias:
            prepared["series"] = str(key_alias["rev2_series"]).strip()
            prepared["size"] = str(key_alias.get("rev2_size", item["size"])).strip()
            prepared["model_key"] = _model_key(prepared["series"], prepared["size"])
            prepared["key_alias_applied"] = {
                "reason": key_alias.get("reason"),
                "from_model_key": item["model_key"],
                "to_model_key": prepared["model_key"],
            }
        else:
            prepared["key_alias_applied"] = None

        logic_by_key[prepared["model_key"]].append(prepared)

    keys = sorted(set(rev2_by_key) | set(logic_by_key))
    result: list[dict[str, Any]] = []

    for key in keys:
        rev2 = rev2_by_key.get(key)
        logic_list = logic_by_key.get(key, [])

        normalized_logic_ids = {
            x["normalized_model_identifier"].strip().upper()
            for x in logic_list
        }
        if len(normalized_logic_ids) > 1:
            issues.append(
                ReconciliationIssue(
                    "Error",
                    "DUPLICATE_LOGIC_MODEL_KEY",
                    f"Pump Options has multiple effective model identifiers for model key {key}: "
                    f"{sorted(x['original_model_identifier'] for x in logic_list)}.",
                    key,
                )
            )

        logic = logic_list[0] if logic_list else None
        rule_reason = None

        if rev2 and logic:
            if rev2["model_identifier"].strip().upper() == logic["normalized_model_identifier"].strip().upper():
                if logic["key_alias_applied"] or logic["identifier_alias_applied"]:
                    classification = "SHARED_ALIAS_MATCH"
                    rule_reason = (
                        (logic["key_alias_applied"] or {}).get("reason")
                        or "Configured model-reference alias reconciles the Logic workbook to the authoritative Rev2 reference."
                    )
                else:
                    classification = "SHARED_MATCH"
                authoritative = rev2
            else:
                override = _approved_rev2_override(
                    key=key,
                    rev2=rev2,
                    logic=logic,
                    rules=approved_overrides,
                )
                if override:
                    classification = "SHARED_REV2_OVERRIDE"
                    authoritative = rev2
                    rule_reason = override.get("reason")
                else:
                    classification = "MODEL_IDENTIFIER_MISMATCH"
                    authoritative = rev2
                    issues.append(
                        ReconciliationIssue(
                            "Error",
                            "MODEL_REFERENCE_MISMATCH",
                            f"Rev2 maps {key} to {rev2['model_identifier']!r}, while PumpConfiguration_Logic maps it to {logic['original_model_identifier']!r}. Rev2 remains authoritative.",
                            key,
                        )
                    )
        elif rev2:
            classification = "REV2_ONLY"
            authoritative = rev2
        else:
            original_model = logic["original_model_identifier"] if logic else ""
            if allow_logic_only_a_prefix and original_model.strip().upper().startswith(a_prefix.upper()):
                classification = "LOGIC_ONLY_SUPPLEMENTAL_READY"
                authoritative = logic
                rule_reason = (
                    "PumpConfiguration_Logic supplies an A-number model missing from Rev2; "
                    "accepted as supplemental metadata under the configured gap-fill policy."
                )
            else:
                classification = "LOGIC_ONLY_REVIEW"
                authoritative = None
                issues.append(
                    ReconciliationIssue(
                        "Warning",
                        "LOGIC_MODEL_NOT_IN_REV2_REFERENCE",
                        f"PumpConfiguration_Logic contains model key {logic['original_model_key'] if logic else key}, but the authoritative Rev2 model-reference source does not.",
                        logic["original_model_key"] if logic else key,
                    )
                )

        ready_classes = {
            "SHARED_MATCH",
            "SHARED_ALIAS_MATCH",
            "SHARED_REV2_OVERRIDE",
            "LOGIC_ONLY_SUPPLEMENTAL_READY",
        }
        publication_status = "READY" if classification in ready_classes else "REVIEW"

        result.append(
            {
                "model_key": key,
                "classification": classification,
                "authoritative_model_identifier": authoritative["model_identifier"] if authoritative else None,
                "authoritative_base_identifier": authoritative["base_identifier"] if authoritative else None,
                "rev2_model_identifier": rev2["model_identifier"] if rev2 else None,
                "logic_model_identifier": logic["original_model_identifier"] if logic else None,
                "logic_original_model_key": logic["original_model_key"] if logic else None,
                "logic_effective_model_identifier": logic["normalized_model_identifier"] if logic else None,
                "series": (rev2 or logic or {})["series"],
                "size": (rev2 or logic or {})["size"],
                "logic_original_series": logic["original_series"] if logic else None,
                "logic_original_size": logic["original_size"] if logic else None,
                "reconciliation_reason": rule_reason,
                "publication_status": publication_status,
            }
        )

    return result


def compile_configuration_source_reconciliation(
    *,
    project_root: Path,
    profile_path: Path,
) -> dict[str, Any]:
    project_root = project_root.resolve()
    profile_path = profile_path.resolve()
    profile = json.loads(profile_path.read_text(encoding="utf-8"))

    family_code = profile["family_code"].upper()
    issues: list[ReconciliationIssue] = []

    rev2_workbook = _resolve_path(project_root, profile["rev2"]["workbook"])
    configuration_profile = _resolve_path(
        project_root, profile["rev2"]["configuration_compiler_profile"]
    )
    identifier_profile = _resolve_path(project_root, profile["rev2"]["identifier_profile"])
    logic_workbook = _resolve_path(project_root, profile["logic"]["workbook"])

    for required in (rev2_workbook, configuration_profile, identifier_profile, logic_workbook):
        if not required.exists():
            raise FileNotFoundError(required)

    base_rule = profile["base_identifier"]
    aliases = profile.get("field_aliases", {})

    rev2_fields = _read_rev2_fields(rev2_workbook, configuration_profile)
    sequence_validation = _validate_rev2_sequences(
        configuration_profile,
        rev2_fields,
        profile.get("expected_sequence_gaps", []),
        issues,
    )
    rev2_models = _read_rev2_model_references(
        rev2_workbook, identifier_profile, base_rule, issues
    )
    domains = _read_config_option_domains(
        logic_workbook, profile["logic"]["config_options_sheet"], issues
    )
    dependencies = _read_codependencies(
        logic_workbook, profile["logic"]["codependencies_sheet"]
    )
    pump_matrix = _read_matrix(
        logic_workbook,
        profile["logic"]["pump_options_sheet"],
        base_rule,
        issues,
        collect_models=True,
    )
    price_matrix = _read_matrix(
        logic_workbook,
        profile["logic"]["price_options_sheet"],
        base_rule,
        issues,
        collect_models=False,
    )

    field_reconciliation = _reconcile_fields(
        rev2_fields,
        domains,
        dependencies,
        pump_matrix,
        price_matrix,
        aliases,
        issues,
    )
    model_reconciliation = _reconcile_models(
        rev2_models,
        pump_matrix["models"],
        issues,
        profile.get("model_reference_reconciliation", {}),
    )

    classification_counts = Counter(
        item["classification"] for item in field_reconciliation
    )
    model_classification_counts = Counter(
        item["classification"] for item in model_reconciliation
    )

    report = {
        "family_code": family_code,
        "reconciliation_version": profile["reconciliation_version"],
        "policy": profile.get("publication_policy", {}),
        "summary": {
            "rev2_field_count": len(rev2_fields),
            "rev2_sequence_validation": sequence_validation,
            "logic_option_domain_count": len(domains),
            "logic_option_row_count": sum(item["row_count"] for item in domains),
            "dependency_table_count": len(dependencies),
            "dependency_row_count": sum(item["row_count"] for item in dependencies),
            "pump_matrix_model_count": pump_matrix["model_row_count"],
            "pump_matrix_group_count": len(pump_matrix["groups"]),
            "price_matrix_group_count": len(price_matrix["groups"]),
            "field_reconciliation_count": len(field_reconciliation),
            "field_classification_counts": dict(sorted(classification_counts.items())),
            "rev2_model_reference_count": len(rev2_models),
            "model_reconciliation_count": len(model_reconciliation),
            "model_classification_counts": dict(sorted(model_classification_counts.items())),
            "issue_count": len(issues),
            "error_count": sum(1 for issue in issues if issue.severity == "Error"),
            "warning_count": sum(1 for issue in issues if issue.severity == "Warning"),
            "info_count": sum(1 for issue in issues if issue.severity == "Info"),
        },
        "sequence_validation": sequence_validation,
        "rev2_fields": rev2_fields,
        "logic_option_domains": domains,
        "dependencies": dependencies,
        "pump_options_matrix": pump_matrix,
        "price_options_matrix": price_matrix,
        "field_reconciliation": field_reconciliation,
        "model_reconciliation": model_reconciliation,
        "issues": [asdict(issue) for issue in issues],
    }
    return report


def _write_csv(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    rows = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    flattened: list[dict[str, Any]] = []
    for row in rows:
        item: dict[str, Any] = {}
        for key, value in row.items():
            if isinstance(value, (list, dict, tuple, set)):
                item[key] = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            else:
                item[key] = value
        flattened.append(item)

    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(flattened[0].keys()))
        writer.writeheader()
        writer.writerows(flattened)


def save_reconciliation_outputs(report: dict[str, Any], export_dir: Path) -> dict[str, Path]:
    export_dir.mkdir(parents=True, exist_ok=True)

    json_path = export_dir / "m023_dean_source_reconciliation.json"
    field_csv = export_dir / "m023_dean_field_reconciliation.csv"
    model_csv = export_dir / "m023_dean_model_reconciliation.csv"
    dependency_csv = export_dir / "m023_dean_dependency_candidates.csv"
    issue_csv = export_dir / "m023_dean_reconciliation_issues.csv"

    json_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    _write_csv(field_csv, report["field_reconciliation"])
    _write_csv(model_csv, report["model_reconciliation"])
    _write_csv(dependency_csv, report["dependencies"])
    _write_csv(issue_csv, report["issues"])

    return {
        "json": json_path,
        "fields": field_csv,
        "models": model_csv,
        "dependencies": dependency_csv,
        "issues": issue_csv,
    }
