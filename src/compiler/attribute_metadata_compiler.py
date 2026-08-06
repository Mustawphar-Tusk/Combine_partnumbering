from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


INVALID_VALUES = {"ERR", "TBD__", "____"}


@dataclass(frozen=True)
class AttributeValueCandidate:
    family_code: str
    workbook_role: str
    workbook_name: str
    worksheet_name: str
    field_code: str
    field_name: str
    display_value: str
    identifier_code: str
    display_order: int
    source_display_cell: str
    source_code_cell: str
    source_profile: str


@dataclass(frozen=True)
class AttributeValueIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    field_code: str | None = None
    source_reference: str | None = None


@dataclass(frozen=True)
class AttributeValueReport:
    field_count: int
    value_count: int
    issue_count: int
    values: tuple[AttributeValueCandidate, ...]
    issues: tuple[AttributeValueIssue, ...]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_workbook(
    project_root: Path,
    discovery: dict[str, Any],
    family_code: str,
    workbook_role: str,
) -> Path:
    matches = [
        row for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == workbook_role
        and row.get("discovery_status") == "discovered"
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one workbook for {family_code}/{workbook_role}; "
            f"found {len(matches)}."
        )
    return project_root / matches[0]["relative_path"]


def compile_attribute_metadata(
    project_root: Path,
    discovery_path: Path,
    profile_path: Path,
) -> AttributeValueReport:
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))
    profile = json.loads(profile_path.read_text(encoding="utf-8"))

    family_code = profile["family_code"].upper()
    role = profile["workbook_role"]
    sheet_name = profile["worksheet_name"]
    row_from = int(profile["row_from"])
    row_to = int(profile.get("row_to", 100000))
    blank_limit = int(profile.get("max_consecutive_blank_rows", 100))

    workbook_path = _resolve_workbook(
        project_root,
        discovery,
        family_code,
        role,
    )

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=False,
        keep_links=False,
    )

    values: list[AttributeValueCandidate] = []
    issues: list[AttributeValueIssue] = []
    seen: set[tuple[str, str]] = set()

    try:
        if sheet_name not in workbook.sheetnames:
            return AttributeValueReport(
                field_count=0,
                value_count=0,
                issue_count=1,
                values=(),
                issues=(
                    AttributeValueIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="WORKSHEET_NOT_FOUND",
                        message=f"Worksheet '{sheet_name}' was not found.",
                    ),
                ),
            )

        worksheet = workbook[sheet_name]

        for attribute in profile["attributes"]:
            field_code = attribute["field_code"]
            field_name = attribute["field_name"]
            display_column = attribute["display_column"]
            code_column = attribute["code_column"]
            blank_count = 0
            display_order = 0

            for row in range(row_from, min(row_to, worksheet.max_row or row_to) + 1):
                display = _text(worksheet[f"{display_column}{row}"].value)
                code = _text(worksheet[f"{code_column}{row}"].value)

                if not display and not code:
                    blank_count += 1
                    if blank_count >= blank_limit:
                        break
                    continue

                blank_count = 0

                if not display or not code:
                    issues.append(
                        AttributeValueIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="INCOMPLETE_ATTRIBUTE_ROW",
                            message=(
                                f"{field_code} row {row} is missing "
                                "display value or identifier code."
                            ),
                            field_code=field_code,
                            source_reference=f"{sheet_name}!{row}",
                        )
                    )
                    continue

                if display.upper() in INVALID_VALUES or code.upper() in INVALID_VALUES:
                    continue

                key = (field_code, display.casefold())

                if key in seen:
                    issues.append(
                        AttributeValueIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="DUPLICATE_DISPLAY_VALUE",
                            message=(
                                f"Duplicate display value '{display}' "
                                f"for {field_code}."
                            ),
                            field_code=field_code,
                            source_reference=f"{sheet_name}!{display_column}{row}",
                        )
                    )
                    continue

                seen.add(key)
                display_order += 10

                values.append(
                    AttributeValueCandidate(
                        family_code=family_code,
                        workbook_role=role,
                        workbook_name=workbook_path.name,
                        worksheet_name=sheet_name,
                        field_code=field_code,
                        field_name=field_name,
                        display_value=display,
                        identifier_code=code,
                        display_order=display_order,
                        source_display_cell=f"{display_column}{row}",
                        source_code_cell=f"{code_column}{row}",
                        source_profile=str(profile_path),
                    )
                )

    finally:
        workbook.close()

    return AttributeValueReport(
        field_count=len({row.field_code for row in values}),
        value_count=len(values),
        issue_count=len(issues),
        values=tuple(values),
        issues=tuple(issues),
    )


def save_json(report: AttributeValueReport, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(report: AttributeValueReport, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    _write_csv(
        directory / "fybroc_attribute_candidates.csv",
        [asdict(row) for row in report.values],
    )
    _write_csv(
        directory / "fybroc_attribute_issues.csv",
        [asdict(row) for row in report.issues],
    )


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
