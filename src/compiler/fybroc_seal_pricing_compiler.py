from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from src.compiler.pricing_metadata_compiler import (
    PriceCandidate,
    PriceCandidateCondition,
    PricingCompilationReport,
    PricingIssue,
)


def _text(value: Any) -> str | None:
    if value is None:
        return None

    value = str(value).strip()
    return value or None


def _canonical_size(value: str) -> str:
    return value.split(" (", 1)[0].strip().lower()


def _lookup_key(value: str) -> str:
    return re.sub(
        r"\s+",
        "",
        value,
    ).casefold()


def _table_headers(
    worksheet,
    table_name: str,
) -> tuple[
    dict[str, int],
    int,
    int,
]:
    if table_name not in worksheet.tables:
        raise KeyError(
            f"Table {table_name!r} was not found "
            f"on worksheet {worksheet.title!r}."
        )

    table = worksheet.tables[table_name]

    (
        min_col,
        min_row,
        max_col,
        max_row,
    ) = range_boundaries(
        table.ref
    )

    headers: dict[str, int] = {}

    for col in range(
        min_col,
        max_col + 1,
    ):
        header = _text(
            worksheet.cell(
                min_row,
                col,
            ).value
        )

        if header:
            headers[header] = col

    return (
        headers,
        min_row,
        max_row,
    )


def _load_size_groups(
    workbook,
    profile: dict[str, Any],
) -> dict[
    str,
    tuple[int, str],
]:
    cfg = profile["table4"]

    worksheet = workbook[
        cfg["worksheet_name"]
    ]

    (
        headers,
        header_row,
        max_row,
    ) = _table_headers(
        worksheet,
        cfg["table_name"],
    )

    size_col = headers.get(
        cfg["size_header"]
    )
    group_col = headers.get(
        cfg["group_header"]
    )

    if size_col is None:
        raise KeyError(
            f"{cfg['table_name']} has no "
            f"{cfg['size_header']!r} column."
        )

    if group_col is None:
        raise KeyError(
            f"{cfg['table_name']} has no "
            f"{cfg['group_header']!r} column."
        )

    result: dict[
        str,
        tuple[int, str],
    ] = {}

    for row in range(
        header_row + 1,
        max_row + 1,
    ):
        source_size = _text(
            worksheet.cell(
                row,
                size_col,
            ).value
        )

        raw_group = worksheet.cell(
            row,
            group_col,
        ).value

        if not source_size:
            continue

        if not isinstance(
            raw_group,
            (int, float),
        ):
            continue

        group_no = int(raw_group)

        if group_no not in {
            1,
            2,
            3,
        }:
            continue

        canonical = _canonical_size(
            source_size
        )

        existing = result.get(
            canonical
        )

        if (
            existing is not None
            and existing[0] != group_no
        ):
            raise ValueError(
                "Conflicting Table4 group assignments "
                f"for size {canonical!r}: "
                f"{existing[0]} vs {group_no}."
            )

        result[canonical] = (
            group_no,
            source_size,
        )

    return result


def _load_supported_series_sizes(
    compilation_path: Path,
    supported_source_series: set[str],
) -> tuple[
    tuple[str, str, str],
    ...
]:
    data = json.loads(
        compilation_path.read_text(
            encoding="utf-8-sig",
        )
    )

    pairs: set[
        tuple[str, str, str]
    ] = set()

    for candidate in data.get(
        "candidates",
        [],
    ):
        source_series = _text(
            candidate.get(
                "source_series_code"
            )
        )

        runtime_series = _text(
            candidate.get(
                "series_code"
            )
        )

        size_value = _text(
            candidate.get(
                "size_value"
            )
        )

        if (
            source_series
            not in supported_source_series
        ):
            continue

        if (
            runtime_series is None
            or size_value is None
        ):
            continue

        pairs.add(
            (
                source_series,
                runtime_series,
                size_value,
            )
        )

    return tuple(
        sorted(
            pairs,
            key=lambda row: (
                row[0],
                row[2],
                row[1],
            ),
        )
    )


def _load_runtime_combinations(
    nomenclature_workbook,
    profile: dict[str, Any],
) -> tuple[
    tuple[
        str,
        str,
        str,
        str,
    ],
    ...
]:
    cfg = profile[
        "seal_assembly"
    ]

    worksheet = nomenclature_workbook[
        cfg["worksheet_name"]
    ]

    columns = cfg["columns"]

    combinations: set[
        tuple[
            str,
            str,
            str,
            str,
        ]
    ] = set()

    for row in range(
        int(cfg["first_data_row"]),
        int(cfg["last_data_row"])
        + 1,
    ):
        option = _text(
            worksheet[
                f"{columns['SEAL_OPTION']}{row}"
            ].value
        )
        seal_type = _text(
            worksheet[
                f"{columns['SEAL_TYPE']}{row}"
            ].value
        )
        materials = _text(
            worksheet[
                f"{columns['SEAL_MATERIALS']}{row}"
            ].value
        )
        elastomers = _text(
            worksheet[
                f"{columns['SEAL_ELASTOMERS']}{row}"
            ].value
        )

        if option is None:
            continue

        if (
            seal_type is None
            or materials is None
            or elastomers is None
        ):
            continue

        combinations.add(
            (
                option,
                seal_type,
                materials,
                elastomers,
            )
        )

    return tuple(
        sorted(combinations)
    )


def _build_seal_mapping(
    profile: dict[str, Any],
) -> dict[
    tuple[str, str],
    str,
]:
    result: dict[
        tuple[str, str],
        str,
    ] = {}

    for row in profile[
        "seal_type_material_mappings"
    ]:
        key = (
            str(row["seal_type"]),
            str(row["seal_materials"]),
        )

        if key in result:
            raise ValueError(
                "Duplicate seal type/material mapping "
                f"in profile: {key!r}"
            )

        result[key] = str(
            row["legacy_seal_key"]
        )

    return result


def _load_price_lookup(
    price_workbook,
    *,
    worksheet_name: str,
    table_name: str,
) -> dict[
    tuple[str, int],
    tuple[Any, str],
]:
    worksheet = price_workbook[
        worksheet_name
    ]

    (
        headers,
        header_row,
        max_row,
    ) = _table_headers(
        worksheet,
        table_name,
    )

    seal_key_col = headers.get(
        "Seal Type"
    )

    if seal_key_col is None:
        raise KeyError(
            f"{table_name} has no "
            "'Seal Type' column."
        )

    group_columns: dict[int, int] = {}

    for group_no in (
        1,
        2,
        3,
    ):
        header = f"Group {group_no}"
        col = headers.get(header)

        if col is None:
            raise KeyError(
                f"{table_name} has no "
                f"{header!r} column."
            )

        group_columns[
            group_no
        ] = col

    result: dict[
        tuple[str, int],
        tuple[Any, str],
    ] = {}

    for row in range(
        header_row + 1,
        max_row + 1,
    ):
        source_key = _text(
            worksheet.cell(
                row,
                seal_key_col,
            ).value
        )

        if not source_key:
            continue

        normalized_key = _lookup_key(
            source_key
        )

        for (
            group_no,
            col,
        ) in group_columns.items():
            cell = worksheet.cell(
                row,
                col,
            )

            result[
                (
                    normalized_key,
                    group_no,
                )
            ] = (
                cell.value,
                cell.coordinate,
            )

    return result


def _parse_price(
    raw_value: Any,
    status_map: dict[str, str],
) -> tuple[
    float | None,
    str | None,
]:
    if isinstance(
        raw_value,
        (int, float),
    ):
        return (
            float(raw_value),
            "found",
        )

    source_text = _text(
        raw_value
    )

    if source_text is None:
        return (
            None,
            None,
        )

    mapped_status = status_map.get(
        source_text.upper()
    )

    if mapped_status:
        return (
            None,
            mapped_status,
        )

    return (
        None,
        None,
    )


def compile_fybroc_seal_pricing(
    price_workbook_path: Path,
    nomenclature_workbook_path: Path,
    base_pump_compilation_path: Path,
    profile_path: Path,
) -> PricingCompilationReport:
    profile = json.loads(
        profile_path.read_text(
            encoding="utf-8-sig",
        )
    )

    family_code = str(
        profile["family_code"]
    ).strip().upper()

    component_code = str(
        profile["component_code"]
    ).strip().upper()

    currency_code = str(
        profile.get(
            "currency_code",
            "USD",
        )
    ).strip().upper()

    mechanical_option = str(
        profile[
            "seal_assembly"
        ][
            "mechanical_seal_option"
        ]
    )

    supported_source_series = {
        str(value).strip()
        for value in profile[
            "supported_source_series"
        ]
    }

    status_map = {
        str(key).strip().upper():
            str(value).strip()
        for (
            key,
            value,
        ) in profile.get(
            "price_value_status_map",
            {},
        ).items()
    }

    candidates: list[
        PriceCandidate
    ] = []

    issues: list[
        PricingIssue
    ] = []

    price_workbook = load_workbook(
        price_workbook_path,
        data_only=True,
        read_only=False,
        keep_vba=True,
    )

    nomenclature_workbook = (
        load_workbook(
            nomenclature_workbook_path,
            data_only=False,
            read_only=False,
            keep_vba=True,
        )
    )

    try:
        size_groups = (
            _load_size_groups(
                price_workbook,
                profile,
            )
        )

        series_sizes = (
            _load_supported_series_sizes(
                base_pump_compilation_path,
                supported_source_series,
            )
        )

        combinations = (
            _load_runtime_combinations(
                nomenclature_workbook,
                profile,
            )
        )

        seal_mapping = (
            _build_seal_mapping(
                profile
            )
        )

        elastomer_mapping = profile[
            "runtime_elastomer_to_legacy"
        ]

        non_mechanical_mapping = (
            profile.get(
                "non_mechanical_option_mappings",
                {},
            )
        )

        price_lookups: dict[
            str,
            dict[
                tuple[str, int],
                tuple[Any, str],
            ],
        ] = {}

        for (
            table_code,
            table_cfg,
        ) in profile[
            "price_tables"
        ].items():
            price_lookups[
                table_code
            ] = _load_price_lookup(
                price_workbook,
                worksheet_name=(
                    table_cfg[
                        "worksheet_name"
                    ]
                ),
                table_name=(
                    table_cfg[
                        "table_name"
                    ]
                ),
            )

        seen_keys: set[
            tuple[
                str,
                str,
                tuple[
                    tuple[
                        str,
                        str,
                        str | None,
                    ],
                    ...
                ],
            ]
        ] = set()

        for (
            source_series_code,
            series_code,
            size_value,
        ) in series_sizes:
            group_record = (
                size_groups.get(
                    size_value.casefold()
                )
            )

            if group_record is None:
                issues.append(
                    PricingIssue(
                        family_code=(
                            family_code
                        ),
                        component_code=(
                            component_code
                        ),
                        series_code=(
                            source_series_code
                        ),
                        severity="Error",
                        issue_code=(
                            "SEAL_SIZE_GROUP_NOT_FOUND"
                        ),
                        message=(
                            "Table4 has no numeric "
                            "Group 1/2/3 mapping "
                            f"for size {size_value!r}."
                        ),
                        source_reference=(
                            "Tables!Table4"
                        ),
                    )
                )
                continue

            (
                group_no,
                source_size_value,
            ) = group_record

            for (
                seal_option,
                seal_type,
                seal_materials,
                seal_elastomers,
            ) in combinations:
                source_configuration_value: str | None = None

                if seal_option == mechanical_option:
                    legacy_seal_key = (
                        seal_mapping.get(
                            (
                                seal_type,
                                seal_materials,
                            )
                        )
                    )

                    if legacy_seal_key is None:
                        issues.append(
                            PricingIssue(
                                family_code=family_code,
                                component_code=component_code,
                                series_code=source_series_code,
                                severity="Error",
                                issue_code=(
                                    "SEAL_VOCABULARY_MAPPING_MISSING"
                                ),
                                message=(
                                    "No pricing vocabulary "
                                    "mapping exists for "
                                    f"{seal_type!r} / "
                                    f"{seal_materials!r}."
                                ),
                                source_reference=(
                                    nomenclature_workbook_path.name
                                    + "!Seal Assembly"
                                ),
                            )
                        )
                        continue

                    elastomer_cfg = (
                        elastomer_mapping.get(
                            seal_elastomers
                        )
                    )

                    if elastomer_cfg is None:
                        issues.append(
                            PricingIssue(
                                family_code=family_code,
                                component_code=component_code,
                                series_code=source_series_code,
                                severity="Error",
                                issue_code=(
                                    "SEAL_ELASTOMER_MAPPING_MISSING"
                                ),
                                message=(
                                    "No pricing elastomer "
                                    "mapping exists for "
                                    f"{seal_elastomers!r}."
                                ),
                                source_reference=(
                                    nomenclature_workbook_path.name
                                    + "!Seal Assembly"
                                ),
                            )
                        )
                        continue

                    table_code = str(
                        elastomer_cfg[
                            "price_table"
                        ]
                    )

                    suffix = str(
                        elastomer_cfg.get(
                            "seal_key_suffix",
                            "",
                        )
                    )

                    source_price_key = (
                        legacy_seal_key
                        + suffix
                    )

                    source_elastomer = str(
                        elastomer_cfg[
                            "legacy_elastomer"
                        ]
                    )

                else:
                    option_cfg = (
                        non_mechanical_mapping.get(
                            seal_option
                        )
                    )

                    if option_cfg is None:
                        issues.append(
                            PricingIssue(
                                family_code=family_code,
                                component_code=component_code,
                                series_code=source_series_code,
                                severity="Error",
                                issue_code=(
                                    "SEAL_OPTION_MAPPING_MISSING"
                                ),
                                message=(
                                    "No non-mechanical pricing "
                                    "mapping exists for "
                                    f"{seal_option!r}."
                                ),
                                source_reference=(
                                    nomenclature_workbook_path.name
                                    + "!Seal Assembly"
                                ),
                            )
                        )
                        continue

                    if (
                        seal_type != "-"
                        or seal_materials != "-"
                        or seal_elastomers != "-"
                    ):
                        issues.append(
                            PricingIssue(
                                family_code=family_code,
                                component_code=component_code,
                                series_code=source_series_code,
                                severity="Error",
                                issue_code=(
                                    "NON_MECHANICAL_SEAL_DETAIL_NOT_EMPTY"
                                ),
                                message=(
                                    f"{seal_option!r} is non-mechanical "
                                    "but its detail fields are not all '-'."
                                ),
                                source_reference=(
                                    nomenclature_workbook_path.name
                                    + "!Seal Assembly"
                                ),
                            )
                        )
                        continue

                    legacy_seal_key = str(
                        option_cfg[
                            "legacy_seal_key"
                        ]
                    )

                    source_price_key = (
                        legacy_seal_key
                    )

                    table_code = str(
                        option_cfg.get(
                            "price_table",
                            "standard",
                        )
                    )

                    source_configuration_value = str(
                        option_cfg[
                            "source_configuration_value"
                        ]
                    )

                    source_elastomer = "-"

                lookup = (
                    price_lookups[
                        table_code
                    ]
                )

                price_record = (
                    lookup.get(
                        (
                            _lookup_key(
                                source_price_key
                            ),
                            group_no,
                        )
                    )
                )

                if price_record is None:
                    issues.append(
                        PricingIssue(
                            family_code=family_code,
                            component_code=component_code,
                            series_code=source_series_code,
                            severity="Error",
                            issue_code=(
                                "SEAL_PRICE_KEY_NOT_FOUND"
                            ),
                            message=(
                                f"{source_price_key!r} "
                                f"Group {group_no} "
                                "was not found in "
                                f"{table_code!r}."
                            ),
                            source_reference=(
                                "Pricebook"
                            ),
                        )
                    )
                    continue

                (
                    raw_price,
                    source_cell,
                ) = price_record

                (
                    amount,
                    pricing_status,
                ) = _parse_price(
                    raw_price,
                    status_map,
                )

                if pricing_status is None:
                    issues.append(
                        PricingIssue(
                            family_code=family_code,
                            component_code=component_code,
                            series_code=source_series_code,
                            severity="Error",
                            issue_code=(
                                "SEAL_PRICE_VALUE_UNSUPPORTED"
                            ),
                            message=(
                                "Unsupported seal price "
                                f"value {raw_price!r} "
                                f"at Pricebook!"
                                f"{source_cell}."
                            ),
                            source_reference=(
                                "Pricebook!"
                                + source_cell
                            ),
                        )
                    )
                    continue

                conditions = (
                    PriceCandidateCondition(
                        sequence_no=1,
                        field_code="SIZE",
                        comparison_operator="EQ",
                        comparison_value=size_value,
                        source_value=source_size_value,
                    ),
                    PriceCandidateCondition(
                        sequence_no=2,
                        field_code="SEAL_OPTION",
                        comparison_operator="EQ",
                        comparison_value=seal_option,
                        source_value=(
                            source_configuration_value
                            or seal_option
                        ),
                    ),
                    PriceCandidateCondition(
                        sequence_no=3,
                        field_code="SEAL_TYPE",
                        comparison_operator="EQ",
                        comparison_value=seal_type,
                        source_value=legacy_seal_key,
                    ),
                    PriceCandidateCondition(
                        sequence_no=4,
                        field_code="SEAL_MATERIALS",
                        comparison_operator="EQ",
                        comparison_value=seal_materials,
                        source_value=legacy_seal_key,
                    ),
                    PriceCandidateCondition(
                        sequence_no=5,
                        field_code="SEAL_ELASTOMERS",
                        comparison_operator="EQ",
                        comparison_value=seal_elastomers,
                        source_value=source_elastomer,
                    ),
                )

                identity = (
                    component_code,
                    series_code,
                    tuple(
                        (
                            condition.field_code,
                            condition.comparison_operator,
                            condition.comparison_value,
                        )
                        for condition
                        in conditions
                    ),
                )

                if identity in seen_keys:
                    issues.append(
                        PricingIssue(
                            family_code=family_code,
                            component_code=component_code,
                            series_code=source_series_code,
                            severity="Error",
                            issue_code=(
                                "DUPLICATE_SEAL_PRICING_KEY"
                            ),
                            message=(
                                "Duplicate canonical "
                                "SEAL pricing condition "
                                "set was generated."
                            ),
                            source_reference=(
                                "Pricebook!"
                                + source_cell
                            ),
                        )
                    )
                    continue

                seen_keys.add(
                    identity
                )

                table_name = (
                    profile[
                        "price_tables"
                    ][
                        table_code
                    ][
                        "table_name"
                    ]
                )

                candidates.append(
                    PriceCandidate(
                        family_code=family_code,
                        component_code=component_code,
                        series_code=series_code,
                        source_series_code=source_series_code,
                        size_value=size_value,
                        source_size_value=source_size_value,
                        option_field_code=None,
                        option_value=None,
                        source_option_value=(
                            source_price_key
                            + " | "
                            + source_elastomer
                        ),
                        amount=amount,
                        pricing_status=pricing_status,
                        source_price_value=raw_price,
                        currency_code=currency_code,
                        workbook_name=price_workbook_path.name,
                        worksheet_name=(
                            profile[
                                "price_tables"
                            ][
                                table_code
                            ][
                                "worksheet_name"
                            ]
                        ),
                        table_name=table_name,
                        source_cell=source_cell,
                        conditions=conditions,
                    )
                )

        return PricingCompilationReport(
            candidate_count=len(
                candidates
            ),
            issue_count=len(
                issues
            ),
            candidates=tuple(
                candidates
            ),
            issues=tuple(
                issues
            ),
        )

    finally:
        price_workbook.close()
        nomenclature_workbook.close()
