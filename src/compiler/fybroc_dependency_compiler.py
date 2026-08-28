from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class DependencyCandidate:
    family_code: str
    dependency_code: str
    target_field_code: str
    target_display_value: str
    target_identifier_code: str | None
    series_code: str | None
    context_json: str
    source_workbook: str
    source_worksheet: str
    source_reference: str


@dataclass(frozen=True)
class DependencyIssue:
    severity: str
    issue_code: str
    message: str
    source_reference: str | None = None


@dataclass(frozen=True)
class DependencyCompilationReport:
    candidate_count: int
    trim_relation_count: int
    motor_modification_relation_count: int
    issue_count: int
    candidates: tuple[DependencyCandidate, ...]
    issues: tuple[DependencyIssue, ...]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _format_trim(value: Any) -> str | None:
    if value is None:
        return None

    try:
        return f"{float(value):.3f}"
    except (TypeError, ValueError):
        return None


def _normal_description(value: str) -> str:
    text = value.casefold()
    text = text.replace("&", "and")
    text = text.replace("auxilliary", "auxiliary")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _catalog_prefixes(value: Any) -> set[str]:
    if value is None:
        return set()

    return {
        item.upper()
        for item in re.findall(
            r"\bM[0-9A-Z]+(?:-[0-9]+)?\b",
            str(value).upper(),
        )
    }


def _frame_class(value: str | None) -> str | None:
    if not value:
        return None

    text = value.upper().replace(" ", "")

    for candidate in (
        "56-180T", "210T", "250T", "280T",
        "320T", "360T", "400T", "440T",
        "5000", "5800",
    ):
        if candidate.replace(" ", "") in text:
            return candidate

    return None


def _find_trim_matrix(
    worksheet,
    *,
    size_pattern: re.Pattern[str],
    minimum_headers: int,
) -> tuple[int, list[tuple[int, str]]] | None:
    best: tuple[int, list[tuple[int, str]]] | None = None

    for row in worksheet.iter_rows():
        matches: list[tuple[int, str]] = []

        for cell in row:
            value = _text(cell.value)

            if value and size_pattern.match(value):
                matches.append((cell.column, value))

        if len(matches) < minimum_headers:
            continue

        if best is None or len(matches) > len(best[1]):
            best = (row[0].row, matches)

    return best


def _compile_trim_dependencies(
    *,
    project_root: Path,
    profile: dict[str, Any],
) -> tuple[list[DependencyCandidate], list[DependencyIssue]]:
    """Compile allowed (SIZE, IMPELLER_TRIM) pairs from Rev0.3 ConstraintTable4.

    Rev0.3 is the authoritative constraint source and does NOT have the legacy
    per-series size x trim sheets. Its authoritative allowed-trim table is
    ConstraintTable4 (Alt Size + ImpellerTrim + Allowed?), compiled into
    FYBROC_CONSTRAINT_MODEL.json. ConstraintTable4 is series-independent, so
    each allowed (size, trim) pair applies to any series that offers that size;
    series_code is left None on the emitted candidate to reflect that.
    """
    settings = profile["trim"]
    candidates: list[DependencyCandidate] = []
    issues: list[DependencyIssue] = []

    model_path = project_root / settings["constraint_model_path"]
    if not model_path.exists():
        issues.append(
            DependencyIssue(
                severity="Error",
                issue_code="CONSTRAINT_MODEL_MISSING",
                message=(
                    f"Constraint model not found: {settings['constraint_model_path']}. "
                    "Run scripts/compile_fybroc_constraint_model.py first."
                ),
            )
        )
        return candidates, issues

    model = json.loads(model_path.read_text(encoding="utf-8"))

    table_name = settings["constraint_table_name"]
    resolved = None
    for entry in model.get("constraint_index", []):
        if entry.get("table_name") == table_name and entry.get("resolved_table"):
            resolved = entry["resolved_table"]
            break

    if resolved is None:
        issues.append(
            DependencyIssue(
                severity="Error",
                issue_code="CONSTRAINT_TABLE_MISSING",
                message=(
                    f"{table_name} was not found (resolved) in the constraint model."
                ),
                source_reference=str(model_path.name),
            )
        )
        return candidates, issues

    size_header = settings["size_header"]
    trim_header = settings["trim_header"]
    allowed_header = settings["allowed_header"]
    allowed_value = settings["allowed_value"].strip().casefold()

    # Deduplicate (size, trim) while preserving deterministic output.
    seen: set[tuple[str, str]] = set()
    per_size: dict[str, set[str]] = {}

    for row in resolved.get("rows", []):
        size_value = _text(row.get(size_header))
        trim_value = _format_trim(row.get(trim_header))
        verdict = _text(row.get(allowed_header))

        if not size_value or not trim_value:
            continue
        # Only emit ALLOWED pairs (skip explicit "Not Allowed" rows).
        if verdict is not None and verdict.strip().casefold() != allowed_value:
            continue

        key = (size_value, trim_value)
        if key in seen:
            continue
        seen.add(key)
        per_size.setdefault(size_value, set()).add(trim_value)

    if not per_size:
        issues.append(
            DependencyIssue(
                severity="Error",
                issue_code="NO_TRIMS_FOR_SIZE",
                message=(
                    f"{table_name} produced no allowed (size, trim) pairs."
                ),
                source_reference=table_name,
            )
        )
        return candidates, issues

    for size_value in sorted(per_size):
        for trim in sorted(per_size[size_value], key=float):
            candidates.append(
                DependencyCandidate(
                    family_code=profile["family_code"],
                    dependency_code="SIZE_IMPELLER_TRIM",
                    target_field_code="IMPELLER_TRIM",
                    target_display_value=trim,
                    target_identifier_code=None,
                    series_code=None,
                    context_json=json.dumps(
                        {"SIZE": size_value},
                        sort_keys=True,
                        separators=(",", ":"),
                    ),
                    source_workbook="Fybroc Configuration Rev0.3.xlsx",
                    source_worksheet="Feasible Constraints",
                    source_reference=f"{table_name}:{size_value}",
                )
            )

    return candidates, issues


def _compile_motor_modification_dependencies(
    *,
    project_root: Path,
    profile: dict[str, Any],
) -> tuple[list[DependencyCandidate], list[DependencyIssue]]:
    settings = profile["motor_modifications"]
    nomenclature_path = (
        project_root / profile["nomenclature_workbook"]
    )
    price_path = project_root / profile["price_workbook"]

    nomenclature = load_workbook(
        nomenclature_path,
        read_only=False,
        data_only=True,
        keep_links=False,
    )
    price = load_workbook(
        price_path,
        read_only=False,
        data_only=True,
        keep_links=False,
    )

    candidates: list[DependencyCandidate] = []
    issues: list[DependencyIssue] = []

    try:
        attributes = nomenclature[settings["attributes_sheet"]]
        modifications: dict[str, tuple[str, str]] = {}

        for row_number in range(
            int(settings["attribute_row_from"]),
            int(settings["attribute_row_to"]) + 1,
        ):
            description = _text(
                attributes[
                    f"{settings['description_column']}{row_number}"
                ].value
            )
            code = _text(
                attributes[
                    f"{settings['identifier_code_column']}{row_number}"
                ].value
            )
            if description and code:
                modifications[_normal_description(description)] = (
                    description,
                    code,
                )

        aliases = {
            _normal_description(source): target
            for source, target
            in settings.get("description_aliases", {}).items()
        }

        rules_sheet = price[settings["new_rules_sheet"]]
        vendor_support: dict[str, dict[str, Any]] = {}

        for row_number in range(
            int(settings["new_rules_row_from"]),
            int(settings["new_rules_row_to"]) + 1,
        ):
            source_description = _text(
                rules_sheet[
                    f"{settings['new_rules_description_column']}"
                    f"{row_number}"
                ].value
            )
            if not source_description:
                continue

            source_key = _normal_description(source_description)
            target_description = aliases.get(
                source_key,
                source_description,
            )
            target_key = _normal_description(target_description)

            if target_key not in modifications:
                issues.append(
                    DependencyIssue(
                        severity="Warning",
                        issue_code="MODIFICATION_DESCRIPTION_UNMATCHED",
                        message=(
                            "Price rule modification was not found in "
                            f"nomenclature: {source_description}"
                        ),
                        source_reference=(
                            f"{rules_sheet.title}!"
                            f"{settings['new_rules_description_column']}"
                            f"{row_number}"
                        ),
                    )
                )
                continue

            vendor_support[target_key] = {
                vendor: rules_sheet[f"{column}{row_number}"].value
                for vendor, column
                in settings["vendor_columns"].items()
            }

        master = price[settings["mod_master_sheet"]]
        master_by_prefix: dict[str, list[str | None]] = {}

        for row_number in range(
            int(settings["mod_master_row_from"]),
            master.max_row + 1,
        ):
            catalog = _text(
                master[
                    f"{settings['mod_master_catalog_column']}{row_number}"
                ].value
            )
            frame_value = _text(
                master[
                    f"{settings['mod_master_frame_column']}{row_number}"
                ].value
            )
            if not catalog:
                continue

            prefix = catalog.split("_", 1)[0].upper()
            master_by_prefix.setdefault(prefix, []).append(frame_value)

        frame_map = settings["frame_stem_to_class"]
        special_rules = settings.get("special_rules", {})
        all_frame_when_supported = set(
            settings.get("all_frame_when_supported", [])
        )
        no_modification = settings["no_modification_value"]
        included_motor = settings["included_motor_value"]
        no_mod_key = _normal_description(no_modification)

        if no_mod_key not in modifications:
            raise ValueError(
                "No Modification was not found in nomenclature attributes."
            )

        motor_sheet = nomenclature[settings["motor_sheet"]]
        motor_columns = settings["motor_columns"]
        seen_contexts: set[str] = set()

        for row_number in range(
            int(settings["motor_row_from"]),
            int(settings["motor_row_to"]) + 1,
        ):
            context = {
                field: _text(
                    motor_sheet[f"{column}{row_number}"].value
                )
                for field, column in motor_columns.items()
            }

            if any(value is None for value in context.values()):
                continue

            context = {
                key: str(value)
                for key, value in context.items()
            }
            context_json = json.dumps(
                context,
                sort_keys=True,
                separators=(",", ":"),
            )

            if context_json in seen_contexts:
                continue
            seen_contexts.add(context_json)

            no_mod_display, no_mod_code = modifications[no_mod_key]
            candidates.append(
                DependencyCandidate(
                    family_code=profile["family_code"],
                    dependency_code="MOTOR_CONTEXT_MODIFICATION",
                    target_field_code="MOTOR_MODIFICATIONS",
                    target_display_value=no_mod_display,
                    target_identifier_code=no_mod_code,
                    series_code=None,
                    context_json=context_json,
                    source_workbook=nomenclature_path.name,
                    source_worksheet=motor_sheet.title,
                    source_reference=f"{motor_sheet.title}!{row_number}",
                )
            )

            if context["MOTOR_OPTION"] != included_motor:
                continue

            manufacturer = context["MOTOR_MANUFACTURER"]
            frame = context["MOTOR_FRAME"]
            orientation = context["MOTOR_ORIENTATION"]

            if manufacturer == "Fybroc Choice*":
                required_vendors = settings["fybroc_choice_vendors"]
            else:
                mapped = settings["manufacturer_to_vendor"].get(
                    manufacturer
                )
                required_vendors = [mapped] if mapped else []

            if not required_vendors:
                continue

            for mod_key, (display_value, identifier_code) in (
                modifications.items()
            ):
                if mod_key == no_mod_key:
                    continue

                support = vendor_support.get(mod_key)
                if support is None:
                    continue

                if not all(
                    support.get(vendor)
                    for vendor in required_vendors
                ):
                    continue

                special = special_rules.get(display_value, {})
                required_orientation = special.get(
                    "required_orientation"
                )
                required_frame = special.get("required_frame")

                if (
                    required_orientation
                    and orientation != required_orientation
                ):
                    continue

                if required_frame and frame != required_frame:
                    continue

                prefixes: set[str] = set()
                for vendor in required_vendors:
                    prefixes.update(
                        _catalog_prefixes(support.get(vendor))
                    )

                allowed_classes: set[str] = set()
                global_support = False

                for prefix in prefixes:
                    for frame_value in master_by_prefix.get(prefix, []):
                        if frame_value is None:
                            global_support = True
                            continue

                        parsed = _frame_class(frame_value)
                        if parsed:
                            allowed_classes.add(parsed)

                if display_value in all_frame_when_supported:
                    global_support = True

                if required_frame:
                    frame_supported = True
                elif global_support:
                    frame_supported = True
                elif allowed_classes:
                    frame_supported = (
                        frame_map.get(frame) in allowed_classes
                    )
                else:
                    frame_supported = False

                if not frame_supported:
                    continue

                candidates.append(
                    DependencyCandidate(
                        family_code=profile["family_code"],
                        dependency_code="MOTOR_CONTEXT_MODIFICATION",
                        target_field_code="MOTOR_MODIFICATIONS",
                        target_display_value=display_value,
                        target_identifier_code=identifier_code,
                        series_code=None,
                        context_json=context_json,
                        source_workbook=price_path.name,
                        source_worksheet=rules_sheet.title,
                        source_reference=(
                            f"{rules_sheet.title}:{display_value}"
                        ),
                    )
                )

    finally:
        nomenclature.close()
        price.close()

    return candidates, issues


def compile_dependencies(
    *,
    project_root: Path,
    profile_path: Path,
) -> DependencyCompilationReport:
    profile = json.loads(
        profile_path.read_text(encoding="utf-8")
    )

    trim_candidates, trim_issues = _compile_trim_dependencies(
        project_root=project_root,
        profile=profile,
    )
    motor_candidates, motor_issues = (
        _compile_motor_modification_dependencies(
            project_root=project_root,
            profile=profile,
        )
    )

    unique: dict[
        tuple[str, str, str | None, str],
        DependencyCandidate,
    ] = {}

    for candidate in trim_candidates + motor_candidates:
        key = (
            candidate.target_field_code,
            candidate.target_display_value,
            candidate.series_code,
            candidate.context_json,
        )
        unique.setdefault(key, candidate)

    candidates = tuple(unique.values())
    issues = tuple(trim_issues + motor_issues)

    return DependencyCompilationReport(
        candidate_count=len(candidates),
        trim_relation_count=sum(
            item.target_field_code == "IMPELLER_TRIM"
            for item in candidates
        ),
        motor_modification_relation_count=sum(
            item.target_field_code == "MOTOR_MODIFICATIONS"
            for item in candidates
        ),
        issue_count=len(issues),
        candidates=candidates,
        issues=issues,
    )


def save_report(
    report: DependencyCompilationReport,
    output_dir: Path,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    (
        output_dir / "fybroc_dependency_candidates.json"
    ).write_text(
        json.dumps(
            asdict(report),
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    _write_csv(
        output_dir / "fybroc_dependency_candidates.csv",
        [asdict(item) for item in report.candidates],
    )
    _write_csv(
        output_dir / "fybroc_dependency_issues.csv",
        [asdict(item) for item in report.issues],
    )


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(rows[0]),
        )
        writer.writeheader()
        writer.writerows(rows)
