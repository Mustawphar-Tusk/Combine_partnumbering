from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


@dataclass(frozen=True)
class SectionCandidate:
    family_code: str
    workbook_role: str
    worksheet_name: str
    section_code: str
    section_name: str
    display_order: int
    sequence_from: int | None
    sequence_to: int | None
    source_profile: str


@dataclass(frozen=True)
class FieldCandidate:
    family_code: str
    workbook_role: str
    worksheet_name: str
    section_code: str
    field_code: str
    field_name: str
    display_order: int
    source_cell: str | None
    source_sequence: int | None
    extraction_strategy: str
    confidence: str
    source_profile: str


@dataclass(frozen=True)
class CompilationIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    worksheet_name: str | None = None
    source_reference: str | None = None


@dataclass(frozen=True)
class ConfigurationModelReport:
    section_count: int
    field_count: int
    issue_count: int
    sections: tuple[SectionCandidate, ...]
    fields: tuple[FieldCandidate, ...]
    issues: tuple[CompilationIssue, ...]


def normalize_code(value: str) -> str:
    code = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_").upper()
    return code or "UNNAMED"


def load_profiles(profile_directory: Path) -> tuple[dict[str, Any], ...]:
    profiles = []
    for path in sorted(profile_directory.glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["_profile_path"] = str(path)
        profiles.append(payload)

    if not profiles:
        raise FileNotFoundError(
            f"No configuration compiler profiles found in {profile_directory}"
        )

    return tuple(profiles)


def _resolve_workbook(
    project_root: Path,
    discovery: dict[str, Any],
    family_code: str,
    workbook_role: str,
) -> Path:
    matches = [
        record
        for record in discovery["records"]
        if record.get("family_code") == family_code
        and record.get("role") == workbook_role
        and record.get("discovery_status") == "discovered"
    ]

    if len(matches) != 1:
        raise ValueError(
            f"Expected exactly one discovered workbook for "
            f"{family_code}/{workbook_role}; found {len(matches)}."
        )

    return project_root / matches[0]["relative_path"]


def _section_for_sequence(
    sequence: int,
    sections: Iterable[dict[str, Any]],
) -> dict[str, Any] | None:
    for section in sections:
        if int(section["sequence_from"]) <= sequence <= int(section["sequence_to"]):
            return section
    return None


def _nearest_label(
    worksheet,
    row: int,
    sequence_column: int,
    search_width: int,
) -> tuple[str | None, str | None]:
    for column in range(sequence_column + 1, sequence_column + search_width + 1):
        value = worksheet.cell(row=row, column=column).value
        if isinstance(value, str) and value.strip():
            return value.strip(), worksheet.cell(row=row, column=column).coordinate
    return None, None


def compile_numbered_rows(
    project_root: Path,
    discovery: dict[str, Any],
    profile: dict[str, Any],
) -> tuple[list[SectionCandidate], list[FieldCandidate], list[CompilationIssue]]:
    family_code = profile["family_code"].upper()
    workbook_role = profile["workbook_role"]
    worksheet_name = profile["worksheet_name"]
    strategy = profile["strategy"]
    profile_path = profile["_profile_path"]

    workbook_path = _resolve_workbook(
        project_root,
        discovery,
        family_code,
        workbook_role,
    )

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        keep_links=True,
    )

    sections: list[SectionCandidate] = []
    fields: list[FieldCandidate] = []
    issues: list[CompilationIssue] = []

    try:
        if worksheet_name not in workbook.sheetnames:
            return [], [], [
                CompilationIssue(
                    family_code=family_code,
                    severity="Error",
                    issue_code="WORKSHEET_NOT_FOUND",
                    message=f"Worksheet '{worksheet_name}' was not found.",
                    worksheet_name=worksheet_name,
                )
            ]

        worksheet = workbook[worksheet_name]

        section_definitions = strategy["sections"]
        for order, section in enumerate(section_definitions, start=1):
            sections.append(
                SectionCandidate(
                    family_code=family_code,
                    workbook_role=workbook_role,
                    worksheet_name=worksheet_name,
                    section_code=section["section_code"],
                    section_name=section["section_name"],
                    display_order=order * 10,
                    sequence_from=int(section["sequence_from"]),
                    sequence_to=int(section["sequence_to"]),
                    source_profile=profile_path,
                )
            )

        minimum = int(strategy["minimum_sequence"])
        maximum = int(strategy["maximum_sequence"])
        scan_columns = int(strategy.get("scan_max_columns", 20))
        label_search_width = int(strategy.get("label_search_width", 6))
        scan_max_rows = int(strategy.get("scan_max_rows", worksheet.max_row or 500))

        found_sequences: set[int] = set()

        for row in range(1, scan_max_rows + 1):
            for column in range(1, scan_columns + 1):
                value = worksheet.cell(row=row, column=column).value
                if isinstance(value, bool):
                    continue
                if isinstance(value, int) and minimum <= value <= maximum:
                    sequence = value
                elif isinstance(value, float) and value.is_integer() and minimum <= value <= maximum:
                    sequence = int(value)
                else:
                    continue

                if sequence in found_sequences:
                    continue

                label, label_cell = _nearest_label(
                    worksheet,
                    row=row,
                    sequence_column=column,
                    search_width=label_search_width,
                )

                section = _section_for_sequence(sequence, section_definitions)
                if section is None:
                    issues.append(
                        CompilationIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="SECTION_NOT_RESOLVED",
                            message=f"No section covers sequence {sequence}.",
                            worksheet_name=worksheet_name,
                            source_reference=worksheet.cell(row=row, column=column).coordinate,
                        )
                    )
                    continue

                if label is None:
                    label = f"Field {sequence}"
                    confidence = "low"
                    issues.append(
                        CompilationIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="FIELD_LABEL_NOT_FOUND",
                            message=f"No adjacent label was found for sequence {sequence}.",
                            worksheet_name=worksheet_name,
                            source_reference=worksheet.cell(row=row, column=column).coordinate,
                        )
                    )
                else:
                    confidence = "high"

                fields.append(
                    FieldCandidate(
                        family_code=family_code,
                        workbook_role=workbook_role,
                        worksheet_name=worksheet_name,
                        section_code=section["section_code"],
                        field_code=f"{normalize_code(label)}_{sequence:03d}",
                        field_name=label,
                        display_order=sequence,
                        source_cell=label_cell,
                        source_sequence=sequence,
                        extraction_strategy="numbered_rows",
                        confidence=confidence,
                        source_profile=profile_path,
                    )
                )
                found_sequences.add(sequence)

        for sequence in range(minimum, maximum + 1):
            if sequence not in found_sequences:
                issues.append(
                    CompilationIssue(
                        family_code=family_code,
                        severity="Warning",
                        issue_code="SEQUENCE_NOT_FOUND",
                        message=f"Sequence {sequence} was not located in the worksheet.",
                        worksheet_name=worksheet_name,
                    )
                )

    finally:
        workbook.close()

    return sections, fields, issues


def compile_column_workflow(
    project_root: Path,
    discovery: dict[str, Any],
    profile: dict[str, Any],
) -> tuple[list[SectionCandidate], list[FieldCandidate], list[CompilationIssue]]:
    family_code = profile["family_code"].upper()
    workbook_role = profile["workbook_role"]
    worksheet_name = profile["worksheet_name"]
    profile_path = profile["_profile_path"]
    strategy = profile["strategy"]

    workbook_path = _resolve_workbook(
        project_root,
        discovery,
        family_code,
        workbook_role,
    )

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        keep_links=True,
    )

    sections: list[SectionCandidate] = []
    fields: list[FieldCandidate] = []
    issues: list[CompilationIssue] = []

    try:
        if worksheet_name not in workbook.sheetnames:
            return [], [], [
                CompilationIssue(
                    family_code=family_code,
                    severity="Error",
                    issue_code="WORKSHEET_NOT_FOUND",
                    message=f"Worksheet '{worksheet_name}' was not found.",
                    worksheet_name=worksheet_name,
                )
            ]

        for section_order, section in enumerate(strategy["sections"], start=1):
            sections.append(
                SectionCandidate(
                    family_code=family_code,
                    workbook_role=workbook_role,
                    worksheet_name=worksheet_name,
                    section_code=section["section_code"],
                    section_name=section["section_name"],
                    display_order=section_order * 10,
                    sequence_from=None,
                    sequence_to=None,
                    source_profile=profile_path,
                )
            )

            for field in section["fields"]:
                columns = field["columns"]
                source_reference = ":".join(columns) if len(columns) > 1 else columns[0]
                fields.append(
                    FieldCandidate(
                        family_code=family_code,
                        workbook_role=workbook_role,
                        worksheet_name=worksheet_name,
                        section_code=section["section_code"],
                        field_code=field["field_code"],
                        field_name=field["field_name"],
                        display_order=int(field["display_order"]),
                        source_cell=source_reference,
                        source_sequence=None,
                        extraction_strategy="column_workflow",
                        confidence="high",
                        source_profile=profile_path,
                    )
                )

                for column in columns:
                    try:
                        column_index_from_string(column)
                    except ValueError:
                        issues.append(
                            CompilationIssue(
                                family_code=family_code,
                                severity="Error",
                                issue_code="INVALID_COLUMN_REFERENCE",
                                message=f"Invalid Excel column reference '{column}'.",
                                worksheet_name=worksheet_name,
                                source_reference=column,
                            )
                        )

    finally:
        workbook.close()

    return sections, fields, issues


def compile_numbered_column_pairs(
    project_root: Path,
    discovery: dict[str, Any],
    profile: dict[str, Any],
) -> tuple[
    list[SectionCandidate],
    list[FieldCandidate],
    list[CompilationIssue],
]:
    family_code = profile["family_code"].upper()
    workbook_role = profile["workbook_role"]
    worksheet_name = profile["worksheet_name"]
    strategy = profile["strategy"]
    profile_path = profile["_profile_path"]

    workbook_path = _resolve_workbook(
        project_root,
        discovery,
        family_code,
        workbook_role,
    )

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        keep_links=True,
    )

    sections: list[SectionCandidate] = []
    fields: list[FieldCandidate] = []
    issues: list[CompilationIssue] = []

    try:
        if worksheet_name not in workbook.sheetnames:
            return [], [], [
                CompilationIssue(
                    family_code=family_code,
                    severity="Error",
                    issue_code="WORKSHEET_NOT_FOUND",
                    message=f"Worksheet '{worksheet_name}' was not found.",
                    worksheet_name=worksheet_name,
                )
            ]

        worksheet = workbook[worksheet_name]
        section_definitions = strategy["sections"]

        for order, section in enumerate(section_definitions, start=1):
            sections.append(
                SectionCandidate(
                    family_code=family_code,
                    workbook_role=workbook_role,
                    worksheet_name=worksheet_name,
                    section_code=section["section_code"],
                    section_name=section["section_name"],
                    display_order=order * 10,
                    sequence_from=int(section["sequence_from"]),
                    sequence_to=int(section["sequence_to"]),
                    source_profile=profile_path,
                )
            )

        found_sequences: set[int] = set()

        for pair in strategy["column_pairs"]:
            sequence_column = column_index_from_string(
                pair["sequence_column"]
            )
            label_column = column_index_from_string(
                pair["label_column"]
            )

            sequence_from = int(pair["sequence_from"])
            sequence_to = int(pair["sequence_to"])
            row_from = int(pair.get("row_from", 1))
            row_to = int(pair.get("row_to", worksheet.max_row or 500))

            for row in range(row_from, row_to + 1):
                raw_sequence = worksheet.cell(
                    row=row,
                    column=sequence_column,
                ).value

                if isinstance(raw_sequence, bool):
                    continue

                if isinstance(raw_sequence, int):
                    sequence = raw_sequence
                elif (
                    isinstance(raw_sequence, float)
                    and raw_sequence.is_integer()
                ):
                    sequence = int(raw_sequence)
                elif (
                    isinstance(raw_sequence, str)
                    and raw_sequence.strip().isdigit()
                ):
                    sequence = int(raw_sequence.strip())
                else:
                    continue

                if not sequence_from <= sequence <= sequence_to:
                    continue

                if sequence in found_sequences:
                    issues.append(
                        CompilationIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="DUPLICATE_SEQUENCE",
                            message=(
                                f"Sequence {sequence} was found more than once."
                            ),
                            worksheet_name=worksheet_name,
                            source_reference=worksheet.cell(
                                row=row,
                                column=sequence_column,
                            ).coordinate,
                        )
                    )
                    continue

                label_cell = worksheet.cell(
                    row=row,
                    column=label_column,
                )
                raw_label = label_cell.value

                if raw_label is None or str(raw_label).strip() == "":
                    label = f"Field {sequence}"
                    confidence = "low"

                    issues.append(
                        CompilationIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="FIELD_LABEL_NOT_FOUND",
                            message=(
                                f"No label was found for sequence {sequence}."
                            ),
                            worksheet_name=worksheet_name,
                            source_reference=label_cell.coordinate,
                        )
                    )
                else:
                    label = str(raw_label).strip()
                    confidence = "high"

                section = _section_for_sequence(
                    sequence,
                    section_definitions,
                )

                if section is None:
                    issues.append(
                        CompilationIssue(
                            family_code=family_code,
                            severity="Error",
                            issue_code="SECTION_NOT_RESOLVED",
                            message=(
                                f"No section covers sequence {sequence}."
                            ),
                            worksheet_name=worksheet_name,
                            source_reference=label_cell.coordinate,
                        )
                    )
                    continue

                fields.append(
                    FieldCandidate(
                        family_code=family_code,
                        workbook_role=workbook_role,
                        worksheet_name=worksheet_name,
                        section_code=section["section_code"],
                        field_code=(
                            f"{normalize_code(label)}_{sequence:03d}"
                        ),
                        field_name=label,
                        display_order=sequence,
                        source_cell=label_cell.coordinate,
                        source_sequence=sequence,
                        extraction_strategy="numbered_column_pairs",
                        confidence=confidence,
                        source_profile=profile_path,
                    )
                )

                found_sequences.add(sequence)

        minimum = int(strategy["minimum_sequence"])
        maximum = int(strategy["maximum_sequence"])

        for sequence in range(minimum, maximum + 1):
            if sequence not in found_sequences:
                issues.append(
                    CompilationIssue(
                        family_code=family_code,
                        severity="Warning",
                        issue_code="SEQUENCE_NOT_FOUND",
                        message=f"Sequence {sequence} was not located.",
                        worksheet_name=worksheet_name,
                    )
                )

    finally:
        workbook.close()

    return sections, fields, issues


def compile_configuration_model(
    project_root: Path,
    discovery_path: Path,
    profile_directory: Path,
) -> ConfigurationModelReport:
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))
    profiles = load_profiles(profile_directory)

    sections: list[SectionCandidate] = []
    fields: list[FieldCandidate] = []
    issues: list[CompilationIssue] = []

    # strategy_handlers = {
    #     "numbered_rows": compile_numbered_rows,
    #     "column_workflow": compile_column_workflow,
    # }

    strategy_handlers = {
        "numbered_rows": compile_numbered_rows,
        "numbered_column_pairs": compile_numbered_column_pairs,
        "column_workflow": compile_column_workflow,
            }

    for profile in profiles:
        strategy_type = profile["strategy"]["type"]
        handler = strategy_handlers.get(strategy_type)

        if handler is None:
            issues.append(
                CompilationIssue(
                    family_code=profile["family_code"].upper(),
                    severity="Error",
                    issue_code="UNKNOWN_STRATEGY",
                    message=f"Unsupported strategy '{strategy_type}'.",
                )
            )
            continue

        profile_sections, profile_fields, profile_issues = handler(
            project_root,
            discovery,
            profile,
        )
        sections.extend(profile_sections)
        fields.extend(profile_fields)
        issues.extend(profile_issues)

    return ConfigurationModelReport(
        section_count=len(sections),
        field_count=len(fields),
        issue_count=len(issues),
        sections=tuple(sections),
        fields=tuple(fields),
        issues=tuple(issues),
    )


def save_report(report: ConfigurationModelReport, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv_exports(report: ConfigurationModelReport, export_directory: Path) -> None:
    export_directory.mkdir(parents=True, exist_ok=True)
    _write_csv(
        export_directory / "configuration_sections_candidates.csv",
        [asdict(item) for item in report.sections],
    )
    _write_csv(
        export_directory / "configuration_fields_candidates.csv",
        [asdict(item) for item in report.fields],
    )
    _write_csv(
        export_directory / "configuration_model_issues.csv",
        [asdict(item) for item in report.issues],
    )


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
