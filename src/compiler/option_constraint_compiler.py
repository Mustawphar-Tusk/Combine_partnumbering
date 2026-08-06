from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer


LEGACY_INVALID_MARKERS = {"ERR", "TBD__", "____"}


@dataclass(frozen=True)
class OptionCandidate:
    family_code: str
    workbook_role: str
    worksheet_name: str
    field_code: str
    option_code: str
    option_description: str
    hex_code: str | None
    source_cell: str
    source_range: str | None
    source_type: str
    confidence: str
    source_profile: str


@dataclass(frozen=True)
class DependencyCandidate:
    family_code: str
    parent_field_code: str
    child_field_code: str
    dependency_type: str
    source_formula: str | None
    source_reference: str
    confidence: str
    source_profile: str


@dataclass(frozen=True)
class ConstraintCandidate:
    family_code: str
    rule_code: str
    rule_name: str
    rule_type: str
    condition_json: str
    action_json: str
    source_workbook_role: str
    source_worksheet: str
    source_reference: str
    confidence: str
    source_profile: str


@dataclass(frozen=True)
class LegacyMarkerDiagnostic:
    family_code: str
    worksheet_name: str
    source_reference: str
    marker: str
    formula: str | None
    meaning: str


@dataclass(frozen=True)
class CompilerIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    worksheet_name: str | None = None
    source_reference: str | None = None


@dataclass(frozen=True)
class OptionConstraintReport:
    option_count: int
    dependency_count: int
    constraint_count: int
    legacy_marker_count: int
    issue_count: int
    options: tuple[OptionCandidate, ...]
    dependencies: tuple[DependencyCandidate, ...]
    constraints: tuple[ConstraintCandidate, ...]
    legacy_markers: tuple[LegacyMarkerDiagnostic, ...]
    issues: tuple[CompilerIssue, ...]


def normalize_code(value: str) -> str:
    code = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_").upper()
    return code or "UNNAMED"


def load_profiles(profile_directory: Path) -> tuple[dict[str, Any], ...]:
    profiles: list[dict[str, Any]] = []

    for path in sorted(profile_directory.glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["_profile_path"] = str(path)
        profiles.append(payload)

    if not profiles:
        raise FileNotFoundError(
            f"No option/constraint compiler profiles found in {profile_directory}"
        )

    return tuple(profiles)


def _resolve_workbook(
    project_root: Path,
    discovery: dict[str, Any],
    family_code: str,
    workbook_role: str,
) -> Path:
    matches = [
        row
        for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == workbook_role
        and row.get("discovery_status") == "discovered"
    ]

    if len(matches) != 1:
        raise ValueError(
            f"Expected one workbook for {family_code}/{workbook_role}; "
            f"found {len(matches)}."
        )

    return project_root / matches[0]["relative_path"]


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _formula_references(formula: str) -> tuple[str, ...]:
    references: list[str] = []

    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return ()

    for token in tokens:
        if token.subtype != "RANGE":
            continue

        value = token.value.replace("$", "")
        if "!" in value:
            references.append(value)
        elif re.fullmatch(r"[A-Z]{1,3}\d+(?::[A-Z]{1,3}\d+)?", value):
            references.append(value)

    return tuple(dict.fromkeys(references))


def _extract_list_pairs(
    workbook,
    family_code: str,
    workbook_role: str,
    profile: dict[str, Any],
) -> tuple[list[OptionCandidate], list[CompilerIssue]]:
    options: list[OptionCandidate] = []
    issues: list[CompilerIssue] = []
    source_profile = profile["_profile_path"]

    for source in profile.get("option_sources", []):
        worksheet_name = source["worksheet_name"]

        if worksheet_name not in workbook.sheetnames:
            issues.append(
                CompilerIssue(
                    family_code=family_code,
                    severity="Error",
                    issue_code="OPTION_WORKSHEET_NOT_FOUND",
                    message=f"Worksheet '{worksheet_name}' was not found.",
                    worksheet_name=worksheet_name,
                )
            )
            continue

        worksheet = workbook[worksheet_name]
        field_code = source["field_code"]
        description_column = source["description_column"]
        code_column = source.get("option_code_column")
        hex_column = source.get("hex_column")
        row_from = int(source["row_from"])
        row_to = int(source.get("row_to", worksheet.max_row or row_from))
        source_range = (
            f"{description_column}{row_from}:"
            f"{hex_column or code_column or description_column}{row_to}"
        )

        seen: set[tuple[str, str | None]] = set()

        for row in range(row_from, row_to + 1):
            description_cell = worksheet[f"{description_column}{row}"]
            description = _cell_text(description_cell.value)

            if not description:
                continue

            raw_option_code = (
                _cell_text(worksheet[f"{code_column}{row}"].value)
                if code_column
                else None
            )
            hex_code = (
                _cell_text(worksheet[f"{hex_column}{row}"].value)
                if hex_column
                else None
            )

            if hex_code in LEGACY_INVALID_MARKERS:
                continue

            option_code = raw_option_code or normalize_code(description)
            key = (option_code, hex_code)

            if key in seen:
                continue

            seen.add(key)

            options.append(
                OptionCandidate(
                    family_code=family_code,
                    workbook_role=workbook_role,
                    worksheet_name=worksheet_name,
                    field_code=field_code,
                    option_code=option_code,
                    option_description=description,
                    hex_code=hex_code,
                    source_cell=description_cell.coordinate,
                    source_range=source_range,
                    source_type="profile_list",
                    confidence="high",
                    source_profile=source_profile,
                )
            )

    return options, issues


def _extract_dependency_sources(
    workbook,
    family_code: str,
    profile: dict[str, Any],
) -> tuple[
    list[DependencyCandidate],
    list[ConstraintCandidate],
    list[LegacyMarkerDiagnostic],
    list[CompilerIssue],
]:
    dependencies: list[DependencyCandidate] = []
    constraints: list[ConstraintCandidate] = []
    legacy_markers: list[LegacyMarkerDiagnostic] = []
    issues: list[CompilerIssue] = []
    source_profile = profile["_profile_path"]

    for source in profile.get("dependency_sources", []):
        worksheet_name = source["worksheet_name"]

        if worksheet_name not in workbook.sheetnames:
            issues.append(
                CompilerIssue(
                    family_code=family_code,
                    severity="Error",
                    issue_code="DEPENDENCY_WORKSHEET_NOT_FOUND",
                    message=f"Worksheet '{worksheet_name}' was not found.",
                    worksheet_name=worksheet_name,
                )
            )
            continue

        worksheet = workbook[worksheet_name]
        parent_fields = tuple(source.get("parent_field_codes", []))
        child_field = source["child_field_code"]
        formula_cells = source.get("formula_cells", [])
        dependency_type = source.get("dependency_type", "filters_options")

        for coordinate in formula_cells:
            raw_value = worksheet[coordinate].value
            formula = (
                raw_value
                if isinstance(raw_value, str) and raw_value.startswith("=")
                else None
            )
            references = _formula_references(formula or "")

            for parent_field in parent_fields:
                dependencies.append(
                    DependencyCandidate(
                        family_code=family_code,
                        parent_field_code=parent_field,
                        child_field_code=child_field,
                        dependency_type=dependency_type,
                        source_formula=formula,
                        source_reference=coordinate,
                        confidence="medium" if formula else "low",
                        source_profile=source_profile,
                    )
                )

            if formula:
                upper_formula = formula.upper()

                for marker in LEGACY_INVALID_MARKERS:
                    if marker not in upper_formula:
                        continue

                    legacy_markers.append(
                        LegacyMarkerDiagnostic(
                            family_code=family_code,
                            worksheet_name=worksheet_name,
                            source_reference=coordinate,
                            marker=marker,
                            formula=formula,
                            meaning=(
                                "Legacy workbook invalid/incomplete-state signal. "
                                "Not valid configuration metadata."
                            ),
                        )
                    )

                    rule_code = (
                        f"PREVENT_{normalize_code(child_field)}_"
                        f"{normalize_code(marker)}_{coordinate}"
                    )

                    constraints.append(
                        ConstraintCandidate(
                            family_code=family_code,
                            rule_code=rule_code,
                            rule_name=(
                                f"Prevent invalid {child_field} selection"
                            ),
                            rule_type="reject",
                            condition_json=json.dumps(
                                {
                                    "requires_resolution": True,
                                    "legacy_marker": marker,
                                    "formula_references": references,
                                },
                                separators=(",", ":"),
                            ),
                            action_json=json.dumps(
                                {
                                    "action": "reject_configuration",
                                    "target_field_code": child_field,
                                    "message": (
                                        f"Selection for {child_field} is "
                                        "incompatible or incomplete."
                                    ),
                                },
                                separators=(",", ":"),
                            ),
                            source_workbook_role=profile["workbook_role"],
                            source_worksheet=worksheet_name,
                            source_reference=coordinate,
                            confidence="low",
                            source_profile=source_profile,
                        )
                    )

    for explicit_rule in profile.get("explicit_constraints", []):
        constraints.append(
            ConstraintCandidate(
                family_code=family_code,
                rule_code=explicit_rule["rule_code"],
                rule_name=explicit_rule["rule_name"],
                rule_type=explicit_rule["rule_type"],
                condition_json=json.dumps(
                    explicit_rule["condition"],
                    separators=(",", ":"),
                ),
                action_json=json.dumps(
                    explicit_rule["action"],
                    separators=(",", ":"),
                ),
                source_workbook_role=profile["workbook_role"],
                source_worksheet=explicit_rule["source_worksheet"],
                source_reference=explicit_rule["source_reference"],
                confidence=explicit_rule.get("confidence", "high"),
                source_profile=source_profile,
            )
        )

    return dependencies, constraints, legacy_markers, issues


def compile_option_constraints(
    project_root: Path,
    discovery_path: Path,
    profile_directory: Path,
) -> OptionConstraintReport:
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))
    profiles = load_profiles(profile_directory)

    all_options: list[OptionCandidate] = []
    all_dependencies: list[DependencyCandidate] = []
    all_constraints: list[ConstraintCandidate] = []
    all_legacy_markers: list[LegacyMarkerDiagnostic] = []
    all_issues: list[CompilerIssue] = []

    for profile in profiles:
        family_code = profile["family_code"].upper()
        workbook_role = profile["workbook_role"]
        workbook_path = _resolve_workbook(
            project_root,
            discovery,
            family_code,
            workbook_role,
        )

        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=False,
            keep_vba=workbook_path.suffix.lower() == ".xlsm",
            keep_links=True,
        )

        try:
            options, option_issues = _extract_list_pairs(
                workbook,
                family_code,
                workbook_role,
                profile,
            )
            (
                dependencies,
                constraints,
                legacy_markers,
                dependency_issues,
            ) = _extract_dependency_sources(
                workbook,
                family_code,
                profile,
            )

            all_options.extend(options)
            all_dependencies.extend(dependencies)
            all_constraints.extend(constraints)
            all_legacy_markers.extend(legacy_markers)
            all_issues.extend(option_issues)
            all_issues.extend(dependency_issues)

        finally:
            workbook.close()

    return OptionConstraintReport(
        option_count=len(all_options),
        dependency_count=len(all_dependencies),
        constraint_count=len(all_constraints),
        legacy_marker_count=len(all_legacy_markers),
        issue_count=len(all_issues),
        options=tuple(all_options),
        dependencies=tuple(all_dependencies),
        constraints=tuple(all_constraints),
        legacy_markers=tuple(all_legacy_markers),
        issues=tuple(all_issues),
    )


def save_json(report: OptionConstraintReport, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(asdict(report), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def save_csv(report: OptionConstraintReport, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)

    _write_csv(
        directory / "configuration_option_candidates.csv",
        [asdict(row) for row in report.options],
    )
    _write_csv(
        directory / "configuration_dependency_candidates.csv",
        [asdict(row) for row in report.dependencies],
    )
    _write_csv(
        directory / "constraint_rule_candidates.csv",
        [asdict(row) for row in report.constraints],
    )
    _write_csv(
        directory / "legacy_invalid_marker_diagnostics.csv",
        [asdict(row) for row in report.legacy_markers],
    )
    _write_csv(
        directory / "option_constraint_issues.csv",
        [asdict(row) for row in report.issues],
    )


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
