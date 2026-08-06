from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class SeriesFieldOptionCandidate:
    family_code: str
    source_field_code: str
    field_code: str
    option_value: str
    series_code: str
    workbook_name: str
    worksheet_name: str
    source_row: int
    source_field_cell: str
    source_value_cell: str
    source_series_cell: str
    source_profile: str


@dataclass(frozen=True)
class SeriesFieldOptionIssue:
    family_code: str
    severity: str
    issue_code: str
    message: str
    source_reference: str | None = None


@dataclass(frozen=True)
class SeriesConstraintReport:
    field_count: int
    option_count: int
    relation_count: int
    issue_count: int
    candidates: tuple[SeriesFieldOptionCandidate, ...]
    issues: tuple[SeriesFieldOptionIssue, ...]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_workbook(
    project_root: Path,
    discovery_path: Path,
    family_code: str,
    workbook_role: str,
) -> Path:
    discovery = json.loads(
        discovery_path.read_text(encoding="utf-8")
    )

    matches = [
        row
        for row in discovery["records"]
        if row.get("family_code") == family_code
        and row.get("role") == workbook_role
        and row.get("discovery_status") == "discovered"
    ]

    if len(matches) != 1:
        raise ValueError(
            f"Expected one discovered workbook for "
            f"{family_code}/{workbook_role}; found {len(matches)}."
        )

    return project_root / matches[0]["relative_path"]


def compile_series_constraints(
    project_root: Path,
    discovery_path: Path,
    profile_path: Path,
) -> SeriesConstraintReport:
    profile = json.loads(
        profile_path.read_text(encoding="utf-8")
    )

    family_code = profile["family_code"]
    workbook_path = _resolve_workbook(
        project_root,
        discovery_path,
        family_code,
        profile["workbook_role"],
    )

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    candidates: list[SeriesFieldOptionCandidate] = []
    issues: list[SeriesFieldOptionIssue] = []
    seen: set[tuple[str, str, str]] = set()

    try:
        sheet_name = profile["worksheet_name"]

        if sheet_name not in workbook.sheetnames:
            return SeriesConstraintReport(
                field_count=0,
                option_count=0,
                relation_count=0,
                issue_count=1,
                candidates=(),
                issues=(
                    SeriesFieldOptionIssue(
                        family_code=family_code,
                        severity="Error",
                        issue_code="WORKSHEET_NOT_FOUND",
                        message=f"Worksheet '{sheet_name}' was not found.",
                    ),
                ),
            )

        worksheet = workbook[sheet_name]
        field_column = profile["field_code_column"]
        value_column = profile["option_value_column"]
        field_map = profile.get("field_code_map", {})
        series_columns = profile["series_columns"]

        for row in range(
            int(profile["row_from"]),
            (worksheet.max_row or 0) + 1,
        ):
            source_field_code = _text(
                worksheet[f"{field_column}{row}"].value
            )
            option_value = _text(
                worksheet[f"{value_column}{row}"].value
            )

            if not source_field_code and not option_value:
                continue

            if not source_field_code or not option_value:
                issues.append(
                    SeriesFieldOptionIssue(
                        family_code=family_code,
                        severity="Warning",
                        issue_code="INCOMPLETE_MATRIX_ROW",
                        message=(
                            f"Row {row} is missing field code or option value."
                        ),
                        source_reference=f"{sheet_name}!{row}",
                    )
                )
                continue

            field_code = field_map.get(
                source_field_code,
                source_field_code.removeprefix("F_").upper(),
            )

            relation_found = False

            for column, configured_series in series_columns.items():
                marker = _text(
                    worksheet[f"{column}{row}"].value
                )

                if marker is None:
                    continue

                relation_found = True
                series_code = configured_series

                if marker.isdigit() and marker != configured_series:
                    issues.append(
                        SeriesFieldOptionIssue(
                            family_code=family_code,
                            severity="Warning",
                            issue_code="SERIES_MARKER_MISMATCH",
                            message=(
                                f"Configured series {configured_series} "
                                f"does not match marker {marker}."
                            ),
                            source_reference=f"{sheet_name}!{column}{row}",
                        )
                    )
                    series_code = marker

                key = (
                    field_code,
                    option_value.casefold(),
                    series_code,
                )

                if key in seen:
                    continue

                seen.add(key)

                candidates.append(
                    SeriesFieldOptionCandidate(
                        family_code=family_code,
                        source_field_code=source_field_code,
                        field_code=field_code,
                        option_value=option_value,
                        series_code=series_code,
                        workbook_name=workbook_path.name,
                        worksheet_name=sheet_name,
                        source_row=row,
                        source_field_cell=f"{field_column}{row}",
                        source_value_cell=f"{value_column}{row}",
                        source_series_cell=f"{column}{row}",
                        source_profile=str(profile_path),
                    )
                )

            if not relation_found:
                issues.append(
                    SeriesFieldOptionIssue(
                        family_code=family_code,
                        severity="Warning",
                        issue_code="NO_SERIES_RELATION",
                        message=(
                            f"No series applicability marker was found for "
                            f"{source_field_code}='{option_value}'."
                        ),
                        source_reference=f"{sheet_name}!{row}",
                    )
                )

    finally:
        workbook.close()

    return SeriesConstraintReport(
        field_count=len(
            {row.field_code for row in candidates}
        ),
        option_count=len(
            {
                (row.field_code, row.option_value.casefold())
                for row in candidates
            }
        ),
        relation_count=len(candidates),
        issue_count=len(issues),
        candidates=tuple(candidates),
        issues=tuple(issues),
    )


def save_report(
    report: SeriesConstraintReport,
    output_dir: Path,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    (
        output_dir
        / "fybroc_series_constraint_candidates.json"
    ).write_text(
        json.dumps(
            asdict(report),
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    _write_csv(
        output_dir
        / "fybroc_series_constraint_candidates.csv",
        [asdict(row) for row in report.candidates],
    )

    _write_csv(
        output_dir
        / "fybroc_series_constraint_issues.csv",
        [asdict(row) for row in report.issues],
    )


def _write_csv(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(rows[0]),
        )
        writer.writeheader()
        writer.writerows(rows)
