from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from src.compiler.model_option_applicability_compiler import (
    compile_model_option_applicability,
)


def _write_fixture(tmp_path: Path) -> Path:
    workbooks = tmp_path / "workbooks" / "Dean"
    exports = tmp_path / "exports"
    profiles = tmp_path / "config" / "applicability_profiles"

    workbooks.mkdir(parents=True)
    exports.mkdir(parents=True)
    profiles.mkdir(parents=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pump Options"

    ws["A3"] = "A Number"
    ws["B3"] = "Series"
    ws["C3"] = "Size"

    # First accepted model
    ws["A4"] = "A779"
    ws["B4"] = "RA2096"
    ws["C4"] = "1x1.5x6"

    # Pump Material spans D:G. D itself is the FIRST option column.
    ws["D1"] = "Pump Material"
    ws["D4"] = "STD"
    ws["E4"] = "X"
    ws["F4"] = None
    ws["G4"] = "O"

    # Second accepted model
    ws["A5"] = "A770"
    ws["B5"] = "RA3146"
    ws["C5"] = "1x2x11.5"
    ws["D5"] = "X"
    ws["E5"] = "STD"

    # Blocked model: must never appear in candidates.
    ws["A6"] = "MDL1-.75"
    ws["B6"] = "Deanline"
    ws["C6"] = "0.75x0.75"
    ws["D6"] = "STD"

    workbook_path = workbooks / "PumpConfiguration_Logic.xlsx"
    wb.save(workbook_path)

    reconciliation = {
        "reconciliation_version": "M023.1.1-TEST",
        "summary": {"error_count": 0},
        "field_reconciliation": [
            {
                "canonical_field_code": "PUMP_MATERIAL",
                "labels": ["Pump Material"],
            }
        ],
        "logic_option_domains": [
            {
                "table_name": "Table88",
                "field_label": "Pump Material",
                "headers": ["Pump Material"],
                "row_count": 4,
                "rows": [
                    {"Pump Material": "Ductile Iron"},
                    {"Pump Material": "Cast Steel"},
                    {"Pump Material": "316 S/S"},
                    {"Pump Material": "Custom"},
                ],
            }
        ],
        "pump_options_matrix": {
            "groups": [
                {
                    "field_label": "Pump Material",
                    "canonical_field_code": "PUMP_MATERIAL",
                    "group_column": "D",
                    "span_end_column": "G",
                    "option_columns": [],
                }
            ],
            "models": [
                {
                    "model_key": "RA2096|1X1.5X6",
                    "model_identifier": "A779",
                    "base_identifier": "D779",
                    "series": "RA2096",
                    "size": "1x1.5x6",
                    "source_row": 4,
                },
                {
                    "model_key": "RA3146|1X2X11.5",
                    "model_identifier": "A770",
                    "base_identifier": "D770",
                    "series": "RA3146",
                    "size": "1x2x11.5",
                    "source_row": 5,
                },
                {
                    "model_key": "DEANLINE|0.75X0.75",
                    "model_identifier": "MDL1-.75",
                    "base_identifier": None,
                    "series": "Deanline",
                    "size": "0.75x0.75",
                    "source_row": 6,
                },
            ],
        },
        "model_reconciliation": [
            {
                "model_key": "RA2096|1X1.5X6",
                "classification": "SHARED_MATCH",
                "authoritative_model_identifier": "A779",
                "authoritative_base_identifier": "D779",
                "logic_model_identifier": "A779",
                "series": "RA2096",
                "size": "1x1.5x6",
            },
            {
                "model_key": "RA3146|1X2X11.5",
                "classification": "LOGIC_ONLY_SUPPLEMENTAL_READY",
                "authoritative_model_identifier": "A770",
                "authoritative_base_identifier": "D770",
                "logic_model_identifier": "A770",
                "series": "RA3146",
                "size": "1x2x11.5",
            },
            {
                "model_key": "DEANLINE|0.75X0.75",
                "classification": "LOGIC_ONLY_REVIEW",
                "authoritative_model_identifier": None,
                "authoritative_base_identifier": None,
                "logic_model_identifier": "MDL1-.75",
                "series": "Deanline",
                "size": "0.75x0.75",
            },
        ],
    }

    (exports / "m023_dean_source_reconciliation.json").write_text(
        json.dumps(reconciliation),
        encoding="utf-8",
    )

    profile = {
        "family_code": "DEAN",
        "version": "M023.2-TEST",
        "reconciliation_json": "exports/m023_dean_source_reconciliation.json",
        "workbook": "workbooks/Dean/PumpConfiguration_Logic.xlsx",
        "worksheet": "Pump Options",
        "accepted_model_classifications": [
            "SHARED_MATCH",
            "SHARED_ALIAS_MATCH",
            "SHARED_REV2_OVERRIDE",
            "LOGIC_ONLY_SUPPLEMENTAL_READY",
        ],
        "blocked_model_classifications": ["LOGIC_ONLY_REVIEW"],
        "marker_semantics": {
            "STD": {
                "status": "STANDARD",
                "is_allowed": True,
                "is_standard": True,
                "publication_status": "READY",
            },
            "X": {
                "status": "AVAILABLE",
                "is_allowed": True,
                "is_standard": False,
                "publication_status": "READY",
            },
            "O": {
                "status": "REVIEW",
                "is_allowed": False,
                "is_standard": False,
                "publication_status": "BLOCKED_REVIEW",
            },
        },
        "unknown_marker_policy": {
            "status": "REVIEW",
            "is_allowed": False,
            "is_standard": False,
            "publication_status": "BLOCKED_REVIEW",
        },
        "normalization": {"uppercase_marker": True},
        "safety": {
            "require_reconciled_base_identifier": True,
        },
    }

    profile_path = profiles / "dean.json"
    profile_path.write_text(json.dumps(profile), encoding="utf-8")
    return profile_path


def test_group_start_column_is_first_option(tmp_path: Path) -> None:
    profile = _write_fixture(tmp_path)
    report = compile_model_option_applicability(
        project_root=tmp_path,
        profile_path=profile,
    )

    rows = [
        row
        for row in report["candidates"]
        if row["model_key"] == "RA2096|1X1.5X6"
    ]

    by_option = {row["option_value"]: row for row in rows}

    assert by_option["Ductile Iron"]["marker_normalized"] == "STD"
    assert by_option["Ductile Iron"]["is_standard"] is True
    assert by_option["Ductile Iron"]["is_allowed"] is True

    assert by_option["Cast Steel"]["marker_normalized"] == "X"
    assert by_option["Cast Steel"]["is_allowed"] is True

    assert "316 S/S" not in by_option

    assert by_option["Custom"]["marker_normalized"] == "O"
    assert by_option["Custom"]["is_allowed"] is False
    assert by_option["Custom"]["publication_status"] == "BLOCKED_REVIEW"


def test_supplemental_a_number_gets_reconciled_base(tmp_path: Path) -> None:
    profile = _write_fixture(tmp_path)
    report = compile_model_option_applicability(
        project_root=tmp_path,
        profile_path=profile,
    )

    rows = [
        row
        for row in report["candidates"]
        if row["model_key"] == "RA3146|1X2X11.5"
    ]

    assert rows
    assert {row["base_identifier"] for row in rows} == {"D770"}


def test_blocked_deanline_never_enters_candidates(tmp_path: Path) -> None:
    profile = _write_fixture(tmp_path)
    report = compile_model_option_applicability(
        project_root=tmp_path,
        profile_path=profile,
    )

    candidate_keys = {row["model_key"] for row in report["candidates"]}
    assert "DEANLINE|0.75X0.75" not in candidate_keys

    blocked = [
        row for row in report["reviews"] if row["review_type"] == "BLOCKED_MODEL"
    ]
    assert len(blocked) == 1
    assert blocked[0]["model_key"] == "DEANLINE|0.75X0.75"


def test_publication_gate_passes_when_only_o_requires_review(tmp_path: Path) -> None:
    profile = _write_fixture(tmp_path)
    report = compile_model_option_applicability(
        project_root=tmp_path,
        profile_path=profile,
    )

    assert report["summary"]["error_count"] == 0
    assert report["summary"]["publication_gate"] == "PASS"
    assert report["summary"]["marker_counts"]["O"] == 1


def test_logic_original_model_key_resolves_reconciled_series_alias(tmp_path: Path) -> None:
    profile = _write_fixture(tmp_path)

    reconciliation_path = (
        tmp_path / "exports" / "m023_dean_source_reconciliation.json"
    )
    reconciliation = json.loads(
        reconciliation_path.read_text(encoding="utf-8")
    )

    # Replace the second accepted model with the same pattern used by the
    # real Dean PH2170/PH3170 reconciliation: authoritative/effective key
    # differs from the original Pump Options key.
    reconciliation["pump_options_matrix"]["models"][1] = {
        "model_key": "PH3170|1.5X3X13.5",
        "model_identifier": "A415",
        "base_identifier": "D415",
        "series": "PH3170",
        "size": "1.5x3x13.5",
        "source_row": 5,
    }

    reconciliation["model_reconciliation"][1] = {
        "model_key": "PH2170|1.5X3X13.5",
        "classification": "SHARED_ALIAS_MATCH",
        "authoritative_model_identifier": "A415",
        "authoritative_base_identifier": "D415",
        "rev2_model_identifier": "A415",
        "logic_model_identifier": "A415",
        "logic_original_model_key": "PH3170|1.5X3X13.5",
        "logic_effective_model_identifier": "A415",
        "series": "PH2170",
        "size": "1.5x3x13.5",
        "logic_original_series": "PH3170",
        "logic_original_size": "1.5x3x13.5",
        "publication_status": "READY",
    }

    reconciliation_path.write_text(
        json.dumps(reconciliation),
        encoding="utf-8",
    )

    workbook_path = (
        tmp_path / "workbooks" / "Dean" / "PumpConfiguration_Logic.xlsx"
    )

    from openpyxl import load_workbook

    wb = load_workbook(workbook_path)
    try:
        ws = wb["Pump Options"]
        ws["A5"] = "A415"
        ws["B5"] = "PH3170"
        ws["C5"] = "1.5x3x13.5"
        wb.save(workbook_path)
    finally:
        wb.close()

    report = compile_model_option_applicability(
        project_root=tmp_path,
        profile_path=profile,
    )

    rows = [
        row
        for row in report["candidates"]
        if row["model_key"] == "PH3170|1.5X3X13.5"
    ]

    assert rows
    assert {row["base_identifier"] for row in rows} == {"D415"}
    assert {
        row["model_reconciliation_classification"] for row in rows
    } == {"SHARED_ALIAS_MATCH"}
