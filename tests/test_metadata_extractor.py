from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from src.compiler.metadata_extractor import extract_workbook_metadata


def test_metadata_extractor(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Configurator"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=SUM(A1:A2)"
    ws.merge_cells("B1:C1")
    validation = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add("D1")
    wb.create_named_range("TotalCell", ws, "$A$3")
    wb.save(path)

    result = extract_workbook_metadata(path, "TEST", "sample.xlsx")
    sheet = result.worksheets[0]
    assert sheet.formula_count == 1
    assert sheet.merged_range_count == 1
    assert sheet.validation_count == 1
    assert len(result.defined_names) == 1
