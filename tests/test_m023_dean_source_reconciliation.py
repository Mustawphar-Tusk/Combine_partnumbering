from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.compiler.configuration_source_reconciliation import (
    canonical_field_code,
    compile_configuration_source_reconciliation,
)


def _add_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def test_canonical_field_code_is_stable() -> None:
    assert canonical_field_code("Cooling Plan Extras") == "COOLING_PLAN_EXTRAS"
    assert canonical_field_code("Flush / Barrier & Cooling") == "FLUSH_BARRIER_AND_COOLING"


def test_dean_reconciliation_preserves_rev2_authority_and_logic_gap(tmp_path: Path) -> None:
    workbooks = tmp_path / "workbooks" / "Dean"
    config = tmp_path / "config"
    compiler_profiles = config / "compiler_profiles"
    identifier_profiles = config / "identifier_profiles"
    reconciliation_profiles = config / "reconciliation_profiles"

    for path in (
        workbooks,
        compiler_profiles,
        identifier_profiles,
        reconciliation_profiles,
    ):
        path.mkdir(parents=True, exist_ok=True)

    rev2_path = workbooks / "Dean Data Sheet Rev 2.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sheet"
    ws["B2"] = 1
    ws["C2"] = "Pump Material"
    ws["B3"] = 2
    ws["C3"] = "Rev Only Field"

    ref = wb.create_sheet("Reference Data")
    ref["A2"] = "A779"
    ref["B2"] = "RA2096"
    ref["C2"] = "1x1.5x6"
    wb.save(rev2_path)

    logic_path = workbooks / "PumpConfiguration_Logic.xlsx"
    wb = Workbook()
    cfg = wb.active
    cfg.title = "Config Options"
    cfg["A1"] = "Pump Material"
    cfg["A2"] = "Ductile Iron"
    cfg["A3"] = "316 S/S"
    _add_table(cfg, "A1:A3", "Table88")
    cfg["C1"] = "Logic Only Field"
    cfg["C2"] = "Option A"
    cfg["C3"] = "Option B"
    _add_table(cfg, "C1:C3", "Table200")

    dep = wb.create_sheet("Codependencies")
    dep["A1"] = "Pump Material"
    dep["B1"] = "Logic Only Field"
    dep["A2"] = "Ductile Iron"
    dep["B2"] = "Option A"
    _add_table(dep, "A1:B2", "Table300")

    pump = wb.create_sheet("Pump Options")
    pump["A3"] = "A Number"
    pump["B3"] = "Series"
    pump["C3"] = "Size"
    pump["D1"] = "=Table88[[#Headers],[Pump Material]]"
    pump["E2"] = "Ductile Iron"
    pump["F2"] = "316 S/S"
    pump["H1"] = "=Table200[[#Headers],[Logic Only Field]]"
    pump["I2"] = "Option A"
    pump["J2"] = "Option B"
    pump["A4"] = "A779"
    pump["B4"] = "RA2096"
    pump["C4"] = "1x1.5x6"
    pump["E4"] = "STD"
    pump["F4"] = "X"
    pump["I4"] = "STD"
    pump["J4"] = "X"

    price = wb.create_sheet("Price Options")
    price["A3"] = "A Number"
    price["B3"] = "Series"
    price["C3"] = "Size"
    price["D1"] = "=Table88[[#Headers],[Pump Material]]"
    price["E2"] = "Ductile Iron"
    price["F2"] = "316 S/S"
    price["A4"] = "A779"
    price["B4"] = "RA2096"
    price["C4"] = "1x1.5x6"
    price["E4"] = 0
    price["F4"] = 100
    wb.save(logic_path)

    compiler_profile = {
        "family_code": "DEAN",
        "worksheet_name": "Data Sheet",
        "strategy": {
            "column_pairs": [
                {
                    "sequence_column": "B",
                    "label_column": "C",
                    "sequence_from": 1,
                    "sequence_to": 2,
                    "row_from": 1,
                    "row_to": 5,
                }
            ],
            "sections": [
                {
                    "section_code": "TEST",
                    "section_name": "Test",
                    "sequence_from": 1,
                    "sequence_to": 2,
                }
            ],
        },
    }
    (compiler_profiles / "dean_configuration.json").write_text(
        json.dumps(compiler_profile), encoding="utf-8"
    )

    identifier_profile = {
        "family_code": "DEAN",
        "model_reference": {
            "worksheet_name": "Reference Data",
            "row_from": 2,
            "columns": {
                "model_identifier": "A",
                "series": "B",
                "size": "C",
            },
        },
    }
    (identifier_profiles / "dean.json").write_text(
        json.dumps(identifier_profile), encoding="utf-8"
    )

    reconciliation_profile = {
        "family_code": "DEAN",
        "reconciliation_version": "M023.1-TEST",
        "rev2": {
            "workbook": str(rev2_path.relative_to(tmp_path)),
            "configuration_compiler_profile": str(
                (compiler_profiles / "dean_configuration.json").relative_to(tmp_path)
            ),
            "identifier_profile": str(
                (identifier_profiles / "dean.json").relative_to(tmp_path)
            ),
        },
        "logic": {
            "workbook": str(logic_path.relative_to(tmp_path)),
            "config_options_sheet": "Config Options",
            "codependencies_sheet": "Codependencies",
            "pump_options_sheet": "Pump Options",
            "price_options_sheet": "Price Options",
        },
        "base_identifier": {
            "from_prefix": "A",
            "to_prefix": "D",
        },
        "field_aliases": {},
    }
    profile_path = reconciliation_profiles / "dean.json"
    profile_path.write_text(json.dumps(reconciliation_profile), encoding="utf-8")

    report = compile_configuration_source_reconciliation(
        project_root=tmp_path,
        profile_path=profile_path,
    )

    by_code = {
        item["canonical_field_code"]: item
        for item in report["field_reconciliation"]
    }

    assert by_code["PUMP_MATERIAL"]["classification"] == "SHARED"
    assert by_code["REV_ONLY_FIELD"]["classification"] == "REV2_ONLY"
    assert by_code["LOGIC_ONLY_FIELD"]["classification"] == "LOGIC_ONLY_SUPPLEMENTAL"

    models = report["model_reconciliation"]
    assert len(models) == 1
    assert models[0]["classification"] == "SHARED_MATCH"
    assert models[0]["authoritative_model_identifier"] == "A779"
    assert models[0]["authoritative_base_identifier"] == "D779"
    assert report["summary"]["error_count"] == 0


def test_model_reconciliation_rules_are_metadata_driven() -> None:
    from src.compiler.configuration_source_reconciliation import _reconcile_models

    rev2_models = [
        {
            "model_key": "PHP2140|2X2X10",
            "model_identifier": "A351P",
            "base_identifier": "D351P",
            "series": "PHP2140",
            "size": "2x2x10",
        },
        {
            "model_key": "PH2170|1.5X3X13.5",
            "model_identifier": "A415",
            "base_identifier": "D415",
            "series": "PH2170",
            "size": "1.5x3x13.5",
        },
        {
            "model_key": "RA3246|6X8X15.5",
            "model_identifier": "A799",
            "base_identifier": "D799",
            "series": "RA3246",
            "size": "6x8x15.5",
        },
    ]
    logic_models = [
        {
            "model_key": "PHP2140|2X2X10",
            "model_identifier": "A351-PHP",
            "base_identifier": "D351-PHP",
            "series": "PHP2140",
            "size": "2x2x10",
        },
        {
            "model_key": "PH3170|1.5X3X13.5",
            "model_identifier": "A415",
            "base_identifier": "D415",
            "series": "PH3170",
            "size": "1.5x3x13.5",
        },
        {
            "model_key": "RA3246|6X8X15.5",
            "model_identifier": "A798",
            "base_identifier": "D798",
            "series": "RA3246",
            "size": "6x8x15.5",
        },
        {
            "model_key": "RA3146|1X2X11.5",
            "model_identifier": "A770",
            "base_identifier": "D770",
            "series": "RA3146",
            "size": "1x2x11.5",
        },
        {
            "model_key": "DEANLINE|0.75X0.75",
            "model_identifier": "MDL1-.75",
            "base_identifier": None,
            "series": "Deanline",
            "size": "0.75x0.75",
        },
    ]
    issues = []
    rules = {
        "a_number_prefix": "A",
        "allow_logic_only_a_prefix_models": True,
        "logic_identifier_aliases": {"A351-PHP": "A351P"},
        "logic_key_aliases": [
            {
                "logic_model_identifier": "A415",
                "logic_series": "PH3170",
                "logic_size": "1.5x3x13.5",
                "rev2_series": "PH2170",
                "rev2_size": "1.5x3x13.5",
                "reason": "approved series alias",
            }
        ],
        "approved_rev2_overrides": [
            {
                "series": "RA3246",
                "size": "6x8x15.5",
                "rev2_model_identifier": "A799",
                "logic_model_identifier": "A798",
                "reason": "Rev2 verified in two sources",
            }
        ],
    }

    rows = _reconcile_models(rev2_models, logic_models, issues, rules)
    by_key = {row["model_key"]: row for row in rows}

    assert by_key["PHP2140|2X2X10"]["classification"] == "SHARED_ALIAS_MATCH"
    assert by_key["PH2170|1.5X3X13.5"]["classification"] == "SHARED_ALIAS_MATCH"
    assert by_key["RA3246|6X8X15.5"]["classification"] == "SHARED_REV2_OVERRIDE"
    assert by_key["RA3246|6X8X15.5"]["authoritative_model_identifier"] == "A799"
    assert by_key["RA3146|1X2X11.5"]["classification"] == "LOGIC_ONLY_SUPPLEMENTAL_READY"
    assert by_key["RA3146|1X2X11.5"]["authoritative_base_identifier"] == "D770"
    assert by_key["DEANLINE|0.75X0.75"]["classification"] == "LOGIC_ONLY_REVIEW"
    assert not any(issue.severity == "Error" for issue in issues)


def test_expected_rev2_sequence_gap_is_not_an_error(tmp_path: Path) -> None:
    from src.compiler.configuration_source_reconciliation import _validate_rev2_sequences

    profile_path = tmp_path / "dean_configuration.json"
    profile_path.write_text(
        json.dumps(
            {
                "strategy": {
                    "minimum_sequence": 49,
                    "maximum_sequence": 55,
                    "column_pairs": [
                        {
                            "sequence_from": 49,
                            "sequence_to": 55,
                        }
                    ],
                }
            }
        ),
        encoding="utf-8",
    )
    fields = [{"sequence": value} for value in (49, 50, 51, 52, 54, 55)]
    issues = []

    result = _validate_rev2_sequences(profile_path, fields, [53], issues)

    assert result["missing_sequences"] == [53]
    assert result["unexpected_missing_sequences"] == []
    assert not any(issue.severity == "Error" for issue in issues)


def test_field_aliases_merge_naming_variants() -> None:
    from src.compiler.configuration_source_reconciliation import _canonical_with_alias

    aliases = {
        "Auxillary Nameplate": "Auxiliary Nameplate",
        "Casing Wear Ring": "Casing Wear Rings",
        "Frame Size": "Frame",
    }

    assert _canonical_with_alias("Auxillary Nameplate", aliases) == "AUXILIARY_NAMEPLATE"
    assert _canonical_with_alias("Casing Wear Ring", aliases) == "CASING_WEAR_RINGS"
    assert _canonical_with_alias("Frame Size", aliases) == "FRAME"
