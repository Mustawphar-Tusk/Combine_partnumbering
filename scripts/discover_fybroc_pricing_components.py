from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

WORKBOOK = (
    ROOT
    / "workbooks"
    / "Fybroc"
    / "Price Estimator-Fybroc.xlsm"
)

OUTPUT = (
    ROOT
    / "exports"
    / "fybroc_price_check_component_inventory.json"
)

SHEET = "Price Check"

START_ROW = 80
END_ROW = 102

DISPLAY_COLUMNS = (
    "A",
    "B",
    "C",
    "D",
    "E",
    "F",
    "G",
    "H",
)


CELL_REFERENCE_RE = re.compile(
    r"""
    (?:
        '
        (?P<quoted_sheet>[^']+)
        '
        |
        (?P<plain_sheet>
            [A-Za-z0-9_\- ]+
        )
    )
    !
    (?P<cell>
        \$?[A-Z]{1,3}
        \$?\d+
    )
    """,
    re.VERBOSE,
)


LOCAL_REFERENCE_RE = re.compile(
    r"""
    (?<![A-Za-z0-9_!])
    (?P<cell>
        \$?[A-Z]{1,3}
        \$?\d+
    )
    """,
    re.VERBOSE,
)


def clean_value(value):
    if value is None:
        return None

    if isinstance(
        value,
        (
            str,
            int,
            float,
            bool,
        ),
    ):
        return value

    return str(value)


def direct_references(
    formula: object,
    current_sheet: str,
) -> list[dict[str, str]]:

    if (
        not isinstance(formula, str)
        or not formula.startswith("=")
    ):
        return []

    refs: list[dict[str, str]] = []
    occupied_spans: list[tuple[int, int]] = []

    for match in CELL_REFERENCE_RE.finditer(
        formula
    ):
        sheet = (
            match.group("quoted_sheet")
            or match.group("plain_sheet")
        )

        refs.append(
            {
                "sheet": sheet.strip(),
                "cell": (
                    match
                    .group("cell")
                    .replace("$", "")
                ),
            }
        )

        occupied_spans.append(
            match.span()
        )

    def inside_sheet_reference(
        start: int,
        end: int,
    ) -> bool:
        return any(
            start >= left
            and end <= right
            for left, right in occupied_spans
        )

    for match in LOCAL_REFERENCE_RE.finditer(
        formula
    ):
        if inside_sheet_reference(
            *match.span()
        ):
            continue

        refs.append(
            {
                "sheet": current_sheet,
                "cell": (
                    match
                    .group("cell")
                    .replace("$", "")
                ),
            }
        )

    unique: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    for ref in refs:
        key = (
            ref["sheet"],
            ref["cell"],
        )

        if key in seen:
            continue

        seen.add(key)
        unique.append(ref)

    return unique


def referenced_values(
    workbook,
    refs: list[dict[str, str]],
) -> list[dict[str, object]]:

    results = []

    for ref in refs:
        sheet = ref["sheet"]
        cell = ref["cell"]

        if sheet not in workbook.sheetnames:
            results.append(
                {
                    **ref,
                    "exists": False,
                    "value": None,
                }
            )
            continue

        value = workbook[sheet][cell].value

        results.append(
            {
                **ref,
                "exists": True,
                "value": clean_value(value),
            }
        )

    return results


def main() -> None:

    if not WORKBOOK.exists():
        raise FileNotFoundError(
            WORKBOOK
        )

    formula_wb = load_workbook(
        WORKBOOK,
        data_only=False,
        keep_vba=True,
        read_only=False,
    )

    value_wb = load_workbook(
        WORKBOOK,
        data_only=True,
        keep_vba=True,
        read_only=False,
    )

    if SHEET not in formula_wb.sheetnames:
        raise RuntimeError(
            f"Worksheet not found: {SHEET}"
        )

    formula_ws = formula_wb[SHEET]
    value_ws = value_wb[SHEET]

    inventory = []

    for row in range(
        START_ROW,
        END_ROW + 1,
    ):

        row_values = {}

        for col in DISPLAY_COLUMNS:
            address = f"{col}{row}"

            row_values[col] = {
                "formula_or_value":
                    clean_value(
                        formula_ws[
                            address
                        ].value
                    ),
                "cached_value":
                    clean_value(
                        value_ws[
                            address
                        ].value
                    ),
            }

        pricing_cells = {}

        for col in ("D", "F"):
            address = f"{col}{row}"

            formula = (
                formula_ws[address].value
            )

            refs = direct_references(
                formula,
                SHEET,
            )

            pricing_cells[col] = {
                "address": address,
                "formula":
                    clean_value(formula),
                "cached_value":
                    clean_value(
                        value_ws[
                            address
                        ].value
                    ),
                "direct_references":
                    referenced_values(
                        formula_wb,
                        refs,
                    ),
            }

        inventory.append(
            {
                "row": row,
                "cells": row_values,
                "pricing": pricing_cells,
            }
        )

    payload = {
        "workbook": WORKBOOK.name,
        "worksheet": SHEET,
        "start_row": START_ROW,
        "end_row": END_ROW,
        "rows": inventory,
    }

    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    print("=" * 100)
    print(
        "M022.1 FYBROC PRICE CHECK COMPONENT INVENTORY"
    )
    print("=" * 100)

    for item in inventory:
        row = item["row"]
        cells = item["cells"]

        print()
        print("-" * 100)
        print(f"ROW {row}")
        print("-" * 100)

        for col in DISPLAY_COLUMNS:
            info = cells[col]

            print(
                f"{col}{row:<3} "
                f"formula/value = "
                f"{info['formula_or_value']!r}"
            )

            if (
                info["cached_value"]
                != info["formula_or_value"]
            ):
                print(
                    f"     "
                    f"cached       = "
                    f"{info['cached_value']!r}"
                )

        for col in ("D", "F"):
            pricing = (
                item["pricing"][col]
            )

            print()
            print(
                f"{pricing['address']} "
                "direct references:"
            )

            if not pricing[
                "direct_references"
            ]:
                print(
                    "    <none>"
                )
                continue

            for ref in pricing[
                "direct_references"
            ]:
                print(
                    "    "
                    f"{ref['sheet']}!"
                    f"{ref['cell']} "
                    f"= {ref['value']!r}"
                )

    print()
    print("=" * 100)
    print(
        f"Inventory rows : "
        f"{len(inventory)}"
    )
    print(
        f"Output         : {OUTPUT}"
    )
    print("=" * 100)


if __name__ == "__main__":
    main()
