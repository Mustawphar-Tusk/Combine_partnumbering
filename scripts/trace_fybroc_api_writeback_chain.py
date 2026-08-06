from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKBOOK_PATH = Path(
    r"workbooks\Fybroc\Price Estimator-Fybroc.xlsm"
)

OUTPUT_PATH = Path(
    r"output\m020_precise_writeback_chain.txt"
)

TARGETS = (
    ("Price Check", "F5"),
    ("Tables", "F33"),
    ("Price Check", "F6"),
    ("Price Check", "F7"),
    ("Price Check", "F9"),
    ("Price Check", "F18"),
)


def display(value: Any) -> str:
    if value is None:
        return "<blank>"

    return str(value).replace("\n", "\\n")


def merged_range(
    worksheet,
    coordinate: str,
) -> str | None:
    for item in worksheet.merged_cells.ranges:
        if coordinate in item:
            return str(item)

    return None


def split_coordinate(
    coordinate: str,
) -> tuple[str, str]:
    match = re.fullmatch(
        r"([A-Z]+)([0-9]+)",
        coordinate.upper(),
    )

    if match is None:
        raise ValueError(
            f"Unsupported coordinate: {coordinate}"
        )

    return match.group(1), match.group(2)


def qualified_reference_pattern(
    sheet_name: str,
    coordinate: str,
) -> re.Pattern[str]:
    column, row = split_coordinate(coordinate)

    escaped_sheet = re.escape(sheet_name)
    escaped_column = re.escape(column)
    escaped_row = re.escape(row)

    return re.compile(
        rf"(?<![A-Z0-9_])"
        rf"(?:'{escaped_sheet}'|{escaped_sheet})!"
        rf"\$?{escaped_column}\$?{escaped_row}"
        rf"(?![0-9])",
        flags=re.IGNORECASE,
    )


def local_reference_pattern(
    coordinate: str,
) -> re.Pattern[str]:
    column, row = split_coordinate(coordinate)

    return re.compile(
        rf"(?<![A-Z0-9_!])"
        rf"\$?{re.escape(column)}"
        rf"\$?{re.escape(row)}"
        rf"(?![0-9])",
        flags=re.IGNORECASE,
    )


def formula_references(
    *,
    formula: str,
    formula_sheet: str,
    target_sheet: str,
    target_coordinate: str,
) -> bool:
    qualified = qualified_reference_pattern(
        target_sheet,
        target_coordinate,
    )

    if qualified.search(formula):
        return True

    if formula_sheet != target_sheet:
        return False

    local = local_reference_pattern(
        target_coordinate
    )

    return local.search(formula) is not None


def find_dependents(
    workbook,
    *,
    target_sheet: str,
    target_coordinate: str,
) -> list[tuple[str, str, str]]:
    results: list[tuple[str, str, str]] = []

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                formula = cell.value

                if not (
                    isinstance(formula, str)
                    and formula.startswith("=")
                ):
                    continue

                if formula_references(
                    formula=formula,
                    formula_sheet=worksheet.title,
                    target_sheet=target_sheet,
                    target_coordinate=target_coordinate,
                ):
                    results.append(
                        (
                            worksheet.title,
                            cell.coordinate,
                            formula,
                        )
                    )

    return results


def find_formal_quote_labels(
    workbook,
) -> list[tuple[str, str]]:
    if "Formal Quote" not in workbook.sheetnames:
        return []

    worksheet = workbook["Formal Quote"]

    pattern = re.compile(
        r"\bpart\s*(?:number|no\.?|#)"
        r"|\bsku\b"
        r"|stock keeping",
        flags=re.IGNORECASE,
    )

    results: list[tuple[str, str]] = []

    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value

            if value is None:
                continue

            text = str(value)

            if pattern.search(text):
                results.append(
                    (
                        cell.coordinate,
                        text,
                    )
                )

    return results


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Workbook not found: {WORKBOOK_PATH}"
        )

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=False,
        data_only=False,
        keep_vba=True,
        keep_links=True,
    )

    lines: list[str] = []

    lines.append("=" * 100)
    lines.append(
        "M020 PRECISE FYBROC WRITEBACK CHAIN"
    )
    lines.append("=" * 100)
    lines.append(f"Workbook: {WORKBOOK_PATH}")
    lines.append("")

    for sheet_name, coordinate in TARGETS:
        lines.append("=" * 100)
        lines.append(
            f"TARGET: {sheet_name}!{coordinate}"
        )
        lines.append("=" * 100)

        if sheet_name not in workbook.sheetnames:
            lines.append("Worksheet not found.")
            lines.append("")
            continue

        worksheet = workbook[sheet_name]
        cell = worksheet[coordinate]

        lines.append(
            f"Value: {display(cell.value)}"
        )
        lines.append(
            f"Data type: {cell.data_type}"
        )
        lines.append(
            "Merged range: "
            f"{merged_range(worksheet, coordinate)}"
        )

        dependents = find_dependents(
            workbook,
            target_sheet=sheet_name,
            target_coordinate=coordinate,
        )

        lines.append(
            f"Exact dependents: {len(dependents)}"
        )

        for (
            dependent_sheet,
            dependent_coordinate,
            formula,
        ) in dependents:
            lines.append(
                f"  {dependent_sheet}!"
                f"{dependent_coordinate}"
            )
            lines.append(
                f"    {display(formula)}"
            )

        lines.append("")

    lines.append("=" * 100)
    lines.append(
        "FORMAL QUOTE PART-NUMBER / SKU TEXT"
    )
    lines.append("=" * 100)

    formal_quote_matches = (
        find_formal_quote_labels(workbook)
    )

    lines.append(
        f"Matches: {len(formal_quote_matches)}"
    )

    for coordinate, value in formal_quote_matches:
        lines.append(
            f"  Formal Quote!{coordinate}: "
            f"{display(value)}"
        )

    OUTPUT_PATH.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )

    print("\n".join(lines))
    print()
    print(f"Report saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
