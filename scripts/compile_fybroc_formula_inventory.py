"""Fybroc Authoritative Formula Inventory.

Extracts, for BOTH authoritative workbooks (Rev0.3 + V6), the complete logic
layer that drives configuration constraints, so constraint-building is guided
by the workbook's own formulas rather than inferred from data shapes:

  - defined names (LAMBDA functions + named ranges/spills)
  - per-sheet distinct formulas (deduplicated by a normalized shape), with a
    representative cell, a count, and the functions used
  - data validations (dropdown/list sources) per sheet
  - table (ListObject) definitions per sheet

Read-only. Writes:
  docs/evidence/F120/FYBROC_FORMULA_INVENTORY.json
  docs/evidence/F120/FYBROC_FORMULA_INVENTORY.txt
"""
from __future__ import annotations

import json
import re
import subprocess
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import warnings
warnings.simplefilter("ignore")

ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = {
    "Rev0.3": "workbooks/Fybroc/Fybroc Configuration Rev0.3.xlsx",
    "V6": "workbooks/Fybroc/Nomenclature_V6.xlsm",
}
OUT_DIR = ROOT / "docs" / "evidence" / "F120"

# Bound per-sheet scanning so huge data sheets (100k+ rows) don't stall - the
# LOGIC (formulas, headers, validations) lives in the top band; data rows below
# are repetitive. We still capture every DISTINCT formula shape from the band.
MAX_SCAN_ROWS = 260
FUNC_RE = re.compile(r"_xlfn\.(?:_xlws\.)?([A-Z][A-Z0-9.]*)|(?<![A-Za-z0-9_.])([A-Z][A-Z0-9]+)\(")


def git_info() -> tuple[str, bool]:
    try:
        commit = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT,
                                capture_output=True, text=True, check=True).stdout.strip()
        status = subprocess.run(["git", "status", "--porcelain"], cwd=ROOT,
                                capture_output=True, text=True, check=True).stdout
        return commit, (status.strip() == "")
    except Exception:
        return "unknown", False


def _formula_text(cell) -> str | None:
    v = cell.value
    if v is None:
        return None
    if hasattr(v, "text"):  # ArrayFormula
        return str(v.text)
    if isinstance(v, str) and v.startswith("="):
        return v
    return None


def _functions_used(formula: str) -> list[str]:
    funcs = set()
    for m in FUNC_RE.finditer(formula):
        name = m.group(1) or m.group(2)
        if name:
            funcs.add(name.rstrip("."))
    return sorted(funcs)


def _normalize_shape(formula: str) -> str:
    """Collapse cell references to '#' so structurally-identical formulas that
    differ only by which cell they point at dedupe to one shape."""
    s = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "#", formula)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def extract_defined_names(wb) -> list[dict]:
    names = []
    dn = wb.defined_names
    for name in dn.keys():
        if name.startswith("_xl"):
            continue
        attr = str(dn[name].attr_text)
        kind = "lambda" if "LAMBDA" in attr else (
            "spill_anchor" if "ANCHORARRAY" in attr else "range_or_formula")
        names.append({"name": name, "kind": kind, "definition": attr})
    return names


def extract_sheet(ws) -> dict[str, Any]:
    max_r = min(ws.max_row or 0, MAX_SCAN_ROWS)
    max_c = ws.max_column or 0

    shapes: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=max_c):
        for cell in row:
            f = _formula_text(cell)
            if not f:
                continue
            shape = _normalize_shape(f)
            entry = shapes.get(shape)
            if entry is None:
                shapes[shape] = {
                    "example_cell": cell.coordinate,
                    "example_formula": f[:400],
                    "functions": _functions_used(f),
                    "count": 1,
                }
            else:
                entry["count"] += 1

    # Data validations (dropdown/list sources)
    validations = []
    try:
        for dv in ws.data_validations.dataValidation:
            validations.append({
                "type": dv.type,
                "formula1": (dv.formula1 or "")[:200],
                "ranges": str(dv.sqref),
            })
    except Exception:
        pass

    # Tables (ListObjects)
    tables = []
    try:
        for tname, tref in getattr(ws, "tables", {}).items():
            tables.append({"name": tname, "ref": str(tref)})
    except Exception:
        pass

    return {
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "scanned_rows": max_r,
        "distinct_formula_shapes": len(shapes),
        "formulas": sorted(shapes.values(), key=lambda e: -e["count"]),
        "data_validations": validations,
        "tables": tables,
    }


def build() -> dict[str, Any]:
    commit, clean = git_info()
    result: dict[str, Any] = {
        "artifact": "FYBROC_FORMULA_INVENTORY",
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit,
        "git_working_tree_clean": clean,
        "note": (
            "Authoritative formula/logic inventory for constraint-building. "
            "Formulas deduped by normalized shape (cell refs -> '#'). Data rows "
            f"scanned to {MAX_SCAN_ROWS}; logic lives in the top band."
        ),
        "workbooks": {},
    }
    for label, rel in WORKBOOKS.items():
        path = ROOT / rel
        keep_vba = path.suffix.lower() == ".xlsm"
        wb = openpyxl.load_workbook(path, data_only=False, keep_vba=keep_vba)
        wb_entry = {
            "file": rel,
            "defined_names": extract_defined_names(wb),
            "sheets": {},
        }
        for sheet in wb.sheetnames:
            wb_entry["sheets"][sheet] = extract_sheet(wb[sheet])
        wb.close()
        result["workbooks"][label] = wb_entry
    return result


def write_outputs(result: dict[str, Any]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "FYBROC_FORMULA_INVENTORY.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    lines = []
    line = "=" * 100
    lines.append(f"{line}\r\nFYBROC AUTHORITATIVE FORMULA INVENTORY\r\n{line}\r\n\r\n")
    lines.append(f"Git commit: {result['git_commit']}  clean={result['git_working_tree_clean']}\r\n\r\n")
    for label, wb in result["workbooks"].items():
        lines.append(f"{line}\r\nWORKBOOK: {label}  ({wb['file']})\r\n{line}\r\n\r\n")
        dn = wb["defined_names"]
        lambdas = [d for d in dn if d["kind"] == "lambda"]
        lines.append(f"Defined names: {len(dn)}  (LAMBDAs: {len(lambdas)})\r\n")
        for d in lambdas:
            lines.append(f"  LAMBDA {d['name']}: {d['definition'][:160]}\r\n")
        lines.append("\r\n")
        for sname, s in wb["sheets"].items():
            if s["distinct_formula_shapes"] == 0 and not s["data_validations"] and not s["tables"]:
                continue
            lines.append(f"--- [{label}] {sname}  ({s['max_row']}x{s['max_column']}, "
                         f"{s['distinct_formula_shapes']} formula shapes) ---\r\n")
            for f in s["formulas"][:12]:
                lines.append(f"    [{f['count']}x] {f['example_cell']}: {f['example_formula'][:120]}\r\n")
            if s["data_validations"]:
                lines.append(f"    data validations: {len(s['data_validations'])}\r\n")
                for dv in s["data_validations"][:8]:
                    lines.append(f"      {dv['ranges']} <- {dv['formula1']}\r\n")
            if s["tables"]:
                lines.append(f"    tables: {', '.join(t['name'] for t in s['tables'])}\r\n")
            lines.append("\r\n")
    (OUT_DIR / "FYBROC_FORMULA_INVENTORY.txt").write_text("".join(lines), encoding="utf-8")


def main() -> int:
    result = build()
    write_outputs(result)
    for label, wb in result["workbooks"].items():
        nlam = sum(1 for d in wb["defined_names"] if d["kind"] == "lambda")
        nform = sum(s["distinct_formula_shapes"] for s in wb["sheets"].values())
        print(f"{label}: {len(wb['sheets'])} sheets, {len(wb['defined_names'])} defined names "
              f"({nlam} LAMBDAs), {nform} distinct formula shapes")
    print(f"Written to {OUT_DIR}/FYBROC_FORMULA_INVENTORY.(json|txt)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
