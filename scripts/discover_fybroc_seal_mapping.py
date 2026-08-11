from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

PRICE_WORKBOOK = (
    ROOT
    / "workbooks"
    / "Fybroc"
    / "Price Estimator-Fybroc.xlsm"
)

NOMENCLATURE_WORKBOOK = (
    ROOT
    / "workbooks"
    / "Fybroc"
    / "Fybroc Nomenclature_V5.xlsm"
)

OUTPUT = (
    ROOT
    / "exports"
    / "fybroc_seal_pricing_mapping_inventory.json"
)


def clean(value):
    if value is None:
        return None

    if isinstance(
        value,
        (
            str,
            int,
            float,
            bool,
        ),
    ):
        return value

    return str(value)


def cell_record(
    formula_ws,
    value_ws,
    coordinate,
):
    return {
        "cell": coordinate,
        "formula_or_value":
            clean(
                formula_ws[
                    coordinate
                ].value
            ),
        "cached_value":
            clean(
                value_ws[
                    coordinate
                ].value
            ),
    }


def resolve_direct_tables_reference(
    formula,
):
    if not isinstance(formula, str):
        return None

    match = re.fullmatch(
        r"=Tables!\$?([A-Z]+)\$?(\d+)",
        formula.strip(),
        flags=re.IGNORECASE,
    )

    if not match:
        return None

    return (
        match.group(1).upper()
        + match.group(2)
    )


def main():
    price_formula = load_workbook(
        PRICE_WORKBOOK,
        data_only=False,
        keep_vba=True,
        read_only=False,
    )

    price_value = load_workbook(
        PRICE_WORKBOOK,
        data_only=True,
        keep_vba=True,
        read_only=False,
    )

    nom_formula = load_workbook(
        NOMENCLATURE_WORKBOOK,
        data_only=False,
        keep_vba=True,
        read_only=False,
    )

    nom_value = load_workbook(
        NOMENCLATURE_WORKBOOK,
        data_only=True,
        keep_vba=True,
        read_only=False,
    )

    tables_f = price_formula["Tables"]
    tables_v = price_value["Tables"]

    pricebook_f = price_formula["Pricebook"]
    pricebook_v = price_value["Pricebook"]

    seal_f = nom_formula["Seal Assembly"]
    seal_v = nom_value["Seal Assembly"]

    payload = {
        "legacy_seal_types": [],
        "legacy_elastomers": [],
        "table79": [],
        "table137": [],
        "runtime_vocabulary": [],
        "runtime_combinations": [],
    }


    # ------------------------------------------------------------
    # Legacy Price Check seal-type vocabulary
    # Tables!B95:B108
    # ------------------------------------------------------------

    print("=" * 100)
    print(
        "LEGACY PRICE CHECK SEAL TYPE VOCABULARY"
    )
    print("=" * 100)

    for row in range(95, 109):
        coordinate = f"B{row}"

        record = cell_record(
            tables_f,
            tables_v,
            coordinate,
        )

        payload[
            "legacy_seal_types"
        ].append(record)

        print(
            f"Tables!{coordinate:<5} "
            f"FORMULA={record['formula_or_value']!r} "
            f"CACHED={record['cached_value']!r}"
        )


    # ------------------------------------------------------------
    # Legacy seal-elastomer vocabulary
    # ------------------------------------------------------------

    print()
    print("=" * 100)
    print(
        "LEGACY PRICE CHECK SEAL ELASTOMERS"
    )
    print("=" * 100)

    for row in range(385, 390):
        coordinate = f"B{row}"

        record = cell_record(
            tables_f,
            tables_v,
            coordinate,
        )

        payload[
            "legacy_elastomers"
        ].append(record)

        print(
            f"Tables!{coordinate:<5} "
            f"FORMULA={record['formula_or_value']!r} "
            f"CACHED={record['cached_value']!r}"
        )


    # ------------------------------------------------------------
    # Table79
    #
    # BR = legacy Seal Type
    # BS = Group 1 list
    # BV = Group 2 list
    # BY = Group 3 list
    # ------------------------------------------------------------

    print()
    print("=" * 100)
    print("TABLE79 PRICING KEYS")
    print("=" * 100)

    for row in range(6, 20):
        seal_type_formula = (
            pricebook_f[
                f"BR{row}"
            ].value
        )

        source_cell = (
            resolve_direct_tables_reference(
                seal_type_formula
            )
        )

        source_record = None

        if source_cell:
            source_record = cell_record(
                tables_f,
                tables_v,
                source_cell,
            )

        record = {
            "row": row,
            "seal_type_formula":
                clean(
                    seal_type_formula
                ),
            "seal_type_cached":
                clean(
                    pricebook_v[
                        f"BR{row}"
                    ].value
                ),
            "source_cell":
                source_cell,
            "source_value":
                (
                    source_record
                    if source_record
                    else None
                ),
            "group_1": clean(
                pricebook_f[
                    f"BS{row}"
                ].value
            ),
            "group_2": clean(
                pricebook_f[
                    f"BV{row}"
                ].value
            ),
            "group_3": clean(
                pricebook_f[
                    f"BY{row}"
                ].value
            ),
        }

        payload[
            "table79"
        ].append(record)

        print(
            f"ROW {row:<3} "
            f"SOURCE={source_cell!r:<8} "
            f"TYPE="
            f"{record['seal_type_cached']!r} "
            f"G1={record['group_1']!r} "
            f"G2={record['group_2']!r} "
            f"G3={record['group_3']!r}"
        )


    # ------------------------------------------------------------
    # Table137
    # ------------------------------------------------------------

    print()
    print("=" * 100)
    print("TABLE137 TEFLON / ELASTOMER PRICING")
    print("=" * 100)

    for row in range(24, 47):
        record = {
            "row": row,
            "seal_type_formula":
                clean(
                    pricebook_f[
                        f"BR{row}"
                    ].value
                ),
            "seal_type_cached":
                clean(
                    pricebook_v[
                        f"BR{row}"
                    ].value
                ),
            "group_1": clean(
                pricebook_f[
                    f"BS{row}"
                ].value
            ),
            "group_2": clean(
                pricebook_f[
                    f"BV{row}"
                ].value
            ),
            "group_3": clean(
                pricebook_f[
                    f"BY{row}"
                ].value
            ),
        }

        payload[
            "table137"
        ].append(record)

        if any(
            record[key] is not None
            for key in (
                "seal_type_formula",
                "seal_type_cached",
                "group_1",
                "group_2",
                "group_3",
            )
        ):
            print(
                f"ROW {row:<3} "
                f"TYPE_FORMULA="
                f"{record['seal_type_formula']!r} "
                f"TYPE_CACHED="
                f"{record['seal_type_cached']!r} "
                f"G1={record['group_1']!r} "
                f"G2={record['group_2']!r} "
                f"G3={record['group_3']!r}"
            )


    # ------------------------------------------------------------
    # Runtime seal vocabulary
    # ------------------------------------------------------------

    print()
    print("=" * 100)
    print("RUNTIME SEAL VOCABULARY")
    print("=" * 100)

    field_columns = {
        "SEAL_OPTION": "D",
        "SEAL_TYPE": "E",
        "SEAL_MATERIALS": "F",
        "SEAL_ELASTOMERS": "G",
    }

    for field_code, column in (
        field_columns.items()
    ):
        values = []

        for row in range(3, 9):
            value = clean(
                seal_f[
                    f"{column}{row}"
                ].value
            )

            if value is not None:
                values.append(value)

        record = {
            "field_code": field_code,
            "values": values,
        }

        payload[
            "runtime_vocabulary"
        ].append(record)

        print(
            f"{field_code:<20}: "
            f"{values}"
        )


    # ------------------------------------------------------------
    # Actual runtime Seal Assembly combinations
    # Rows 14:80
    # ------------------------------------------------------------

    print()
    print("=" * 100)
    print("RUNTIME SEAL ASSEMBLY COMBINATIONS")
    print("=" * 100)

    for row in range(14, 81):
        values = {
            field_code:
                clean(
                    seal_f[
                        f"{column}{row}"
                    ].value
                )
            for (
                field_code,
                column
            ) in field_columns.items()
        }

        if not any(
            value is not None
            for value in values.values()
        ):
            continue

        record = {
            "row": row,
            **values,
        }

        payload[
            "runtime_combinations"
        ].append(record)

        print(
            f"{row:>3}: "
            f"OPTION={values['SEAL_OPTION']!r} | "
            f"TYPE={values['SEAL_TYPE']!r} | "
            f"MATERIALS={values['SEAL_MATERIALS']!r} | "
            f"ELASTOMERS={values['SEAL_ELASTOMERS']!r}"
        )


    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT.write_text(
        json.dumps(
            payload,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 100)
    print(
        "Output:",
        OUTPUT,
    )
    print("=" * 100)

    price_formula.close()
    price_value.close()
    nom_formula.close()
    nom_value.close()


if __name__ == "__main__":
    main()
