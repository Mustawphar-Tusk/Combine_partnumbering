"""F110.2b - Fybroc V5/V6 Value-Domain Diff (Seal Assembly / Options / Motor Assy).

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F110
("V5 to V6 Nomenclature Reconciliation").

F110.2a already diffed `Attributes` - a clean name-to-hex-code lookup
table. These three sheets are structured differently: each column is a
field, each row down that column is one valid option value for that
field. There is no per-value hex code to compare here - Smart Number's
own live example shows the whole Seal Assembly / Motor Assy combination
gets ONE combined code (e.g. "03", "049"), not one code per field per
value. So this step compares value DOMAINS (which options exist), not
hex codes.

Field mapping was verified directly against both real workbooks (row
extents + field names) before writing this, not assumed:

  Seal Assembly -> Seal Assembly - Horizontal: 5 shared fields, same
      78-value extent in both. V6 adds a 6th field, Seal Manufacturer,
      with no V5 equivalent.
  Options -> Options - Horizontal: 5 shared fields, V6 nearly 2x the
      row extent (197 vs 105). V6 adds C-Face Adapter.
  Options -> Options - Vertical: only 4 of V5's 5 fields have an
      equivalent. Baseplate_Option is RENAMED to "Mounting Plate
      Option" (not a simple wording tweak - flagged explicitly, not
      auto-matched). Baseplate_Hardware has NO Vertical equivalent at
      all - flagged as dropped, not silently ignored.
  Motor Assy -> Motor Assy: all 11 fields match 1:1, same 197-value
      extent in both.

Values are matched using the same normalization already established in
F110.2a (strip a trailing "*" default-value marker, "_" <-> " "), so a
cosmetic V6 cleanup doesn't show up as a false DEPRECATED+NEW pair.

No workbook is opened in write mode. Nothing is written back.

Inputs:
  workbooks/Fybroc/Fybroc Nomenclature_V5.xlsm  (read-only)
  workbooks/Fybroc/Nomenclature_V6.xlsm         (read-only)

Outputs:
  docs/evidence/F110/FYBROC_VALUE_DOMAIN_DIFF.{json,txt}
"""

from __future__ import annotations

import argparse
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required (it's already a pinned project dependency - "
        "run this with the project's own .venv interpreter)."
    ) from exc

STEP = "F110.2b"
ROADMAP_VERSION = "1.0"
MILESTONE = "F110"

V5_WORKBOOK_REL = "workbooks/Fybroc/Fybroc Nomenclature_V5.xlsm"
V6_WORKBOOK_REL = "workbooks/Fybroc/Nomenclature_V6.xlsm"

HEADER_ROW = 2
DATA_START_ROW = 3
MAX_SCAN_ROW = 250  # comfortably above the largest confirmed extent (199)

# ---------------------------------------------------------------------------
# Verified (not assumed) field mappings: (v5_sheet, v6_sheet) -> list of
# (v5_field_or_None, v6_field_or_None). None on either side means "no
# equivalent on that side" - reported explicitly, not silently skipped.
# ---------------------------------------------------------------------------
SHEET_PAIRS: list[dict[str, Any]] = [
    {
        "label": "Seal Assembly",
        "v5_sheet": "Seal Assembly",
        "v6_sheet": "Seal Assembly - Horizontal",
        "fields": [
            ("Seal_Option", "Seal Option"),
            ("Seal_Type", "Seal Type"),
            ("Seal_Materials", "Seal Materials"),
            ("Seal_Elastomers", "Seal Elastomers"),
            ("Seal_Guard", "Seal Guard"),
            (None, "Seal Manufacturer"),
        ],
    },
    {
        "label": "Options (vs Horizontal)",
        "v5_sheet": "Options",
        "v6_sheet": "Options - Horizontal",
        "fields": [
            ("Coupling_Option", "Coupling Option"),
            ("Coupling_Guard", "Coupling Guard"),
            ("Baseplate_Option", "Baseplate Option"),
            ("Baseplate_Hardware", "Baseplate Hardware"),
            ("Custom_Nameplate", "Customer Nameplate"),
            (None, "C-Face Adapter"),
        ],
    },
    {
        "label": "Options (vs Vertical)",
        "v5_sheet": "Options",
        "v6_sheet": "Options - Vertical",
        "fields": [
            ("Coupling_Option", "Coupling Option"),
            ("Coupling_Guard", "Coupling Guard"),
            ("Baseplate_Option", "Mounting Plate Option"),  # verified rename, not a guess
            ("Baseplate_Hardware", None),  # verified: no Vertical equivalent exists
            ("Custom_Nameplate", "Customer Nameplate"),
        ],
    },
    {
        "label": "Motor Assy",
        "v5_sheet": " Motor Assy",
        "v6_sheet": " Motor Assy",
        "fields": [
            ("Motor_Option", "Motor Option"),
            ("Motor_Class", "Motor Class"),
            ("Motor_Orientation", "Motor Orientation"),
            ("Motor_Horsepower", "Motor Horsepower"),
            ("Motor_RPM", "Motor RPM"),
            ("Motor_Voltage", "Motor Voltage"),
            ("Motor_Hertz", "Motor Hertz"),
            ("Motor_Frame", "Motor Frame"),
            ("Motor_Enclosure", "Motor Enclosure"),
            ("Motor_Efficiency", "Motor Efficiency"),
            ("Motor_Manufacturer", "Motor Manufacturer"),
        ],
    },
]


def normalize_value(value: Any) -> str:
    """Same normalization as F110.2a - reused, not re-derived."""
    text = str(value).strip()
    if text.endswith("*"):
        text = text[:-1].strip()
    return text.replace("_", " ").strip().lower()


def find_column(ws, field_name: str) -> int | None:
    for col in range(1, 30):
        if ws.cell(row=HEADER_ROW, column=col).value == field_name:
            return col
    return None


def read_value_list(ws, field_name: str) -> dict[str, str]:
    """normalized_value -> raw_value, for one field's column."""
    col = find_column(ws, field_name)
    if col is None:
        return {}
    result: dict[str, str] = {}
    for row in range(DATA_START_ROW, MAX_SCAN_ROW):
        v = ws.cell(row=row, column=col).value
        if v is None or str(v).strip() == "":
            continue
        result[normalize_value(v)] = str(v)
    return result


def git_info(repo_root: Path) -> tuple[str, bool]:
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout.strip()
        status = subprocess.run(
            ["git", "status", "--porcelain"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout
        return commit, (status.strip() == "")
    except Exception:
        return "unknown", False


def diff_pair(repo_root: Path, pair: dict[str, Any]) -> dict[str, Any]:
    v5_path = repo_root / V5_WORKBOOK_REL
    v6_path = repo_root / V6_WORKBOOK_REL
    wb5 = openpyxl.load_workbook(str(v5_path), read_only=True, data_only=True)
    wb6 = openpyxl.load_workbook(str(v6_path), read_only=True, data_only=True)
    ws5 = wb5[pair["v5_sheet"]]
    ws6 = wb6[pair["v6_sheet"]]

    field_results = []
    for v5_field, v6_field in pair["fields"]:
        if v5_field is None:
            v6_values = read_value_list(ws6, v6_field)
            field_results.append({
                "v5_field": None, "v6_field": v6_field,
                "note": "No V5 equivalent - new field in V6.",
                "value_count": len(v6_values),
                "classification_counts": {"NEW": len(v6_values)},
                "values": [{"value": v, "classification": "NEW"} for v in v6_values.values()],
            })
            continue
        if v6_field is None:
            v5_values = read_value_list(ws5, v5_field)
            field_results.append({
                "v5_field": v5_field, "v6_field": None,
                "note": "No V6 equivalent found - verified absent, not just unmatched by name.",
                "value_count": len(v5_values),
                "classification_counts": {"DEPRECATED": len(v5_values)},
                "values": [{"value": v, "classification": "DEPRECATED"} for v in v5_values.values()],
            })
            continue

        v5_values = read_value_list(ws5, v5_field)
        v6_values = read_value_list(ws6, v6_field)
        all_keys = sorted(set(v5_values) | set(v6_values))
        counts = {"UNCHANGED": 0, "NEW": 0, "DEPRECATED": 0}
        rows = []
        for key in all_keys:
            in_v5, in_v6 = key in v5_values, key in v6_values
            if in_v5 and in_v6:
                cls = "UNCHANGED"
            elif in_v6:
                cls = "NEW"
            else:
                cls = "DEPRECATED"
            counts[cls] += 1
            rows.append({
                "value": v5_values.get(key) or v6_values.get(key),
                "v5_raw": v5_values.get(key), "v6_raw": v6_values.get(key),
                "classification": cls,
            })
        field_results.append({
            "v5_field": v5_field, "v6_field": v6_field, "note": None,
            "value_count": len(all_keys),
            "classification_counts": counts,
            "values": rows,
        })

    wb5.close()
    wb6.close()
    return {
        "label": pair["label"], "v5_sheet": pair["v5_sheet"], "v6_sheet": pair["v6_sheet"],
        "fields": field_results,
    }


def build_diff(repo_root: Path) -> dict[str, Any]:
    commit, clean = git_info(repo_root)
    pairs = [diff_pair(repo_root, p) for p in SHEET_PAIRS]
    total: dict[str, int] = {}
    for p in pairs:
        for f in p["fields"]:
            for cls, n in f["classification_counts"].items():
                total[cls] = total.get(cls, 0) + n
    return {
        "step": STEP, "roadmap_version": ROADMAP_VERSION, "milestone": MILESTONE,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit, "git_working_tree_clean": clean,
        "total_classification_counts": total,
        "sheet_pairs": pairs,
    }


def _banner(title: str) -> str:
    line = "=" * 140
    return f"{line}\r\n{title}\r\n{line}\r\n\r\n"


def write_outputs(evidence_dir: Path, result: dict[str, Any]) -> None:
    payload = {"artifact": "FYBROC_VALUE_DOMAIN_DIFF", **result}
    (evidence_dir / "FYBROC_VALUE_DOMAIN_DIFF.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
    )

    lines = [_banner("F110.2b - FYBROC VALUE-DOMAIN DIFF (Seal Assembly / Options / Motor Assy)")]
    lines.append(
        f"Git commit  : {result['git_commit']}\r\nGit clean   : {result['git_working_tree_clean']}\r\n\r\n"
    )
    lines.append("TOTAL CLASSIFICATION COUNTS\r\n" + "-" * 140 + "\r\n")
    for cls, count in sorted(result["total_classification_counts"].items()):
        lines.append(f"  {cls:<15}: {count}\r\n")
    lines.append("\r\n")

    for p in result["sheet_pairs"]:
        lines.append(_banner(f"{p['label']}  ::  {p['v5_sheet']!r} vs {p['v6_sheet']!r}"))
        for f in p["fields"]:
            v5_label = f["v5_field"] or "(none)"
            v6_label = f["v6_field"] or "(none)"
            lines.append(f"[{v5_label}] <-> [{v6_label}]  ({f['value_count']} values)\r\n")
            if f["note"]:
                lines.append(f"    note: {f['note']}\r\n")
            counts_str = "  ".join(f"{k}={v}" for k, v in f["classification_counts"].items() if v)
            lines.append(f"    {counts_str}\r\n")
            for r in f["values"]:
                if r["classification"] == "UNCHANGED":
                    continue
                lines.append(f"      {r['classification']:<12} {r['value']!r}\r\n")
            lines.append("\r\n")

    (evidence_dir / "FYBROC_VALUE_DOMAIN_DIFF.txt").write_text("".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="F110.2b Fybroc value-domain diff")
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument("--repo-root", type=Path, default=default_root)
    parser.add_argument("--evidence-dir", type=Path, default=None)
    args = parser.parse_args()

    repo_root: Path = args.repo_root.resolve()
    evidence_dir: Path = (args.evidence_dir or (repo_root / "docs" / "evidence" / "F110")).resolve()
    evidence_dir.mkdir(parents=True, exist_ok=True)

    result = build_diff(repo_root)
    write_outputs(evidence_dir, result)

    print(json.dumps({
        "step": STEP, "output_dir": str(evidence_dir),
        "total_classification_counts": result["total_classification_counts"],
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())