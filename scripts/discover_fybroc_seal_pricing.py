from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

WORKBOOK = (
    ROOT
    / "workbooks"
    / "Fybroc"
    / "Price Estimator-Fybroc.xlsm"
)

TABLE_NAMES = (
    "Table4",
    "Table79",
    "Table137",
)

LOOKUP_CELLS = (
    "B102",
    "B104",
    "B108",
    "B385",
    "B386",
    "B388",
    "B389",
)


def find_table(workbook, table_name):
    for worksheet in workbook.worksheets:
        if table_name in worksheet.tables:
            return (
                worksheet,
                worksheet.tables[table_name],
            )

    raise KeyError(
        f"Table not found: {table_name}"
    )


def print_table(
    worksheet,
    table,
) -> None:

    print("=" * 100)
    print(
        f"{table.name} "
        f"[{worksheet.title}!{table.ref}]"
    )
    print("=" * 100)

    cells = worksheet[
        table.ref
    ]

    for row_number, row in enumerate(
        cells,
        start=1,
    ):
        values = [
            cell.value
            for cell in row
        ]

        print(
            f"{row_number:>3}: "
            + " | ".join(
                repr(value)
                for value in values
            )
        )

    print()


def main() -> None:

    workbook = load_workbook(
        WORKBOOK,
        data_only=False,
        keep_vba=True,
        read_only=False,
    )

    print("=" * 100)
    print(
        "M022.2 FYBROC SEAL PRICING DISCOVERY"
    )
    print("=" * 100)
    print(
        "Workbook:",
        WORKBOOK,
    )
    print()

    for table_name in TABLE_NAMES:
        worksheet, table = find_table(
            workbook,
            table_name,
        )

        print_table(
            worksheet,
            table,
        )

    print("=" * 100)
    print(
        "TABLES SHEET LOOKUP VOCABULARY"
    )
    print("=" * 100)

    worksheet = workbook[
        "Tables"
    ]

    for cell in LOOKUP_CELLS:
        print(
            f"Tables!{cell} = "
            f"{worksheet[cell].value!r}"
        )

    print()

    price_check = workbook[
        "Price Check"
    ]

    print("=" * 100)
    print(
        "PRICE CHECK SEAL INPUTS"
    )
    print("=" * 100)

    for cell in (
        "B25",
        "B26",
        "B27",
        "B32",
        "B33",
        "B54",
        "D81",
        "J81",
    ):
        print(
            f"Price Check!{cell} = "
            f"{price_check[cell].value!r}"
        )

    workbook.close()


if __name__ == "__main__":
    main()
