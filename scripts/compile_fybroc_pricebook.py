"""F130.1 - Fybroc Price Estimator Pricebook Compiler.

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F130
("Fybroc Pricing & Adders Reconciliation") - existing source review.

The Pricebook sheet in Price Estimator-Fybroc.xlsm is the authoritative
production base-price source. It contains 15+ side-by-side pricing
tables for different series/components:

  1500 Pump Pricing          (table 81)     cols 2-16   25 sizes x 10 materials
  1530 Pump Pricing          (table49)      cols 17-26  25 sizes x 6 materials
  1600 Pump Pricing          (table50)      cols 28-38  sizes x materials
  1630 Pump Pricing          (table91)      cols 40-48  sizes x materials
  2530 Pump Pricing          (table92)      cols 50-58  sizes x materials
  2630 Pump Pricing          (table93)      cols 60-68  sizes x materials
  Seal Pricing               (Table79)      cols 70-79  seal types x groups
  Coupling Pricing           (table 83)     cols 82-85  frames x groups
  5500 Pump Pricing          (table96)      cols 87-115 sizes x settings x materials
  5500/7500 Tailpipe         (---)          cols 116-124
  5530 Pricing               (Table128)     cols 126-130
  7500 Pump Pricing          (table143)     cols 133-155
  7530 Pump Pricing          (Table144)     cols 157-167
  3000 Pump Pricing          (Table152)     cols 169-176
  500 Pump Pricing           (---)          cols 178+

Each horizontal-series block has: Group, Size, ANSI Desig, ANSI Size,
then one price column per material (VR-1, VR-1A, EY-2, BPO/DMA, etc.)

5500 is keyed on Size + Setting Number (vertical pumps have settings).

Inputs:
  workbooks/Fybroc/Price Estimator-Fybroc.xlsm  (read-only)

Outputs:
  docs/evidence/F130/FYBROC_PRICEBOOK.{json,txt}
"""
from __future__ import annotations
import argparse, json, subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl required") from exc

STEP = "F130.1"; ROADMAP_VERSION = "1.1"; MILESTONE = "F130"
WORKBOOK_REL = "workbooks/Fybroc/Price Estimator-Fybroc.xlsm"
SHEET = "Pricebook"
HEADER_ROW = 5
DATA_START = 6

# Series table blocks identified from workbook scan
# (start_col, table_id, title, end_col)
SERIES_BLOCKS = [
    (2,  "table81",  "1500 Pump Pricing",  16),
    (17, "table49",  "1530 Pump Pricing",  26),
    (28, "table50",  "1600 Pump Pricing",  38),
    (40, "table91",  "1630 Pump Pricing",  48),
    (50, "table92",  "2530 Pump Pricing",  58),
    (60, "table93",  "2630 Pump Pricing",  68),
    (70, "Table79",  "Seal Pricing (1500/1530/1600/1630/3000)", 79),
    (82, "table83",  "Coupling Pricing (1500/1600)", 85),
    (87, "table96",  "5500 Pump Pricing",  115),
    (116, None,      "5500/7500 Tailpipe Pricing", 124),
    (126, "Table128","5530 Pricing",       130),
    (133, "table143","7500 Pump Pricing",  155),
    (157, "Table144","7530 Pump Pricing",  167),
    (169, "Table152","3000 Pump Pricing",  176),
    (178, None,      "500 Pump Pricing",   188),
]


def git_info(repo_root):
    try:
        c = subprocess.run(["git","rev-parse","HEAD"], cwd=repo_root, capture_output=True, text=True, check=True).stdout.strip()
        s = subprocess.run(["git","status","--porcelain"], cwd=repo_root, capture_output=True, text=True, check=True).stdout
        return c, (s.strip() == "")
    except: return "unknown", False

def _s(v):
    if v is None: return None
    s = str(v).strip(); return s if s else None

def _num(v):
    if v is None: return None
    try: return float(v)
    except: return None


def compile_block(ws, start_col, end_col, max_row) -> dict[str, Any]:
    """Compile one pricing table block."""
    # Read headers from row 5
    headers = []
    for c in range(start_col, end_col + 1):
        h = _s(ws.cell(row=HEADER_ROW, column=c).value)
        headers.append((c, h))

    # Filter to non-None headers
    active_headers = [(c, h) for c, h in headers if h]
    if not active_headers:
        return {"headers": [], "data_rows": 0, "sample_rows": []}

    # Identify material/price columns (columns after the size/group/desig cols)
    # These are the ones that aren't 'Group', 'SIZE', 'ANSI DESIG.', 'ANSI Size'
    meta_names = {'Group', 'SIZE', 'ANSI DESIG.', 'ANSI Size', 'Setting Number',
                  'ANSI Size : Setting Number', 'Frame', 'Seal Type', 'Suction Pipe Size',
                  'Dia. At Sleeve Bearings'}
    price_cols = [(c, h) for c, h in active_headers if h not in meta_names]
    meta_cols = [(c, h) for c, h in active_headers if h in meta_names]

    # Count data rows
    data_rows = 0
    first_data_col = active_headers[0][0]
    for r in range(DATA_START, max_row + 1):
        v = ws.cell(row=r, column=first_data_col).value
        if v is None:
            break
        data_rows += 1

    # Sample first 5 data rows
    samples = []
    for r in range(DATA_START, min(DATA_START + 5, DATA_START + data_rows)):
        row_data = {}
        for c, h in active_headers:
            v = ws.cell(row=r, column=c).value
            row_data[h] = v
        samples.append(row_data)

    return {
        "headers": [h for _, h in active_headers],
        "meta_columns": [h for _, h in meta_cols],
        "price_columns": [h for _, h in price_cols],
        "price_column_count": len(price_cols),
        "data_rows": data_rows,
        "sample_rows": samples,
    }


def build_model(repo_root):
    commit, clean = git_info(repo_root)
    wb = openpyxl.load_workbook(str(repo_root / WORKBOOK_REL), read_only=True, data_only=True)
    ws = wb[SHEET]

    blocks = []
    for start_col, table_id, title, end_col in SERIES_BLOCKS:
        block_data = compile_block(ws, start_col, end_col, ws.max_row)
        blocks.append({
            "start_col": start_col,
            "end_col": end_col,
            "table_id": table_id,
            "title": title,
            **block_data,
        })

    wb.close()

    total_price_cols = sum(b["price_column_count"] for b in blocks)

    return {
        "step": STEP, "roadmap_version": ROADMAP_VERSION, "milestone": MILESTONE,
        "source_workbook": WORKBOOK_REL, "sheet": SHEET,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit, "git_working_tree_clean": clean,
        "table_block_count": len(blocks),
        "total_price_columns": total_price_cols,
        "blocks": blocks,
    }


def _banner(title):
    return f"{'='*120}\r\n{title}\r\n{'='*120}\r\n\r\n"

def write_outputs(evidence_dir, result):
    payload = {"artifact": "FYBROC_PRICEBOOK", **result}
    (evidence_dir / "FYBROC_PRICEBOOK.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    out = [_banner("F130.1 - FYBROC PRICEBOOK (Base Pump Pricing from Price Estimator)")]
    out.append(f"Git commit       : {result['git_commit']}\r\nGit clean        : {result['git_working_tree_clean']}\r\n\r\n")
    out.append(f"Table blocks     : {result['table_block_count']}\r\n")
    out.append(f"Total price cols : {result['total_price_columns']}\r\n\r\n")

    for b in result["blocks"]:
        out.append(_banner(f"{b['title']} ({b['table_id'] or 'no table id'})  cols {b['start_col']}-{b['end_col']}"))
        out.append(f"  Data rows        : {b['data_rows']}\r\n")
        out.append(f"  Meta columns     : {b['meta_columns']}\r\n")
        out.append(f"  Price columns    : {b['price_columns']}\r\n\r\n")
        if b["sample_rows"]:
            out.append("  SAMPLE (first 5):\r\n")
            for sr in b["sample_rows"]:
                out.append(f"    {sr}\r\n")
            out.append("\r\n")

    (evidence_dir / "FYBROC_PRICEBOOK.txt").write_text("".join(out), encoding="utf-8")


def main():
    p = argparse.ArgumentParser(); root = Path(__file__).resolve().parent.parent
    p.add_argument("--repo-root", type=Path, default=root)
    p.add_argument("--evidence-dir", type=Path, default=None)
    a = p.parse_args(); repo_root = a.repo_root.resolve()
    ev = (a.evidence_dir or (repo_root / "docs" / "evidence" / "F130")).resolve()
    ev.mkdir(parents=True, exist_ok=True)
    result = build_model(repo_root)
    write_outputs(ev, result)
    print(json.dumps({
        "step": STEP, "output_dir": str(ev),
        "table_blocks": result["table_block_count"],
        "total_price_columns": result["total_price_columns"],
        "block_summary": [
            {"title": b["title"], "data_rows": b["data_rows"], "price_cols": b["price_column_count"]}
            for b in result["blocks"]
        ],
    }, indent=2))
    return 0

if __name__ == "__main__": raise SystemExit(main())
