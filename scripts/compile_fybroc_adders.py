"""F130.2 - Fybroc Price Estimator Adders Compiler.

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F130.

The Adders sheet in Price Estimator-Fybroc.xlsm contains all engineering
adder pricing organized by category. Categories are identified by rows
where col C = 'Group-1' (section headers).

Major adder categories:
  - Hardware (Gland, Frame, Baseplate, Casing) by material x group
  - Shaft (5500, 7500) by material x group
  - Flush (1500/1530/1600/1630/3000) by type x group
  - Bearings by type x group
  - Miscellaneous (Gauge Taps, Casing Drains, Cyclone Sep, DIN/JIS, etc.)
  - Elastomers (standard + self-priming) by material x group
  - Journal Sleeve by material x group
  - Alloy Shaft Sleeve / Packing

Each adder has: description, Group-1/2/3 prices, and sometimes a
Horizontal/Vertical/Accessory breakdown.

Inputs:
  workbooks/Fybroc/Price Estimator-Fybroc.xlsm  (read-only)

Outputs:
  docs/evidence/F130/FYBROC_ADDERS.{json,txt}
"""
from __future__ import annotations
import argparse, json, subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl required") from exc

STEP = "F130.2"; ROADMAP_VERSION = "1.1"; MILESTONE = "F130"
WORKBOOK_REL = "workbooks/Fybroc/Price Estimator-Fybroc.xlsm"
SHEET = "Adders"


def git_info(repo_root):
    try:
        c = subprocess.run(["git","rev-parse","HEAD"], cwd=repo_root, capture_output=True, text=True, check=True).stdout.strip()
        s = subprocess.run(["git","status","--porcelain"], cwd=repo_root, capture_output=True, text=True, check=True).stdout
        return c, (s.strip() == "")
    except: return "unknown", False

def _s(v):
    if v is None: return None
    s = str(v).strip(); return s if s else None

def _num(v):
    if v is None: return None
    try: return float(v)
    except: return _s(v)


def build_model(repo_root):
    commit, clean = git_info(repo_root)
    wb = openpyxl.load_workbook(str(repo_root / WORKBOOK_REL), read_only=True, data_only=True)
    ws = wb[SHEET]

    # Identify section boundaries by looking for "Group-1" in col 3 (first 320 rows only - primary data)
    section_starts = []
    for r in range(1, 320):
        c3 = _s(ws.cell(row=r, column=3).value)
        if c3 and "Group" in c3:
            title = _s(ws.cell(row=r, column=2).value) or ""
            section_starts.append((r, title))

    # Also capture MISC section (identified at row 25)
    c2_25 = _s(ws.cell(row=25, column=2).value)
    if c2_25 and "MISC" in c2_25.upper():
        section_starts.append((25, "MISC."))

    section_starts.sort(key=lambda x: x[0])

    # For each section, read adder rows until next section or blank gap
    sections = []
    for i, (start_row, title) in enumerate(section_starts):
        end_row = section_starts[i+1][0] - 1 if i+1 < len(section_starts) else ws.max_row

        adders = []
        for r in range(start_row + 1, min(end_row + 1, start_row + 100)):
            desc = _s(ws.cell(row=r, column=2).value)
            if not desc:
                # Allow one blank row gap, stop at two consecutive blanks
                next_desc = _s(ws.cell(row=r+1, column=2).value) if r+1 <= end_row else None
                if not next_desc:
                    break
                continue
            adder = {
                "description": desc,
                "id": _num(ws.cell(row=r, column=1).value),
                "group_1": _num(ws.cell(row=r, column=3).value),
                "group_2": _num(ws.cell(row=r, column=4).value),
                "group_3": _num(ws.cell(row=r, column=5).value),
                "extra": _num(ws.cell(row=r, column=6).value),
            }
            v7 = _num(ws.cell(row=r, column=7).value)
            if v7 is not None:
                adder["col7"] = v7
            adders.append(adder)

        sections.append({
            "start_row": start_row,
            "category": title,
            "adder_count": len(adders),
            "adders": adders,
        })

    wb.close()

    total_adders = sum(s["adder_count"] for s in sections)

    return {
        "step": STEP, "roadmap_version": ROADMAP_VERSION, "milestone": MILESTONE,
        "source_workbook": WORKBOOK_REL, "sheet": SHEET,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit, "git_working_tree_clean": clean,
        "section_count": len(sections),
        "total_adder_lines": total_adders,
        "sections": sections,
    }


def _banner(title):
    return f"{'='*120}\r\n{title}\r\n{'='*120}\r\n\r\n"

def write_outputs(evidence_dir, result):
    payload = {"artifact": "FYBROC_ADDERS", **result}
    (evidence_dir / "FYBROC_ADDERS.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    out = [_banner("F130.2 - FYBROC ADDERS (from Price Estimator)")]
    out.append(f"Git commit       : {result['git_commit']}\r\nGit clean        : {result['git_working_tree_clean']}\r\n\r\n")
    out.append(f"Sections         : {result['section_count']}\r\n")
    out.append(f"Total adder lines: {result['total_adder_lines']}\r\n\r\n")

    for s in result["sections"]:
        out.append(_banner(f"{s['category']} (row {s['start_row']}, {s['adder_count']} adders)"))
        for a in s["adders"]:
            g1 = a.get("group_1", "—")
            g2 = a.get("group_2", "—")
            g3 = a.get("group_3", "—")
            out.append(f"  {a['description']:<55} G1={str(g1):<8} G2={str(g2):<8} G3={str(g3):<8}\r\n")
        out.append("\r\n")

    (evidence_dir / "FYBROC_ADDERS.txt").write_text("".join(out), encoding="utf-8")


def main():
    p = argparse.ArgumentParser(); root = Path(__file__).resolve().parent.parent
    p.add_argument("--repo-root", type=Path, default=root)
    p.add_argument("--evidence-dir", type=Path, default=None)
    a = p.parse_args(); repo_root = a.repo_root.resolve()
    ev = (a.evidence_dir or (repo_root / "docs" / "evidence" / "F130")).resolve()
    ev.mkdir(parents=True, exist_ok=True)
    result = build_model(repo_root)
    write_outputs(ev, result)
    print(json.dumps({
        "step": STEP, "output_dir": str(ev),
        "section_count": result["section_count"],
        "total_adder_lines": result["total_adder_lines"],
        "sections": [{"category": s["category"], "adders": s["adder_count"]} for s in result["sections"]],
    }, indent=2))
    return 0

if __name__ == "__main__": raise SystemExit(main())
