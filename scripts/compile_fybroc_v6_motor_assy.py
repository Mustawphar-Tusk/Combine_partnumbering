"""F110.9 - Fybroc Nomenclature V6 Motor Assembly Compiler.

Sheet: ' Motor Assy' (note leading space). This sheet has FIVE regions, all of
which are authoritative and are now captured (a prior version read only rows
3-8 of the option band and dropped the '-' value, truncating most domains):

  R1  D2:Q21   - selectable option DOMAINS, one column per attribute. Columns
                 are SPARSE (values may skip rows, e.g. Orientation resumes at
                 row 14) so each column is read across rows 3-21 collecting all
                 non-empty values. '-' IS a legitimate selectable value and is
                 KEPT (it ties into the combination table below).
  R2  U2:V17   - Motor Hertz -> allowable Motor Voltage (dependency/constraint).
  R3  U20:V34  - Motor Hertz -> allowable Motor RPM (dependency/constraint).
  R4  Z2:AD52  - Motor Frame + Horsepower + RPM -> ID -> 2-char base-36 hex
                 (50 rows). Formula: AD = BASE(AC,36) zero-padded to 2.
  R5  C26:P728 - full motor-assembly combination -> 3-char base-36 hex (702
                 rows). Key C = D&E&F&G&H&I&J&K&L&M&N (11 attrs concatenated);
                 hex O = BASE(P,36) zero-padded to 3. This is the motor-assembly
                 segment lookup for the part number.

Defined names on the sheet (Voltage=I3:I5, RPM=H3:H6, etc.) are NARROWER than
the true R1 domains and must NOT be used as the option source - they omit many
values and the '-'. R1 (rows 3-21) is authoritative for the domains.

Inputs:  workbooks/Fybroc/Nomenclature_V6.xlsm
Outputs: docs/evidence/F110/FYBROC_V6_MOTOR_ASSY.{json,txt}
"""
from __future__ import annotations
import argparse, json, subprocess, re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl required") from exc

STEP="F110.9"; ROADMAP_VERSION="1.0"; MILESTONE="F110"
WORKBOOK_REL="workbooks/Fybroc/Nomenclature_V6.xlsm"
SHEET=" Motor Assy"

# --- R1: option domains (columns D..Q, header row 2, values rows 3..21) ---
R1_HEADER_ROW=2; R1_VALUE_START=3; R1_VALUE_END=21
R1_COLS=[
    (4,"Motor Option"),(5,"Motor Class"),(6,"Motor Orientation"),
    (7,"Motor Horsepower"),(8,"Motor RPM"),(9,"Motor Voltage"),
    (10,"Motor Hertz"),(11,"Motor Frame"),(12,"Motor Enclosure"),
    (13,"Motor Efficiency"),(14,"Motor Manufacturer"),(17,"Speed Control"),
]

# --- R2/R3: Hertz -> Voltage / Hertz -> RPM dependency tables (cols U=21,V=22) ---
R2_HEADER_ROW=2;  R2_START=3;  R2_END=17   # Hertz -> Voltage
R3_HEADER_ROW=20; R3_START=21; R3_END=34   # Hertz -> RPM
DEP_KEY_COL=21; DEP_VAL_COL=22

# --- R4: Frame+Hp+RPM -> ID -> hex (cols Z=26..AD=30, rows 3..52) ---
R4_HEADER_ROW=2; R4_START=3; R4_END=52
R4_FRAME=26; R4_HP=27; R4_RPM=28; R4_ID=29; R4_HEX=30

# --- R5: full combination -> hex (key col C=3, fields D..N, hex O=15, id P=16) ---
R5_KEY=3
R5_FIELD_COLS=[
    (4,"Motor Option"),(5,"Motor Class"),(6,"Motor Orientation"),
    (7,"Motor Horsepower"),(8,"Motor RPM"),(9,"Motor Voltage"),
    (10,"Motor Hertz"),(11,"Motor Frame"),(12,"Motor Enclosure"),
    (13,"Motor Efficiency"),(14,"Motor Manufacturer"),
]
R5_HEX=15; R5_ID=16; R5_DATA_START=27; R5_DATA_END=728


def git_info(repo_root):
    try:
        c=subprocess.run(["git","rev-parse","HEAD"],cwd=repo_root,capture_output=True,text=True,check=True).stdout.strip()
        s=subprocess.run(["git","status","--porcelain"],cwd=repo_root,capture_output=True,text=True,check=True).stdout
        return c,(s.strip()=="")
    except Exception:
        return "unknown",False


def _s(v):
    if v is None: return None
    s=str(v).strip(); return s if s else None


def _dedup(seq):
    seen=set(); out=[]
    for x in seq:
        if x not in seen:
            seen.add(x); out.append(x)
    return out


def build_model(repo_root):
    commit,clean=git_info(repo_root)
    wb=openpyxl.load_workbook(str(repo_root/WORKBOOK_REL),read_only=True,data_only=True)
    ws=wb[SHEET]

    # R1: option domains. Read all rows 3-21 per column, KEEP '-', dedup order.
    option_domains={}
    for col,name in R1_COLS:
        vals=[]
        for r in range(R1_VALUE_START,R1_VALUE_END+1):
            v=_s(ws.cell(row=r,column=col).value)
            if v is not None:
                vals.append(v)
        option_domains[name]=_dedup(vals)

    # R2: Hertz -> Voltage (key cell carries forward until next non-blank key)
    def read_dependency(hstart,hend):
        out={}; cur=None
        for r in range(hstart+1,hend+1):
            k=_s(ws.cell(row=r,column=DEP_KEY_COL).value)
            val=_s(ws.cell(row=r,column=DEP_VAL_COL).value)
            if k: cur=k; out.setdefault(cur,[])
            if cur and val: out[cur].append(val)
        return {k:_dedup(v) for k,v in out.items()}

    hertz_to_voltage=read_dependency(R2_HEADER_ROW,R2_END)
    hertz_to_rpm=read_dependency(R3_HEADER_ROW,R3_END)

    # R4: Frame+Hp+RPM -> ID/hex
    frame_hp_rpm_hex=[]
    for r in range(R4_START,R4_END+1):
        frame=_s(ws.cell(row=r,column=R4_FRAME).value)
        hp=ws.cell(row=r,column=R4_HP).value
        rpm=ws.cell(row=r,column=R4_RPM).value
        hx=_s(ws.cell(row=r,column=R4_HEX).value)
        idv=ws.cell(row=r,column=R4_ID).value
        if frame is None and hx is None: continue
        frame_hp_rpm_hex.append({
            "motor_frame":frame,"motor_horsepower":hp,"motor_rpm":rpm,
            "id":idv,"hex":hx,
        })

    # R5: full combination -> hex (the motor-assembly segment lookup)
    combos=[]; hexset=set(); fv={name:set() for _,name in R5_FIELD_COLS}
    total=0
    for r in range(R5_DATA_START,R5_DATA_END+1):
        hx=_s(ws.cell(row=r,column=R5_HEX).value)
        if hx is None: break
        total+=1; hexset.add(hx)
        rd={}
        for col,name in R5_FIELD_COLS:
            v=_s(ws.cell(row=r,column=col).value); rd[name]=v
            if v: fv[name].add(v)
        combos.append({
            "id":ws.cell(row=r,column=R5_ID).value,"hex_code":hx,
            "key":_s(ws.cell(row=r,column=R5_KEY).value),**rd,
        })

    wb.close()
    sh=sorted(hexset)
    return {
        "step":STEP,"roadmap_version":ROADMAP_VERSION,"milestone":MILESTONE,
        "source_workbook":WORKBOOK_REL,"sheet":SHEET,
        "generated_utc":datetime.now(timezone.utc).isoformat(),
        "git_commit":commit,"git_working_tree_clean":clean,
        # R1
        "option_domains":option_domains,
        "note_dash":"'-' is a legitimate selectable value and is preserved in every domain.",
        # R2 / R3
        "hertz_to_voltage":hertz_to_voltage,
        "hertz_to_rpm":hertz_to_rpm,
        # R4
        "frame_hp_rpm_hex_count":len(frame_hp_rpm_hex),
        "frame_hp_rpm_hex":frame_hp_rpm_hex,
        # R5
        "combination_fields":[n for _,n in R5_FIELD_COLS],
        "combination_matrix":{
            "total_combinations":total,"distinct_hex_codes":len(hexset),
            "hex_code_range":{"first":sh[0] if sh else None,"last":sh[-1] if sh else None},
            "distinct_values_per_field":{n:sorted(v) for n,v in fv.items()},
            "all_rows":combos,
        },
        "lookup_key_col":"C","hex_code_col":"O (3-char)",
        "part_number_role":"Motor Assembly segment (3-char hex). VLOOKUP(concat(D..N),col_C,col_O).",
        "hex_formula":"O=BASE(P,36) padded to 3; R4 AD=BASE(AC,36) padded to 2.",
    }


def write_outputs(evidence_dir,result):
    payload={"artifact":"FYBROC_V6_MOTOR_ASSY",**result}
    (evidence_dir/"FYBROC_V6_MOTOR_ASSY.json").write_text(
        json.dumps(payload,indent=2,ensure_ascii=False,default=str),encoding="utf-8")
    cm=result["combination_matrix"]; line="="*120
    out=[f"{line}\r\nF110.9 - FYBROC V6 MOTOR ASSY (5 regions)\r\n{line}\r\n\r\n"]
    out.append(f"Git commit  : {result['git_commit']}\r\nGit clean   : {result['git_working_tree_clean']}\r\n\r\n")

    out.append(f"{line}\r\nR1 - OPTION DOMAINS (D2:Q21, '-' KEPT)\r\n{line}\r\n\r\n")
    for f,opts in result["option_domains"].items():
        out.append(f"  {f} ({len(opts)}): {opts}\r\n")

    out.append(f"\r\n{line}\r\nR2 - MOTOR HERTZ -> ALLOWABLE VOLTAGE\r\n{line}\r\n\r\n")
    for k,v in result["hertz_to_voltage"].items():
        out.append(f"  {k} ({len(v)}): {v}\r\n")
    out.append(f"\r\n{line}\r\nR3 - MOTOR HERTZ -> ALLOWABLE RPM\r\n{line}\r\n\r\n")
    for k,v in result["hertz_to_rpm"].items():
        out.append(f"  {k} ({len(v)}): {v}\r\n")

    out.append(f"\r\n{line}\r\nR4 - FRAME + HP + RPM -> HEX ({result['frame_hp_rpm_hex_count']} rows)\r\n{line}\r\n\r\n")
    for row in result["frame_hp_rpm_hex"]:
        out.append(f"  frame={row['motor_frame']:>6}  hp={row['motor_horsepower']:>5}  rpm={row['motor_rpm']:>5}  id={row['id']:>3}  hex={row['hex']}\r\n")

    out.append(f"\r\n{line}\r\nR5 - COMBINATION -> HEX\r\n{line}\r\n\r\n")
    out.append(f"Total combinations : {cm['total_combinations']}\r\nDistinct hex-codes : {cm['distinct_hex_codes']}\r\nHex range          : {cm['hex_code_range']['first']} -> {cm['hex_code_range']['last']}\r\n\r\n")
    for f,vals in cm["distinct_values_per_field"].items():
        out.append(f"  {f} ({len(vals)}): {vals}\r\n")
    out.append("\r\n  first 20 combination rows:\r\n")
    for r in cm["all_rows"][:20]:
        out.append(f"    id={r['id']:>4} hex={r['hex_code']} key={r['key']}\r\n")
    (evidence_dir/"FYBROC_V6_MOTOR_ASSY.txt").write_text("".join(out),encoding="utf-8")


def main():
    p=argparse.ArgumentParser(); root=Path(__file__).resolve().parent.parent
    p.add_argument("--repo-root",type=Path,default=root); p.add_argument("--evidence-dir",type=Path,default=None)
    a=p.parse_args(); repo_root=a.repo_root.resolve()
    ev=(a.evidence_dir or (repo_root/"docs"/"evidence"/"F110")).resolve(); ev.mkdir(parents=True,exist_ok=True)
    result=build_model(repo_root); write_outputs(ev,result)
    cm=result["combination_matrix"]
    print(json.dumps({
        "step":STEP,"output_dir":str(ev),
        "option_domain_sizes":{k:len(v) for k,v in result["option_domains"].items()},
        "hertz_to_voltage_keys":list(result["hertz_to_voltage"].keys()),
        "hertz_to_rpm_keys":list(result["hertz_to_rpm"].keys()),
        "frame_hp_rpm_hex_count":result["frame_hp_rpm_hex_count"],
        "combination_total":cm["total_combinations"],
        "combination_hex_range":cm["hex_code_range"],
    },indent=2))
    return 0

if __name__=="__main__": raise SystemExit(main())
