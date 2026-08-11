from pathlib import Path
import json

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[1]

WORKBOOK = (
    ROOT
    / "workbooks"
    / "Dean"
    / "PumpConfiguration_Logic.xlsm"
)

EXPORT_DIR = ROOT / "exports"
OUTPUT_DIR = ROOT / "output"

EXPORT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


def table_rows(ws, table):
    min_col, min_row, max_col, max_row = (
        range_boundaries(table.ref)
    )

    headers = [
        ws.cell(
            min_row,
            col,
        ).value
        for col in range(
            min_col,
            max_col + 1,
        )
    ]

    rows = []

    for row in range(
        min_row + 1,
        max_row + 1,
    ):
        values = [
            ws.cell(
                row,
                col,
            ).value
            for col in range(
                min_col,
                max_col + 1,
            )
        ]

        if not any(
            value not in (
                None,
                "",
            )
            for value in values
        ):
            continue

        item = {
            str(headers[i]): values[i]
            for i in range(
                len(headers)
            )
        }

        rows.append(item)

    return {
        "table": table.name,
        "range": table.ref,
        "headers": headers,
        "rows": rows,
    }


def export_structured_tables(wb):
    result = {}

    for sheet_name in (
        "Config Options",
        "Codependencies",
    ):
        ws = wb[sheet_name]

        tables = []

        for table in ws.tables.values():
            tables.append(
                table_rows(
                    ws,
                    table,
                )
            )

        result[
            sheet_name
        ] = tables

    path = (
        EXPORT_DIR
        / "dean_pumpconfiguration_logic.json"
    )

    path.write_text(
        json.dumps(
            result,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )

    return path


def export_raw_sheet(ws, output_name):
    lines = [
        "=" * 120,
        f"SHEET: {ws.title}",
        "=" * 120,
        "",
    ]

    for row in ws.iter_rows():
        populated = []

        for cell in row:
            if cell.value in (
                None,
                "",
            ):
                continue

            populated.append(
                f"{cell.coordinate}={cell.value!r}"
            )

        if populated:
            lines.append(
                " | ".join(populated)
            )

    path = (
        OUTPUT_DIR
        / output_name
    )

    path.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )

    return path


def main():
    wb = load_workbook(
        WORKBOOK,
        read_only=False,
        data_only=False,
        keep_vba=True,
        keep_links=False,
    )

    try:
        metadata_path = (
            export_structured_tables(
                wb
            )
        )

        pump_path = export_raw_sheet(
            wb["Pump Options"],
            "dean_pump_options_trace.txt",
        )

        price_path = export_raw_sheet(
            wb["Price Options"],
            "dean_price_options_trace.txt",
        )

    finally:
        wb.close()

    print(
        "DEAN LOGIC METADATA EXTRACTED"
    )
    print(
        "Structured metadata:",
        metadata_path,
    )
    print(
        "Pump Options:",
        pump_path,
    )
    print(
        "Price Options:",
        price_path,
    )


if __name__ == "__main__":
    main()
