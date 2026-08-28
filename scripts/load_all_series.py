"""Load ALL Fybroc series into SeriesFieldOption (including vertical series).

Series codes are normalized to canonical TEXT via norm_series() because the
V6 Attributes table stores some series as text and 2530/3000/5500 as integers
(see docs/DATA_TYPE_NORMALIZATION.md). No other field types are altered here.
"""
import sys
from pathlib import Path

import openpyxl
import pyodbc

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.compiler.workbook_types import norm_series

conn_str = (
    "DRIVER={ODBC Driver 18 for SQL Server};"
    "SERVER=localhost;"
    "DATABASE=PumpConfiguratorDB;"
    "Trusted_Connection=yes;"
    "Encrypt=yes;"
    "TrustServerCertificate=yes;"
)

def main():
    wb = openpyxl.load_workbook(
        "workbooks/Fybroc/Fybroc Configuration Rev0.3.xlsx",
        read_only=True, data_only=True,
    )
    ws = wb["Selections"]

    # Read series from header (row 1, cols 4-13). Normalize to canonical text
    # so a numeric-stored series (e.g. 5500) becomes "5500", not "5500.0".
    series_codes = []
    for c in range(4, 14):
        s = norm_series(ws.cell(row=1, column=c).value)
        if s:
            series_codes.append(s)
    print(f"Series in Selections: {series_codes}")

    # V6-AUTHORITATIVE flange rule (Nomenclature_V6.xlsm Attributes Series+Flange
    # code table, rows 23-42): DIN and JIS codes exist ONLY for these series.
    # All other series are ANSI-only. V6 governs over Rev0.3 Selections, which is
    # over-permissive (it marks DIN/JIS selectable for 2530 and 5500 too, but V6
    # has no code for those combinations, so we do not offer them).
    V6_DIN_JIS_SERIES = {"1500", "1530", "1600", "1630"}
    NON_ANSI_FLANGE_ANSWERS = {"din/iso flange", "jis flange"}

    # Read all rows
    rows = []
    flange_rows_dropped = 0
    for r in range(2, 679):
        question = ws.cell(row=r, column=2).value
        answer = ws.cell(row=r, column=3).value
        if not question or not answer:
            continue
        question = str(question).strip()
        answer = str(answer).strip()
        field_code = question.upper().replace(" ", "_").replace("-", "_")

        for i, series in enumerate(series_codes):
            marker = ws.cell(row=r, column=4 + i).value
            marker = str(marker).strip() if marker is not None else ""
            marker_upper = marker.upper()
            # Rev0.3 Selections three-state semantics:
            #   "STD"  -> selectable AND the standard/default for this series
            #   "X"    -> selectable (non-default)
            #   blank  -> not selectable (skip)
            if marker_upper in ("X", "STD"):
                # Enforce V6 authority on flange availability.
                if (
                    field_code == "FLANGE_TYPE"
                    and answer.lower() in NON_ANSI_FLANGE_ANSWERS
                    and series not in V6_DIN_JIS_SERIES
                ):
                    flange_rows_dropped += 1
                    continue
                is_standard = 1 if marker_upper == "STD" else 0
                rows.append((field_code, answer.lower(), series, marker_upper, is_standard))

    wb.close()
    std_count = sum(1 for row in rows if row[4] == 1)
    print(f"Total rows: {len(rows)} (standard defaults: {std_count})")
    print(
        f"V6 flange authority: dropped {flange_rows_dropped} DIN/JIS rows for "
        f"non-{sorted(V6_DIN_JIS_SERIES)} series (over-permissive in Rev0.3)."
    )

    # Load into SQL
    conn = pyodbc.connect(conn_str, autocommit=False)
    cursor = conn.cursor()

    pub_id = 2
    family_id = 2

    # Clear and reload
    cursor.execute("DELETE FROM cfg.SeriesFieldOption WHERE MetadataPublicationId=?", pub_id)
    conn.commit()

    # Batch insert
    batch_size = 500
    for i in range(0, len(rows), batch_size):
        batch = rows[i:i + batch_size]
        cursor.executemany(
            "INSERT INTO cfg.SeriesFieldOption "
            "(MetadataPublicationId, PumpFamilyId, SourceFieldCode, FieldCode, OptionValue, "
            " SeriesCode, WorkbookName, WorksheetName, SourceRow, SourceFieldCell, SourceValueCell, SourceSeriesCell, "
            " SelectionMarker, IsStandard) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(pub_id, family_id, "", fc, val, series,
              "Fybroc Configuration Rev0.3.xlsx", "Selections", 0, "", "", "",
              marker, is_standard)
             for fc, val, series, marker, is_standard in batch],
        )
    conn.commit()

    # Verify
    total = cursor.execute(
        "SELECT COUNT(*) FROM cfg.SeriesFieldOption WHERE MetadataPublicationId=?", pub_id
    ).fetchone()[0]
    print(f"\nPublished: {total} rows")

    for r in cursor.execute(
        "SELECT SeriesCode, COUNT(*) FROM cfg.SeriesFieldOption "
        "WHERE MetadataPublicationId=? GROUP BY SeriesCode ORDER BY SeriesCode", pub_id
    ).fetchall():
        print(f"  {r[0]}: {r[1]}")

    std_total = cursor.execute(
        "SELECT COUNT(*) FROM cfg.SeriesFieldOption "
        "WHERE MetadataPublicationId=? AND IsStandard=1", pub_id
    ).fetchone()[0]
    print(f"\nStandard (STD) defaults published: {std_total}")

    conn.close()


if __name__ == "__main__":
    main()
