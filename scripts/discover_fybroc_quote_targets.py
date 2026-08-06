from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKBOOK_PATH = Path(
    r"workbooks\Fybroc\Price Estimator-Fybroc.xlsm"
)

OUTPUT_PATH = Path(
    r"output\m020_fybroc_quote_targets.json"
)

PREFERRED_SHEETS = (
    "Formal Quote",
    "FORMAL QUOTES",
    "MainTemplate",
    "Price Check",
    "Top Level",
    "M-#",
)

SEARCH_TERMS = (
    "part number",
    "part no",
    "part #",
    "smart number",
    "sku",
    "stock keeping",
    "model number",
    "model no",
    "catalog number",
    "quote number",
)


def normalize(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value).strip().lower().split()
    )


def merged_ranges_for_cell(
    worksheet,
    coordinate: str,
) -> list[str]:
    matches: list[str] = []

    for merged_range in worksheet.merged_cells.ranges:
        if coordinate in merged_range:
            matches.append(str(merged_range))

    return matches


def nearby_cells(
    worksheet,
    row_number: int,
    column_number: int,
    *,
    row_radius: int = 2,
    column_radius: int = 5,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []

    minimum_row = max(1, row_number - row_radius)
    maximum_row = min(
        worksheet.max_row,
        row_number + row_radius,
    )

    minimum_column = max(
        1,
        column_number - column_radius,
    )
    maximum_column = min(
        worksheet.max_column,
        column_number + column_radius,
    )

    for row in worksheet.iter_rows(
        min_row=minimum_row,
        max_row=maximum_row,
        min_col=minimum_column,
        max_col=maximum_column,
    ):
        for cell in row:
            if cell.value is None:
                continue

            value = str(cell.value).strip()

            if not value:
                continue

            results.append(
                {
                    "coordinate": cell.coordinate,
                    "value": value,
                    "data_type": cell.data_type,
                    "merged_ranges": (
                        merged_ranges_for_cell(
                            worksheet,
                            cell.coordinate,
                        )
                    ),
                }
            )

    return results


def defined_name_text(defined_name) -> str:
    name = getattr(defined_name, "name", "")
    value = getattr(defined_name, "attr_text", "")

    return f"{name} {value}"


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Workbook not found: {WORKBOOK_PATH}"
        )

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=False,
        data_only=False,
        keep_vba=True,
        keep_links=True,
    )

    sheets_to_scan = [
        sheet_name
        for sheet_name in PREFERRED_SHEETS
        if sheet_name in workbook.sheetnames
    ]

    report: dict[str, Any] = {
        "workbook": str(WORKBOOK_PATH),
        "available_sheets": workbook.sheetnames,
        "scanned_sheets": sheets_to_scan,
        "matches": [],
        "matching_defined_names": [],
    }

    for sheet_name in sheets_to_scan:
        worksheet = workbook[sheet_name]

        for row in worksheet.iter_rows():
            for cell in row:
                normalized_value = normalize(cell.value)

                if not normalized_value:
                    continue

                matched_terms = [
                    term
                    for term in SEARCH_TERMS
                    if term in normalized_value
                ]

                if not matched_terms:
                    continue

                report["matches"].append(
                    {
                        "sheet": sheet_name,
                        "coordinate": cell.coordinate,
                        "value": str(cell.value),
                        "matched_terms": matched_terms,
                        "merged_ranges": (
                            merged_ranges_for_cell(
                                worksheet,
                                cell.coordinate,
                            )
                        ),
                        "nearby_cells": nearby_cells(
                            worksheet,
                            cell.row,
                            cell.column,
                        ),
                    }
                )

    for defined_name in workbook.defined_names.values():
        text = normalize(
            defined_name_text(defined_name)
        )

        matched_terms = [
            term
            for term in SEARCH_TERMS
            if term in text
        ]

        if not matched_terms:
            continue

        report["matching_defined_names"].append(
            {
                "name": defined_name.name,
                "refers_to": defined_name.attr_text,
                "matched_terms": matched_terms,
            }
        )

    OUTPUT_PATH.write_text(
        json.dumps(
            report,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print("=" * 90)
    print("FYBROC QUOTE TARGET DISCOVERY")
    print("=" * 90)
    print(f"Workbook: {WORKBOOK_PATH}")
    print(
        "Sheets scanned: "
        + ", ".join(sheets_to_scan)
    )
    print(
        f"Cell matches: {len(report['matches'])}"
    )
    print(
        "Matching defined names: "
        f"{len(report['matching_defined_names'])}"
    )
    print()

    for match in report["matches"]:
        print("-" * 90)
        print(
            f"{match['sheet']}!"
            f"{match['coordinate']} = "
            f"{match['value']!r}"
        )

        if match["merged_ranges"]:
            print(
                "Merged ranges: "
                + ", ".join(match["merged_ranges"])
            )

        print("Nearby populated cells:")

        for nearby in match["nearby_cells"]:
            print(
                f"  {nearby['coordinate']}: "
                f"{nearby['value']}"
            )

    if report["matching_defined_names"]:
        print()
        print("=" * 90)
        print("MATCHING DEFINED NAMES")
        print("=" * 90)

        for item in report["matching_defined_names"]:
            print(
                f"{item['name']} -> "
                f"{item['refers_to']}"
            )

    print()
    print(f"JSON report: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
