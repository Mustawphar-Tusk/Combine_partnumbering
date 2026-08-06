from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = (
    PROJECT_ROOT
    / "workbooks"
    / "Fybroc"
    / "Fybroc Attributes and Constraints.xlsx"
)
OUTPUT_JSON = (
    PROJECT_ROOT
    / "exports"
    / "fybroc_model_constraint_profile.json"
)
OUTPUT_CSV = (
    PROJECT_ROOT
    / "exports"
    / "fybroc_model_constraint_sheet_profile.csv"
)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def main() -> None:
    if not WORKBOOK.exists():
        raise SystemExit(
            f"Workbook not found: {WORKBOOK}"
        )

    workbook = load_workbook(
        WORKBOOK,
        read_only=True,
        data_only=False,
        keep_links=False,
    )

    sheet_profiles: list[dict[str, Any]] = []
    row_samples: list[dict[str, Any]] = []

    try:
        for worksheet in workbook.worksheets:
            profile = {
                "worksheet_name": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "dimensions": worksheet.calculate_dimension(),
            }
            sheet_profiles.append(profile)

            for row_number, cells in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(
                        worksheet.max_row or 1,
                        30,
                    ),
                    min_col=1,
                    max_col=min(
                        worksheet.max_column or 1,
                        50,
                    ),
                    values_only=False,
                ),
                start=1,
            ):
                populated = {
                    cell.coordinate: _text(cell.value)
                    for cell in cells
                    if _text(cell.value) is not None
                }

                if populated:
                    row_samples.append(
                        {
                            "worksheet_name": worksheet.title,
                            "row_number": row_number,
                            "cells": populated,
                        }
                    )

    finally:
        workbook.close()

    OUTPUT_JSON.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT_JSON.write_text(
        json.dumps(
            {
                "workbook": str(WORKBOOK),
                "sheet_profiles": sheet_profiles,
                "row_samples": row_samples,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    with OUTPUT_CSV.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "worksheet_name",
                "max_row",
                "max_column",
                "dimensions",
            ],
        )
        writer.writeheader()
        writer.writerows(sheet_profiles)

    print(
        f"Worksheets profiled: {len(sheet_profiles)}"
    )
    print(
        f"Populated sample rows recorded: "
        f"{len(row_samples)}"
    )
    print(f"JSON report: {OUTPUT_JSON}")
    print(f"CSV report: {OUTPUT_CSV}")


if __name__ == "__main__":
    main()
