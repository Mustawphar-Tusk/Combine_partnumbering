"""D140 (re-base) — Load the Dean identifier authority into SQL from the NEW
authoritative workbook workbooks/Dean/PumpConfiguration_Logic_0.1.xlsm.

This SUPERSEDES the prior version of this loader (which read the old
`Dean Data Sheet Rev 2.xlsm` numbering sheets and the m023 reconciliation JSON).
Now the offered options (D110), the numbering (this loader), and the identifier
resolver all derive from ONE workbook.

Loads TWO things, both family-scoped to DEAN and additive (never touches Fybroc):

  1. A#/D# model identity  ->  cfg.PumpModelReference
     From `Pump Constraints` cols A(`A Number`)/B(`Series`)/C(`Size`), data rows
     4..209 (206 models). D# (BaseIdentifier) = 'D' + numeric part of the A#
     (e.g. A461 -> D461), matching the D130 convention already in the DB.

  2. Segment String->Code maps  ->  stg.SegmentCombinationImport (DEAN batch),
     the SAME runtime store Fybroc uses. Each numbering sheet is located by
     HEADER SEARCH (Module1/Module2 semantics): find the row containing
     `Permutation` (col C) and `Alphanumeric Code` (last table col); the OPTION
     columns are strictly between them (left->right); the ComboString is the
     '*'-join of those option values (per row) with trailing '*' trimmed; the
     Alphanumeric Code is the literal zero-padded sequential code.

     SelectionsJson is keyed by the CANONICAL DEAN field code (so the resolver's
     _build_combo produces a matching CombinationKey). CombinationKey = the same
     '*'-joined string, so the resolver's exact-hash lookup succeeds.

Segments present as numbering sheets in v0.1 (LOADED here):
  WET_END_OPTIONS, IMPELLER_OPTIONS, POWER_FRAME_OPTIONS, BASEPLATE_OPTIONS,
  FLUSH_PLAN, MOTOR_FRAME (Motor Frame-Size sub-table S/T/U — the live motor code;
  the main Motor code col O is inert '000' and is NOT loaded).

Segments with NO numbering table in v0.1 (retained-but-unbuilt gaps — see
docs/evidence/DEAN_ENGINEERING_QUESTIONS.md §F and D140 evidence):
  BARRIER_PLAN (header-only, 0 rows — F2), COOLING_PLAN (sheet empty — F1),
  TESTING / DOCUMENTATION / ADDITIONAL_OPTIONS (no sheet in v0.1).
The resolver emits gated/placeholder codes for those so the PN never errors.

Idempotent: re-running replaces the DEAN model-reference rows and the DEAN
segment-combination batch for the same source-file hash. Asserts Fybroc row
counts are unchanged.

Usage:
  $env:PYTHONPATH="."
  python scripts/load_dean_identifier_to_sql.py [--dry-run]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import pyodbc
from openpyxl import load_workbook

CONN = (
    "DRIVER={ODBC Driver 18 for SQL Server};SERVER=localhost;"
    "DATABASE=PumpConfiguratorDB;Trusted_Connection=yes;"
    "Encrypt=yes;TrustServerCertificate=yes;"
)

WB_PATH = Path("workbooks/Dean/PumpConfiguration_Logic_0.1.xlsm")
FAMILY = "DEAN"

# ---------------------------------------------------------------------------
# Segment definitions. Each numbering sheet is found by HEADER SEARCH; we only
# hard-code (a) the sheet name, (b) the segment code/name, (c) the expected code
# width, and (d) the mapping from each OPTION-column header label -> canonical
# DEAN field code (None = position has no direct Dean SFO field). The option
# column ORDER is taken from the sheet itself (header row), so the loader tracks
# the workbook rather than a hand-transcribed order.
#
# `header_scan_max` bounds the header search (Flush's header is at row 26).
# `subtable` marks the Motor Frame-Size sub-table (its own Permutation/Frame
# Size/Alphanumeric Code trio in cols S/T/U).
# ---------------------------------------------------------------------------
SEGMENTS = [
    {
        "sheet": "Wet End Numbering",
        "segment_code": "WET_END_OPTIONS",
        "segment_name": "Wet End Options",
        "width": 4,
        "field_map": {
            "Pump Material": "PUMP_MATERIAL",
            "Casing Material": "CASING_MATERIAL",
            "Casing Drain": "CASING_DRAIN",
            "Casing Taps": "CASING_TAPS",
            "Casing Gasket": "CASING_GASKET",
            "Flange Configuration": "FLANGE_CONFIGURATION",
            "Spot Facing": "SPOT_FACING",
            "Casing Wear Ring": "CASING_WEAR_RING",
            "Casing Mounting": "CASING_MOUNTING",
            "Seal Chamber Config": "SEAL_CHAMBER_CONFIG",
        },
    },
    {
        "sheet": "Impeller Numbering",
        "segment_code": "IMPELLER_OPTIONS",
        "segment_name": "Impeller Options",
        "width": 2,
        "field_map": {
            "Impeller Balance": "IMPELLER_BALANCE",
            "Impeller Material": "IMPELLER_MATERIAL",
            "Impeller Wear Ring Material": "IMPELLER_WEAR_RING_MATERIAL",
        },
    },
    {
        "sheet": "Power Frame Numbering",
        "segment_code": "POWER_FRAME_OPTIONS",
        "segment_name": "Power Frame Options",
        "width": 4,
        "field_map": {
            "Shaft Configuration": "SHAFT_CONFIGURATION",
            "Shaft Material": "SHAFT_MATERIAL",
            "Bearing Lubrication": "BEARING_LUBRICATION",
            "Bearing Seal": "BEARING_SEAL",
            "Oiler Options": "OILER_OPTIONS",
            "Sight Glass": "SIGHT_GLASS",
            "Bearing Frame Cooling": "BEARING_FRAME_COOLING",
            "Magnetic Drain": "MAGNETIC_DRAIN",
            "Expansion Chamber": "EXPANSION_CHAMBER",
            "Coupling Type": "COUPLING_TYPE",
            "Coupling Guard": "COUPLING_GUARD",
        },
    },
    {
        "sheet": "Flush Plan Numbering",
        "segment_code": "FLUSH_PLAN",
        "segment_name": "Flush Plan",
        "width": 2,
        "header_scan_max": 40,   # header is at row 26
        "field_map": {
            "Flush Plan": "FLUSH_PLAN",
            # routing/connections/temp/media have no canonical Dean option code;
            # they participate in the ComboString by their literal workbook value.
            "Flush Plan Routing": None,
            "Flush Connections": None,
            "Flush Temperature Range": None,
            "Flush Cooling Media": None,
        },
    },
    {
        "sheet": "Baseplate Numbering",
        "segment_code": "BASEPLATE_OPTIONS",
        "segment_name": "Baseplate Options",
        "width": 2,
        "field_map": {
            "Baseplate Type": "BASEPLATE_TYPE",
            "Drip Pan": "DRIP_PAN",
            "Alignment Lugs": "ALIGNMENT_LUGS",
            "Lifting Lugs": "LIFTING_LUGS",
            "Levelling Screws": "LEVELLING_SCREWS",
            "Grounding Lug": "GROUNDING_LUG",
            "Grout Hole": "GROUT_HOLE",
            "Isolation Pads": "ISOLATION_PADS",
            "Stilts": "STILTS",
        },
    },
    {
        # Motor: the MAIN code column (O) is inert ('000'). The live motor code is
        # the Frame-Size sub-table in cols S(Permutation)/T(Frame Size)/U(code).
        "sheet": "Motor Numbering",
        "segment_code": "MOTOR_FRAME",
        "segment_name": "Motor Frame Size",
        "width": 2,
        "subtable": {"perm_col": 18, "value_col": 19, "code_col": 20,  # S/T/U 0-based
                     "header_row": 16, "value_field": "FRAME_SIZE"},
    },
]

# Segments with NO numbering table in v0.1 (documented, retained-unbuilt).
UNBUILT_SEGMENTS = {
    "BARRIER_PLAN": "header-only skeleton, 0 data rows (engineering gap F2)",
    "COOLING_PLAN": "sheet empty in v0.1 (engineering gap F1)",
    "TESTING": "no numbering sheet in v0.1",
    "DOCUMENTATION": "no numbering sheet in v0.1",
    "ADDITIONAL_OPTIONS": "no numbering sheet in v0.1",
}


def sha256_file(path: Path) -> str:
    d = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            d.update(chunk)
    return d.hexdigest()


def col_letter(i0: int) -> str:
    s = ""
    c = i0 + 1
    while c > 0:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


# ---------------------------------------------------------------------------
# 1. A#/D# identity  (from Pump Constraints A/B/C)
# ---------------------------------------------------------------------------
def _d_number(a_number: str) -> str:
    """D# = 'D' + the numeric part of the A# (A461 -> D461)."""
    m = re.search(r"(\d+)", a_number or "")
    return ("D" + m.group(1)) if m else ""


def read_models(wb):
    """Yield (series, size, a_number, d_number, source_row) from Pump Constraints.
    Identity headers (A Number / Series / Size) are on row 3; data rows 4.."""
    ws = wb["Pump Constraints"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row,
                             max_col=min(ws.max_column, 8), values_only=True))
    row3 = rows[2] if len(rows) > 2 else ()
    ac = sc = zc = None
    for c0, v in enumerate(row3):
        h = ("" if v is None else str(v)).strip().lower()
        if h == "a number":
            ac = c0
        elif h == "series":
            sc = c0
        elif h == "size":
            zc = c0
    if ac is None or sc is None or zc is None:
        raise SystemExit(f"Pump Constraints identity cols not found "
                         f"(A={ac} Series={sc} Size={zc})")
    out = []
    for r_idx, r in enumerate(rows[3:], start=4):
        ser = r[sc] if sc < len(r) else None
        if ser is None or str(ser).strip() == "":
            continue
        a = str(r[ac]).strip() if ac < len(r) and r[ac] is not None else ""
        size = str(r[zc]).strip() if zc < len(r) and r[zc] is not None else ""
        d = _d_number(a)
        if not (a and d and size):
            continue
        out.append((str(ser).strip(), size, a, d, r_idx))
    return out


def load_models(conn, wb, dry_run: bool) -> int:
    rows = read_models(wb)
    print(f"[models] {len(rows)} A#/D# identity rows parsed from Pump Constraints")
    if dry_run:
        for r in rows[:5]:
            print("   sample:", r)
        return len(rows)

    cur = conn.cursor()
    fam_id = cur.execute(
        "SELECT PumpFamilyId FROM cfg.PumpFamily WHERE FamilyCode=?", FAMILY
    ).fetchone()[0]
    cv_id = cur.execute(
        "SELECT ConfigurationVersionId FROM cfg.ConfigurationVersion "
        "WHERE PumpFamilyId=? AND IsCurrent=1", fam_id
    ).fetchone()[0]

    cur.execute(
        "DELETE FROM cfg.PumpModelReference WHERE PumpFamilyId=? AND ConfigurationVersionId=?",
        fam_id, cv_id,
    )
    ins = (
        "INSERT INTO cfg.PumpModelReference "
        "(PumpFamilyId, ConfigurationVersionId, ModelIdentifier, SeriesCode, "
        " SizeCode, BaseIdentifier, IsActive, SourceWorksheet, SourceRow) "
        "VALUES (?, ?, ?, ?, ?, ?, 1, ?, ?)"
    )
    # dedupe on (series,size) — the UX index is (family,cv,series,size) WHERE IsActive=1
    seen = {}
    for series, size, a, d, srow in rows:
        seen[(series.upper(), size.upper())] = (series.upper(), size.upper(), a, d, srow)
    payload = [(fam_id, cv_id, a, s, z, d, "Pump Constraints", srow)
               for (s, z, a, d, srow) in seen.values()]
    cur.fast_executemany = True
    cur.executemany(ins, payload)
    conn.commit()
    print(f"[models] loaded {len(payload)} unique (series,size) rows into "
          f"cfg.PumpModelReference (family={fam_id}, cv={cv_id})")
    return len(payload)


# ---------------------------------------------------------------------------
# 2. Segment String->Code maps  (header-search per sheet)
# ---------------------------------------------------------------------------
def find_table(ws, scan_max=40):
    """Return (header_row1, perm_col0, code_col0) via header search for
    'Permutation' + 'Alphanumeric Code' in the same row. 1-based row, 0-based cols.
    Returns (None, None, None) if not found; (row, perm, None) if code missing."""
    for r in range(1, scan_max + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 30)]
        perm = code = None
        for c0, v in enumerate(vals):
            sv = ("" if v is None else str(v)).strip().lower()
            if sv == "permutation":
                perm = c0
            elif sv == "alphanumeric code":
                code = c0
        if perm is not None:
            return r, perm, code
    return None, None, None


def extract_segment_rows(wb, seg):
    """Yield (source_id, source_row, combo_string, code, selections_dict) for a
    segment. selections_dict is keyed by CANONICAL Dean field code (None-mapped
    option columns are omitted from the dict but STILL included in the ComboString
    by their literal value, matching the resolver's positional _build_combo).

    For the Motor Frame sub-table (subtable mode) there is a single option leg
    (Frame Size) and the code lives in its own U column.
    """
    ws = wb[seg["sheet"]]

    if "subtable" in seg:
        st = seg["subtable"]
        vcol, ccol, width = st["value_col"], st["code_col"], seg["width"]
        vfield = st["value_field"]
        hr = st["header_row"]
        source_id = 0
        for r_idx, row in enumerate(
            ws.iter_rows(min_row=hr + 1, max_col=ccol + 1, values_only=True),
            start=hr + 1,
        ):
            code = row[ccol] if ccol < len(row) else None
            if code is None or str(code).strip() == "":
                continue
            code_s = str(code).strip()
            if len(code_s) != width:
                continue
            val = row[vcol] if vcol < len(row) else None
            val_s = "" if val is None else str(val).strip()
            if val_s == "" or val_s.lower() in ("frame size", "alphanumeric code"):
                continue
            source_id += 1
            yield source_id, r_idx, val_s, code_s, {vfield: val_s}
        return

    scan_max = seg.get("header_scan_max", 20)
    hr, perm, ccol = find_table(ws, scan_max)
    if hr is None or ccol is None:
        return  # empty / header-only -> no rows
    opt_cols = list(range(perm + 1, ccol))
    headers = [("" if ws.cell(row=hr, column=c + 1).value is None
                else str(ws.cell(row=hr, column=c + 1).value).strip())
               for c in opt_cols]
    # per-position canonical code (None if the header has no Dean field code)
    fmap = seg["field_map"]
    pos_code = [fmap.get(h) for h in headers]
    width = seg["width"]
    max_col = ccol + 1

    source_id = 0
    for r_idx, row in enumerate(
        ws.iter_rows(min_row=hr + 1, max_col=max_col, values_only=True),
        start=hr + 1,
    ):
        code = row[ccol] if ccol < len(row) else None
        if code is None:
            continue
        code_s = str(code).strip()
        if code_s == "" or code_s.lower() in ("alphanumeric code", "code"):
            continue
        if len(code_s) != width:  # mirrors SQL CHECK(LEN=ExpectedWidth)
            continue
        parts = []
        for c in opt_cols:
            v = row[c] if c < len(row) else None
            parts.append("" if v is None else str(v).strip())
        combo_s = "*".join(parts)
        while combo_s.endswith("*"):        # Module1/Module2 trailing-* trim
            combo_s = combo_s[:-1]
        if combo_s == "":
            continue
        selections = {}
        for i, val in enumerate(parts):
            code_i = pos_code[i] if i < len(pos_code) else None
            if code_i:
                selections[code_i] = val
        source_id += 1
        yield source_id, r_idx, combo_s, code_s, selections


def load_segments(conn, wb, dry_run: bool) -> dict:
    source_hash = sha256_file(WB_PATH)
    counts = {}
    if dry_run:
        for seg in SEGMENTS:
            n = 0
            sample = None
            for sid, srow, combo_s, code_s, sel in extract_segment_rows(wb, seg):
                n += 1
                if sample is None:
                    sample = (combo_s, code_s)
            counts[seg["segment_code"]] = n
            print(f"[segment] {seg['segment_code']:<22} {seg['sheet']:<24} "
                  f"width={seg['width']}  rows={n:,}  sample={sample}")
        print(f"[segment] unbuilt (retained gaps): {sorted(UNBUILT_SEGMENTS)}")
        return counts

    cur = conn.cursor()
    # Idempotent + supersede: delete ALL prior DEAN batches (this-hash re-runs AND
    # the old-workbook batch, which is no longer authoritative). This keeps the
    # runtime store to exactly one current DEAN batch and prevents the resolver's
    # "TOP 1 ... Status='Loaded' ORDER BY ImportBatchId DESC" from ever seeing a
    # stale one. Fybroc batches (FamilyCode<>DEAN) are never touched.
    old = [int(r[0]) for r in cur.execute(
        "SELECT ImportBatchId FROM stg.SegmentCombinationImportBatch WHERE FamilyCode=?",
        FAMILY).fetchall()]
    for bid in old:
        cur.execute("DELETE FROM stg.SegmentCombinationImport WHERE ImportBatchId=?", bid)
        cur.execute("DELETE FROM stg.SegmentCombinationImportBatch WHERE ImportBatchId=?", bid)
    conn.commit()
    if old:
        print(f"[segment] removed {len(old)} prior DEAN batch(es): {old}")

    batch_id = int(cur.execute(
        "INSERT INTO stg.SegmentCombinationImportBatch "
        "(FamilyCode, SourceFile, SourceFileHash, ExpectedRowCount, Status) "
        "OUTPUT INSERTED.ImportBatchId VALUES (?, ?, ?, 0, 'Loading')",
        FAMILY, str(WB_PATH.resolve()), source_hash,
    ).fetchone()[0])
    conn.commit()

    ins = (
        "INSERT INTO stg.SegmentCombinationImport "
        "(ImportBatchId, FamilyCode, WorkbookRole, WorkbookName, WorksheetName, "
        " SegmentCode, SegmentName, SourceRow, SourceId, SegmentValue, "
        " ExpectedWidth, CombinationKey, SelectionsJson, SourceCellsJson, SourceProfile) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    )
    cur.fast_executemany = True
    total = 0
    for seg in SEGMENTS:
        buf = []
        seg_rows = 0
        for source_id, srow, combo_s, code_s, selections in extract_segment_rows(wb, seg):
            sel_json = json.dumps(selections, separators=(",", ":"))
            src_desc = {"sheet": seg["sheet"], "row": srow,
                        "mode": "subtable" if "subtable" in seg else "header_search"}
            cells_json = json.dumps(src_desc, separators=(",", ":"))
            buf.append((
                batch_id, FAMILY, "configuration_and_quote",
                WB_PATH.name, seg["sheet"], seg["segment_code"], seg["segment_name"],
                srow, source_id, code_s, seg["width"], combo_s, sel_json,
                cells_json, "dean_numbering_v01",
            ))
            if len(buf) >= 2000:
                cur.executemany(ins, buf)
                conn.commit()
                total += len(buf)
                seg_rows += len(buf)
                buf = []
        if buf:
            cur.executemany(ins, buf)
            conn.commit()
            total += len(buf)
            seg_rows += len(buf)
        counts[seg["segment_code"]] = seg_rows
        print(f"[segment] {seg['segment_code']:<22} loaded {seg_rows:,} rows")

    cur.execute(
        "UPDATE stg.SegmentCombinationImportBatch "
        "SET LoadedRowCount=?, ExpectedRowCount=?, Status='Loaded', "
        "    CompletedAt=SYSUTCDATETIME() WHERE ImportBatchId=?",
        total, total, batch_id,
    )
    conn.commit()
    print(f"[segment] DEAN batch {batch_id}: {total:,} total segment rows loaded")
    print(f"[segment] unbuilt (retained gaps, NOT loaded): {sorted(UNBUILT_SEGMENTS)}")
    counts["_batch_id"] = batch_id
    counts["_total"] = total
    return counts


def _fybroc_counts(cur):
    return {
        "SegmentCombinationImport": cur.execute(
            "SELECT COUNT(*) FROM stg.SegmentCombinationImport WHERE FamilyCode<>?",
            FAMILY).fetchone()[0],
        "PumpModelReference": cur.execute(
            "SELECT COUNT(*) FROM cfg.PumpModelReference WHERE PumpFamilyId<>"
            "(SELECT PumpFamilyId FROM cfg.PumpFamily WHERE FamilyCode=?)",
            FAMILY).fetchone()[0],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true",
                    help="parse + count only; no DB writes")
    args = ap.parse_args()

    if not WB_PATH.exists():
        raise SystemExit(f"Workbook not found: {WB_PATH}")

    print(f"Loading workbook {WB_PATH} (read_only) ...", flush=True)
    wb = load_workbook(WB_PATH, read_only=True, data_only=True, keep_vba=False)

    conn = None if args.dry_run else pyodbc.connect(CONN, autocommit=False)
    pre = None
    try:
        if conn is not None:
            pre = _fybroc_counts(conn.cursor())
            print("pre non-DEAN:", pre)
        n_models = load_models(conn, wb, args.dry_run)
        seg_counts = load_segments(conn, wb, args.dry_run)
        if conn is not None:
            post = _fybroc_counts(conn.cursor())
            print("post non-DEAN:", post)
            assert pre == post, f"FYBROC rows changed! pre={pre} post={post}"
            print("FYBROC unchanged.")
    finally:
        wb.close()
        if conn is not None:
            conn.close()

    print("\n=== SUMMARY ===")
    print(f"models (A#/D#): {n_models}")
    for k, v in seg_counts.items():
        if not k.startswith("_"):
            print(f"  {k}: {v:,}")
    if "_total" in seg_counts:
        print(f"  TOTAL segment rows: {seg_counts['_total']:,} "
              f"(batch {seg_counts['_batch_id']})")
    print("DONE")


if __name__ == "__main__":
    main()
