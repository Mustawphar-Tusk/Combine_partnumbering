"""F110.8b - Fybroc Nomenclature V6 Options Vertical Compiler.

Sheet: Options - Vertical
4 fields, 96 combinations, hex-code col H(8), ID col I(9).

Fields:
  D: Coupling Option
  E: Coupling Guard
  F: Mounting Plate Option
  G: Customer Nameplate

Inputs:  workbooks/Fybroc/Nomenclature_V6.xlsm
Outputs: docs/evidence/F110/FYBROC_V6_OPTIONS_VERTICAL.{json,txt}
"""
from __future__ import annotations
import argparse, json, subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl required") from exc

STEP="F110.8b"; ROADMAP_VERSION="1.0"; MILESTONE="F110"
WORKBOOK_REL="workbooks/Fybroc/Nomenclature_V6.xlsm"
SHEET="Options - Vertical"
DISPLAY_HEADER_ROW=2; DISPLAY_DATA_START=3; DISPLAY_DATA_END=6
DATA_END_ROW=107
COL_KEY=3; COL_COUPLING_OPT=4; COL_COUPLING_GRD=5
COL_MOUNT_PLATE=6; COL_NAMEPLATE=7
COL_HEX=8; COL_ID=9
DATA_START_ROW=12  # row 11 is likely header

FIELD_COLS=[(COL_COUPLING_OPT,"Coupling Option"),(COL_COUPLING_GRD,"Coupling Guard"),
            (COL_MOUNT_PLATE,"Mounting Plate Option"),(COL_NAMEPLATE,"Customer Nameplate")]

def git_info(repo_root):
    try:
        c=subprocess.run(["git","rev-parse","HEAD"],cwd=repo_root,capture_output=True,text=True,check=True).stdout.strip()
        s=subprocess.run(["git","status","--porcelain"],cwd=repo_root,capture_output=True,text=True,check=True).stdout
        return c,(s.strip()=="")
    except: return "unknown",False

def _s(v):
    if v is None: return None
    s=str(v).strip(); return s if s else None

def build_model(repo_root):
    commit,clean=git_info(repo_root)
    wb=openpyxl.load_workbook(str(repo_root/WORKBOOK_REL),read_only=True,data_only=True)
    ws=wb[SHEET]

    # Display options
    fields={}
    for c in range(4,8):
        v=_s(ws.cell(row=DISPLAY_HEADER_ROW,column=c).value)
        if v: fields[c]=v
    options={name:[] for name in fields.values()}
    for r in range(DISPLAY_DATA_START,DISPLAY_DATA_END+1):
        for col,fname in fields.items():
            v=_s(ws.cell(row=r,column=col).value)
            if v and v!="-": options[fname].append(v)

    # Find data start
    data_start = DATA_START_ROW

    fv={name:set() for _,name in FIELD_COLS}
    hexcodes=set(); rows=[]; total=0; BASE=COL_KEY
    for row_tuple in ws.iter_rows(min_row=data_start,max_row=DATA_END_ROW,
                                   min_col=COL_KEY,max_col=COL_ID,values_only=True):
        hx=_s(row_tuple[COL_HEX-BASE])
        if hx is None or hx in ("Hex-Code","Base-36 Code"): break
        total+=1; hexcodes.add(hx)
        rd={}
        for col,name in FIELD_COLS:
            v=_s(row_tuple[col-BASE]); rd[name]=v
            if v: fv[name].add(v)
        rows.append({"id":row_tuple[COL_ID-BASE],"hex_code":hx,**rd})

    wb.close()
    sh=sorted(hexcodes)
    return {
        "step":STEP,"roadmap_version":ROADMAP_VERSION,"milestone":MILESTONE,
        "source_workbook":WORKBOOK_REL,"sheet":SHEET,
        "generated_utc":datetime.now(timezone.utc).isoformat(),
        "git_commit":commit,"git_working_tree_clean":clean,
        "fields":[n for _,n in FIELD_COLS],
        "field_options":options,
        "combination_matrix":{
            "total_combinations":total,"distinct_hex_codes":len(hexcodes),
            "hex_code_range":{"first":sh[0] if sh else None,"last":sh[-1] if sh else None},
            "distinct_values_per_field":{n:sorted(v) for n,v in fv.items()},
            "all_rows":rows
        },
        "lookup_key_col":"C","hex_code_col":"H",
        "vs_horizontal":{"vertical_fields":4,"horizontal_fields":6,
            "vertical_only":["Mounting Plate Option"],
            "horizontal_only":["Baseplate Option","Baseplate Hardware","C-Face Adapter"],
            "shared":["Coupling Option","Coupling Guard","Customer Nameplate"]},
        "part_number_role":"Options segment in vertical part number. VLOOKUP(concat_key,col_C,col_H)."
    }

def write_outputs(evidence_dir,result):
    payload={"artifact":"FYBROC_V6_OPTIONS_VERTICAL",**result}
    (evidence_dir/"FYBROC_V6_OPTIONS_VERTICAL.json").write_text(json.dumps(payload,indent=2,ensure_ascii=False,default=str),encoding="utf-8")
    cm=result["combination_matrix"]; line="="*120
    out=[f"{line}\r\nF110.8b - FYBROC V6 OPTIONS VERTICAL\r\n{line}\r\n\r\n"]
    out.append(f"Git commit  : {result['git_commit']}\r\nGit clean   : {result['git_working_tree_clean']}\r\n\r\n")
    out.append(f"Fields             : {len(result['fields'])}\r\nTotal combinations : {cm['total_combinations']}\r\nDistinct hex-codes : {cm['distinct_hex_codes']}\r\nHex range          : {cm['hex_code_range']['first']} -> {cm['hex_code_range']['last']}\r\n\r\n")
    out.append(f"{line}\r\nFIELD OPTIONS\r\n{line}\r\n\r\n")
    for f,opts in result["field_options"].items():
        out.append(f"  {f}:\r\n")
        for o in opts: out.append(f"    - {o}\r\n")
        out.append("\r\n")
    out.append(f"{line}\r\nVERTICAL vs HORIZONTAL\r\n{line}\r\n\r\n")
    v=result["vs_horizontal"]
    out.append(f"  Vertical-only: {v['vertical_only']}\r\n  Horizontal-only: {v['horizontal_only']}\r\n  Shared: {v['shared']}\r\n\r\n")
    out.append(f"{line}\r\nALL COMBINATIONS\r\n{line}\r\n\r\n")
    for r in cm["all_rows"]:
        out.append(f"  ID={r['id']:>3}  hex={r['hex_code']}  "+"  ".join(f"{n}={r.get(n,'—')}" for _,n in FIELD_COLS)+"\r\n")
    (evidence_dir/"FYBROC_V6_OPTIONS_VERTICAL.txt").write_text("".join(out),encoding="utf-8")

def main():
    p=argparse.ArgumentParser(); root=Path(__file__).resolve().parent.parent
    p.add_argument("--repo-root",type=Path,default=root); p.add_argument("--evidence-dir",type=Path,default=None)
    a=p.parse_args(); repo_root=a.repo_root.resolve()
    ev=(a.evidence_dir or (repo_root/"docs"/"evidence"/"F110")).resolve(); ev.mkdir(parents=True,exist_ok=True)
    result=build_model(repo_root); write_outputs(ev,result)
    cm=result["combination_matrix"]
    print(json.dumps({"step":STEP,"output_dir":str(ev),"fields":result["fields"],
        "total_combinations":cm["total_combinations"],"hex_code_range":cm["hex_code_range"]},indent=2))
    return 0

if __name__=="__main__": raise SystemExit(main())
