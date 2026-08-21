"""F120.2 - Fybroc Rev0.3 Selections Compiler.

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F120
("Fybroc Rev0.3 Configuration Model") - "Selections" from the eight
required sources.

Verified structure (not assumed): row 1 has Question / Answers / then
one column per series (1500, 1530, 1600, 1630, 2530, 3000, 5500, 5530,
7500, 8500 - confirmed all 10, not just the roadmap's 7). Each data row
is one (Question, Answer) pair with an 'X' in each series column where
that answer is valid for that series. 56 distinct Questions, 677 answer
rows total (rows 2-678).

This compiles the full per-question, per-answer, per-series validity
table - the option-domain half of the configuration model, parallel to
F120.1's Items (which covers field *applicability* per series, not the
valid *values* within each field).

No workbook is opened in write mode. Nothing is written back.

Inputs:
  workbooks/Fybroc/Fybroc Configuration Rev0.3.xlsx  (read-only)

Outputs:
  docs/evidence/F120/FYBROC_SELECTIONS_MODEL.{json,txt}
"""

from __future__ import annotations

import argparse
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required (it's already a pinned project dependency - "
        "run this with the project's own .venv interpreter)."
    ) from exc

STEP = "F120.2"
ROADMAP_VERSION = "1.0"
MILESTONE = "F120"
WORKBOOK_REL = "workbooks/Fybroc/Fybroc Configuration Rev0.3.xlsx"

HEADER_ROW = 1
DATA_START_ROW = 2
DATA_END_ROW = 678  # verified: last populated row
SERIES_COLS_START = 4
SERIES_COLS_END = 13  # verified: 1500..8500, 10 series


def git_info(repo_root: Path) -> tuple[str, bool]:
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout.strip()
        status = subprocess.run(
            ["git", "status", "--porcelain"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout
        return commit, (status.strip() == "")
    except Exception:
        return "unknown", False


def compile_selections(ws) -> dict[str, Any]:
    series_codes = {}
    for c in range(SERIES_COLS_START, SERIES_COLS_END + 1):
        v = ws.cell(row=HEADER_ROW, column=c).value
        if v is not None:
            series_codes[c] = str(v).strip()

    rows = []
    questions_seen: dict[str, int] = {}
    for row in range(DATA_START_ROW, DATA_END_ROW + 1):
        question = ws.cell(row=row, column=2).value
        answer = ws.cell(row=row, column=3).value
        if question is None or answer is None:
            continue
        question = str(question).strip()
        answer = str(answer).strip()
        questions_seen[question] = questions_seen.get(question, 0) + 1

        valid_series = [
            series_codes[c] for c in range(SERIES_COLS_START, SERIES_COLS_END + 1)
            if ws.cell(row=row, column=c).value == "X"
        ]
        rows.append({
            "question": question,
            "answer": answer,
            "valid_series": valid_series,
            "valid_series_count": len(valid_series),
        })

    return {
        "series_columns": list(series_codes.values()),
        "question_count": len(questions_seen),
        "questions": sorted(questions_seen.keys()),
        "answer_row_count": len(rows),
        "rows": rows,
    }


def build_model(repo_root: Path) -> dict[str, Any]:
    commit, clean = git_info(repo_root)
    wb = openpyxl.load_workbook(str(repo_root / WORKBOOK_REL), read_only=True, data_only=True)
    selections = compile_selections(wb["Selections"])
    wb.close()

    return {
        "step": STEP, "roadmap_version": ROADMAP_VERSION, "milestone": MILESTONE,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit, "git_working_tree_clean": clean,
        "selections": selections,
    }


def _banner(title: str) -> str:
    line = "=" * 140
    return f"{line}\r\n{title}\r\n{line}\r\n\r\n"


def write_outputs(evidence_dir: Path, result: dict[str, Any]) -> None:
    payload = {"artifact": "FYBROC_SELECTIONS_MODEL", **result}
    (evidence_dir / "FYBROC_SELECTIONS_MODEL.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    sel = result["selections"]
    lines = [_banner("F120.2 - FYBROC SELECTIONS MODEL (per-question, per-answer, per-series validity)")]
    lines.append(
        f"Git commit  : {result['git_commit']}\r\nGit clean   : {result['git_working_tree_clean']}\r\n\r\n"
        f"Series columns : {', '.join(sel['series_columns'])}\r\n"
        f"Questions      : {sel['question_count']}\r\n"
        f"Answer rows    : {sel['answer_row_count']}\r\n\r\n"
    )

    # Group rows by question for a readable report.
    by_question: dict[str, list[dict[str, Any]]] = {}
    for r in sel["rows"]:
        by_question.setdefault(r["question"], []).append(r)

    for question in sel["questions"]:
        rows_for_q = by_question.get(question, [])
        lines.append(_banner(f"QUESTION: {question}  ({len(rows_for_q)} answers)"))
        for r in rows_for_q:
            series_str = ", ".join(r["valid_series"]) if r["valid_series"] else "(none)"
            lines.append(f"  {r['answer']!r:<40} valid for: {series_str}\r\n")
        lines.append("\r\n")

    (evidence_dir / "FYBROC_SELECTIONS_MODEL.txt").write_text("".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="F120.2 Fybroc Selections compiler")
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument("--repo-root", type=Path, default=default_root)
    parser.add_argument("--evidence-dir", type=Path, default=None)
    args = parser.parse_args()

    repo_root: Path = args.repo_root.resolve()
    evidence_dir: Path = (args.evidence_dir or (repo_root / "docs" / "evidence" / "F120")).resolve()
    evidence_dir.mkdir(parents=True, exist_ok=True)

    result = build_model(repo_root)
    write_outputs(evidence_dir, result)

    print(json.dumps({
        "step": STEP, "output_dir": str(evidence_dir),
        "series_columns": result["selections"]["series_columns"],
        "question_count": result["selections"]["question_count"],
        "answer_row_count": result["selections"]["answer_row_count"],
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())