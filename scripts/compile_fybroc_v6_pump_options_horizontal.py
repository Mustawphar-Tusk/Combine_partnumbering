"""F110.5a - Fybroc Nomenclature V6 Pump Options Horizontal Compiler.

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F110
("V5 to V6 Nomenclature Reconciliation") - combination rules, source
fields, orientation behavior.

The "Pump Options - Horizontal" sheet in Nomenclature_V6.xlsm defines
the full combination matrix for all horizontal pump option fields. It
has two sections:

DISPLAY SECTION (rows 1-9, cols C-O):
  Shows each field name and its valid option values:
    Col D : Casing Drains
    Col E : Suction Discharge Taps
    Col F : Shaft Material
    Col G : Sleeve
    Col H : Casing Hardware
    Col I : Pump Elastomers
    Col J : Bearing Option
    Col K : Power Frame Hardware
    Col L : Gland Hardware
    Col M : Flush
    Col N : Cyclone Separator
    Col O : Impeller Balance (Dynamic Balancing)

DATA SECTION (rows 11-138251, cols C-T):
  Row 11  = column headers
  Rows 12+= one row per unique combination of all 12 fields
    Col C  : concatenated key (all values joined, used as lookup key)
    Col D  : Casing Drains
    Col E  : Suction Discharge
    Col F  : Shaft Material
    Col G  : Sleeve
    Col H  : Casing Hardware
    Col I  : Pump Elastomers
    Col J  : Bearing Option
    Col K  : Frame Hardware
    Col L  : Gland Hardware
    Col M  : Flush
    Col N  : Cyclone Separator
    Col O  : Dynamic Balancing
    Col P  : Hex-Code (the part number token for this combination)
    Col Q  : ID (sequential integer)
    Col S  : Shaft Material label (derived, e.g. "303 SHAFT")
    Col T  : Flush abbreviation (derived, e.g. "EXT", "INT", "BYP")

Total rows: 138,240 combinations (rows 12-138251).

The hex-code in col P is what gets embedded in the part number.
The Smart Number sheet uses a VLOOKUP against col C (concatenated key)
to retrieve col P (hex-code) for the Pump Options segment.

No workbook is opened in write mode. Nothing is written back.

Inputs:
  workbooks/Fybroc/Nomenclature_V6.xlsm  (read-only)

Outputs:
  docs/evidence/F110/FYBROC_V6_PUMP_OPTIONS_HORIZONTAL.{json,txt}
"""

from __future__ import annotations

import argparse
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required - run with the project's own .venv interpreter."
    ) from exc

STEP = "F110.5a"
ROADMAP_VERSION = "1.0"
MILESTONE = "F110"
WORKBOOK_REL = "workbooks/Fybroc/Nomenclature_V6.xlsm"
SHEET = "Pump Options - Horizontal"

# Display section
DISPLAY_HEADER_ROW  = 2
DISPLAY_DATA_START  = 3
DISPLAY_DATA_END    = 8
DISPLAY_COL_START   = 4   # col D
DISPLAY_COL_END     = 15  # col O

# Data section
DATA_HEADER_ROW = 11
DATA_START_ROW  = 12
DATA_END_ROW    = 138251

# Column indices (1-based)
COL_KEY         = 3   # C - concatenated lookup key
COL_CASING_DRAINS    = 4   # D
COL_SUCTION_DISC     = 5   # E
COL_SHAFT_MATERIAL   = 6   # F
COL_SLEEVE           = 7   # G
COL_CASING_HW        = 8   # H
COL_PUMP_ELASTOMERS  = 9   # I
COL_BEARING_OPTION   = 10  # J
COL_FRAME_HW         = 11  # K
COL_GLAND_HW         = 12  # L
COL_FLUSH            = 13  # M
COL_CYCLONE_SEP      = 14  # N
COL_DYNAMIC_BAL      = 15  # O
COL_HEX_CODE         = 16  # P
COL_ID               = 17  # Q
COL_SHAFT_LABEL      = 19  # S
COL_FLUSH_ABBREV     = 20  # T

FIELD_COLS = [
    (COL_CASING_DRAINS,   "Casing Drains"),
    (COL_SUCTION_DISC,    "Suction Discharge"),
    (COL_SHAFT_MATERIAL,  "Shaft Material"),
    (COL_SLEEVE,          "Sleeve"),
    (COL_CASING_HW,       "Casing Hardware"),
    (COL_PUMP_ELASTOMERS, "Pump Elastomers"),
    (COL_BEARING_OPTION,  "Bearing Option"),
    (COL_FRAME_HW,        "Frame Hardware"),
    (COL_GLAND_HW,        "Gland Hardware"),
    (COL_FLUSH,           "Flush"),
    (COL_CYCLONE_SEP,     "Cyclone Separator"),
    (COL_DYNAMIC_BAL,     "Dynamic Balancing"),
]


def git_info(repo_root: Path) -> tuple[str, bool]:
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout.strip()
        status = subprocess.run(
            ["git", "status", "--porcelain"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout
        return commit, (status.strip() == "")
    except Exception:
        return "unknown", False


def _str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def compile_display_section(ws) -> dict[str, list[str]]:
    """Extract the field -> valid options map from the display section."""
    # Row 2 = field names in cols D-O
    fields: dict[int, str] = {}
    for c in range(DISPLAY_COL_START, DISPLAY_COL_END + 1):
        v = _str(ws.cell(row=DISPLAY_HEADER_ROW, column=c).value)
        if v:
            fields[c] = v

    options: dict[str, list[str]] = {name: [] for name in fields.values()}
    for r in range(DISPLAY_DATA_START, DISPLAY_DATA_END + 1):
        for col, field_name in fields.items():
            v = _str(ws.cell(row=r, column=col).value)
            if v:
                options[field_name].append(v)

    return options


def compile_data_section(ws) -> dict[str, Any]:
    """
    Compile statistics and a sample from the full combination matrix using
    iter_rows for performance (138,240 rows).
    Captures: total count, distinct values per field, distinct hex-codes,
    first 20 sample rows, hex-code range.
    """
    # Column offsets relative to min_col=3 (col C = index 0)
    # C=0, D=1, E=2, F=3, G=4, H=5, I=6, J=7, K=8, L=9, M=10, N=11, O=12
    # P=13(hex), Q=14(id), S=16(shaft_label), T=17(flush_abbrev)
    BASE = COL_KEY  # = 3

    field_values: dict[str, set] = {name: set() for _, name in FIELD_COLS}
    hex_codes: set[str] = set()
    sample_rows: list[dict] = []
    total = 0

    for row_tuple in ws.iter_rows(
        min_row=DATA_START_ROW, max_row=DATA_END_ROW,
        min_col=COL_KEY, max_col=COL_FLUSH_ABBREV,
        values_only=True,
    ):
        # row_tuple index: 0=key, 1=casing_drains .. 12=dynamic_bal,
        #                  13=hex_code, 14=id, 15=blank(R), 16=shaft_label, 17=flush_abbrev
        hex_code = _str(row_tuple[COL_HEX_CODE - BASE])
        if hex_code is None:
            break

        total += 1
        hex_codes.add(hex_code)

        row_data: dict[str, Any] = {}
        for col, name in FIELD_COLS:
            v = _str(row_tuple[col - BASE])
            row_data[name] = v
            if v:
                field_values[name].add(v)

        if total <= 20:
            sample_rows.append({
                "id": row_tuple[COL_ID - BASE],
                "hex_code": hex_code,
                "shaft_label": _str(row_tuple[COL_SHAFT_LABEL - BASE]),
                "flush_abbrev": _str(row_tuple[COL_FLUSH_ABBREV - BASE]),
                **row_data,
            })

    sorted_hex = sorted(hex_codes)

    return {
        "total_combinations": total,
        "distinct_hex_codes": len(hex_codes),
        "hex_code_range": {
            "first": sorted_hex[0] if sorted_hex else None,
            "last": sorted_hex[-1] if sorted_hex else None,
        },
        "distinct_values_per_field": {
            name: sorted(vals) for name, vals in field_values.items()
        },
        "sample_rows": sample_rows,
    }


def build_model(repo_root: Path) -> dict[str, Any]:
    commit, clean = git_info(repo_root)

    wb = openpyxl.load_workbook(
        str(repo_root / WORKBOOK_REL), read_only=True, data_only=True
    )
    ws = wb[SHEET]

    display = compile_display_section(ws)
    data    = compile_data_section(ws)

    wb.close()

    return {
        "step": STEP,
        "roadmap_version": ROADMAP_VERSION,
        "milestone": MILESTONE,
        "source_workbook": WORKBOOK_REL,
        "sheet": SHEET,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit,
        "git_working_tree_clean": clean,
        "fields": list(display.keys()),
        "field_options": display,
        "combination_matrix": data,
        "lookup_key_col": "C (concatenated all field values)",
        "hex_code_col": "P",
        "part_number_role": (
            "The hex-code in col P is the Pump Options segment token "
            "embedded in the Fybroc horizontal part number. "
            "Smart Number uses VLOOKUP(concat_key, col_C, col_P) to "
            "retrieve it."
        ),
    }


def _banner(title: str) -> str:
    line = "=" * 120
    return f"{line}\r\n{title}\r\n{line}\r\n\r\n"


def write_outputs(evidence_dir: Path, result: dict[str, Any]) -> None:
    payload = {"artifact": "FYBROC_V6_PUMP_OPTIONS_HORIZONTAL", **result}
    (evidence_dir / "FYBROC_V6_PUMP_OPTIONS_HORIZONTAL.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )

    cm = result["combination_matrix"]
    lines = [_banner("F110.5a - FYBROC V6 PUMP OPTIONS HORIZONTAL")]
    lines.append(
        f"Git commit           : {result['git_commit']}\r\n"
        f"Git clean            : {result['git_working_tree_clean']}\r\n\r\n"
        f"Fields               : {len(result['fields'])}\r\n"
        f"Total combinations   : {cm['total_combinations']:,}\r\n"
        f"Distinct hex-codes   : {cm['distinct_hex_codes']:,}\r\n"
        f"Hex-code range       : {cm['hex_code_range']['first']} -> {cm['hex_code_range']['last']}\r\n"
        f"Lookup key col       : {result['lookup_key_col']}\r\n"
        f"Hex-code col         : {result['hex_code_col']}\r\n\r\n"
    )

    lines.append(_banner("FIELD OPTIONS"))
    for field, opts in result["field_options"].items():
        lines.append(f"  {field}:\r\n")
        for opt in opts:
            lines.append(f"    - {opt}\r\n")
        lines.append("\r\n")

    lines.append(_banner("DISTINCT VALUES PER FIELD (from full matrix)"))
    for field, vals in cm["distinct_values_per_field"].items():
        lines.append(f"  {field} ({len(vals)}):\r\n")
        for v in vals:
            lines.append(f"    {v}\r\n")
        lines.append("\r\n")

    lines.append(_banner("SAMPLE ROWS (first 20)"))
    for row in cm["sample_rows"]:
        lines.append(f"  ID={row['id']:>6}  hex={row['hex_code']}  shaft={row['shaft_label']}  flush={row['flush_abbrev']}\r\n")
        for field, _ in FIELD_COLS:
            fname = [n for c, n in FIELD_COLS if c == field][0]
            lines.append(f"    {fname:<25} = {row.get(fname) or '—'}\r\n")
        lines.append("\r\n")

    (evidence_dir / "FYBROC_V6_PUMP_OPTIONS_HORIZONTAL.txt").write_text(
        "".join(lines), encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="F110.5a Fybroc V6 Pump Options Horizontal compiler")
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument("--repo-root", type=Path, default=default_root)
    parser.add_argument("--evidence-dir", type=Path, default=None)
    args = parser.parse_args()

    repo_root: Path = args.repo_root.resolve()
    evidence_dir: Path = (
        args.evidence_dir or (repo_root / "docs" / "evidence" / "F110")
    ).resolve()
    evidence_dir.mkdir(parents=True, exist_ok=True)

    result = build_model(repo_root)
    write_outputs(evidence_dir, result)

    cm = result["combination_matrix"]
    print(json.dumps({
        "step": STEP,
        "output_dir": str(evidence_dir),
        "fields": result["fields"],
        "total_combinations": cm["total_combinations"],
        "distinct_hex_codes": cm["distinct_hex_codes"],
        "hex_code_range": cm["hex_code_range"],
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
