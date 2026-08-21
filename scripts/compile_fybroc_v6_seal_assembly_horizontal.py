"""F110.6 - Fybroc Nomenclature V6 Seal Assembly Horizontal Compiler.

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F110
("V5 to V6 Nomenclature Reconciliation") - combination rules, source
fields, orientation behavior.

The "Seal Assembly - Horizontal" sheet in Nomenclature_V6.xlsm defines
the combination matrix for all horizontal seal assembly fields.

DISPLAY SECTION (rows 1-8, cols C-N):
  Row 2 = field names, rows 3-8 = valid options:
    Col D : Seal Option
    Col E : Seal Type
    Col F : Seal Materials
    Col G : Seal Elastomers
    Col H : Seal Guard
    Col M : Seal Manufacturer  (separate lookup: name -> code)
    Col N : Seal Manufacturer Code

  Seal Manufacturer codes:
    Standard Offering -> S
    Flexaseal         -> F
    John Crane        -> J
    Custom            -> C

DATA SECTION (rows 13-80, cols C-N):
  Row 13  = column headers
  Rows 14+= one row per unique combination of the 5 seal fields
    Col C  : concatenated key (VLOOKUP key)
    Col D  : Seal Option
    Col E  : Seal Type
    Col F  : Seal Materials
    Col G  : Seal Elastomers
    Col H  : Seal Guard
    Col I  : Hex-Code (2-digit part number token for this combination)
    Col J  : ID (sequential integer)
    Col M  : Seal Type abbreviation (derived, e.g. "8B2", "RAC")
    Col N  : Seal Materials abbreviation (derived, e.g. "C/CR", "S/S")

Total rows: ~67 combinations (rows 14-80).

The hex-code (col I) is a 2-digit code embedded in the Seal Assembly
segment of the part number. Smart Number concatenates Seal Manufacturer
code + Seal Assembly hex-code to form the full seal segment.

No workbook is opened in write mode. Nothing is written back.

Inputs:
  workbooks/Fybroc/Nomenclature_V6.xlsm  (read-only)

Outputs:
  docs/evidence/F110/FYBROC_V6_SEAL_ASSEMBLY_HORIZONTAL.{json,txt}
"""

from __future__ import annotations

import argparse
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required - run with the project's own .venv interpreter."
    ) from exc

STEP = "F110.6"
ROADMAP_VERSION = "1.0"
MILESTONE = "F110"
WORKBOOK_REL = "workbooks/Fybroc/Nomenclature_V6.xlsm"
SHEET = "Seal Assembly - Horizontal"

DISPLAY_HEADER_ROW = 2
DISPLAY_DATA_START = 3
DISPLAY_DATA_END   = 8
DISPLAY_COL_START  = 4   # D
DISPLAY_COL_END    = 8   # H  (seal fields only)

DATA_HEADER_ROW = 13
DATA_START_ROW  = 14
DATA_END_ROW    = 80

# Column indices (1-based)
COL_KEY           = 3   # C
COL_SEAL_OPTION   = 4   # D
COL_SEAL_TYPE     = 5   # E
COL_SEAL_MATS     = 6   # F
COL_SEAL_ELAST    = 7   # G
COL_SEAL_GUARD    = 8   # H
COL_HEX_CODE      = 9   # I
COL_ID            = 10  # J
COL_SEAL_TYPE_ABB = 13  # M  derived
COL_SEAL_MATS_ABB = 14  # N  derived

# Seal manufacturer display cols
COL_MFR_NAME = 13
COL_MFR_CODE = 14

FIELD_COLS = [
    (COL_SEAL_OPTION, "Seal Option"),
    (COL_SEAL_TYPE,   "Seal Type"),
    (COL_SEAL_MATS,   "Seal Materials"),
    (COL_SEAL_ELAST,  "Seal Elastomers"),
    (COL_SEAL_GUARD,  "Seal Guard"),
]


def git_info(repo_root: Path) -> tuple[str, bool]:
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout.strip()
        status = subprocess.run(
            ["git", "status", "--porcelain"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout
        return commit, (status.strip() == "")
    except Exception:
        return "unknown", False


def _str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def compile_display_section(ws) -> tuple[dict[str, list[str]], list[dict]]:
    """Returns (field_options, seal_manufacturer_codes)."""
    fields: dict[int, str] = {}
    for c in range(DISPLAY_COL_START, DISPLAY_COL_END + 1):
        v = _str(ws.cell(row=DISPLAY_HEADER_ROW, column=c).value)
        if v:
            fields[c] = v

    options: dict[str, list[str]] = {name: [] for name in fields.values()}
    mfr_codes: list[dict] = []

    for r in range(DISPLAY_DATA_START, DISPLAY_DATA_END + 1):
        for col, field_name in fields.items():
            v = _str(ws.cell(row=r, column=col).value)
            if v and v != "-":
                options[field_name].append(v)
        # Seal manufacturer name + code in display cols M(13) and N(14)
        mfr_name = _str(ws.cell(row=r, column=COL_MFR_NAME).value)
        mfr_code = _str(ws.cell(row=r, column=COL_MFR_CODE).value)
        if mfr_name and mfr_code:
            mfr_codes.append({"manufacturer": mfr_name, "code": mfr_code})

    return options, mfr_codes


def compile_data_section(ws) -> dict[str, Any]:
    BASE = COL_KEY

    field_values: dict[str, set] = {name: set() for _, name in FIELD_COLS}
    hex_codes: set[str] = set()
    rows: list[dict] = []
    total = 0

    for row_tuple in ws.iter_rows(
        min_row=DATA_START_ROW, max_row=DATA_END_ROW,
        min_col=COL_KEY, max_col=COL_SEAL_MATS_ABB,
        values_only=True,
    ):
        hex_code = _str(row_tuple[COL_HEX_CODE - BASE])
        if hex_code is None:
            break

        total += 1
        hex_codes.add(hex_code)

        row_data: dict[str, Any] = {}
        for col, name in FIELD_COLS:
            v = _str(row_tuple[col - BASE])
            row_data[name] = v
            if v:
                field_values[name].add(v)

        rows.append({
            "id": row_tuple[COL_ID - BASE],
            "hex_code": hex_code,
            "seal_type_abbrev": _str(row_tuple[COL_SEAL_TYPE_ABB - BASE]),
            "seal_materials_abbrev": _str(row_tuple[COL_SEAL_MATS_ABB - BASE]),
            **row_data,
        })

    sorted_hex = sorted(hex_codes)
    return {
        "total_combinations": total,
        "distinct_hex_codes": len(hex_codes),
        "hex_code_range": {
            "first": sorted_hex[0] if sorted_hex else None,
            "last": sorted_hex[-1] if sorted_hex else None,
        },
        "distinct_values_per_field": {
            name: sorted(vals) for name, vals in field_values.items()
        },
        "all_rows": rows,
    }


def build_model(repo_root: Path) -> dict[str, Any]:
    commit, clean = git_info(repo_root)

    wb = openpyxl.load_workbook(
        str(repo_root / WORKBOOK_REL), read_only=True, data_only=True
    )
    ws = wb[SHEET]

    display, mfr_codes = compile_display_section(ws)
    data = compile_data_section(ws)

    wb.close()

    return {
        "step": STEP,
        "roadmap_version": ROADMAP_VERSION,
        "milestone": MILESTONE,
        "source_workbook": WORKBOOK_REL,
        "sheet": SHEET,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit,
        "git_working_tree_clean": clean,
        "fields": [name for _, name in FIELD_COLS],
        "field_options": display,
        "seal_manufacturer_codes": mfr_codes,
        "combination_matrix": data,
        "lookup_key_col": "C (concatenated all 5 field values)",
        "hex_code_col": "I (2-digit code)",
        "part_number_role": (
            "Seal Assembly segment = Seal Manufacturer code (1 char, from Attributes "
            "or separate mfr lookup) + Seal Assembly hex-code (2 digits from col I). "
            "Smart Number uses VLOOKUP(concat_key, col_C, col_I) for the 2-digit token."
        ),
    }


def _banner(title: str) -> str:
    line = "=" * 120
    return f"{line}\r\n{title}\r\n{line}\r\n\r\n"


def write_outputs(evidence_dir: Path, result: dict[str, Any]) -> None:
    payload = {"artifact": "FYBROC_V6_SEAL_ASSEMBLY_HORIZONTAL", **result}
    (evidence_dir / "FYBROC_V6_SEAL_ASSEMBLY_HORIZONTAL.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )

    cm = result["combination_matrix"]
    lines = [_banner("F110.6 - FYBROC V6 SEAL ASSEMBLY HORIZONTAL")]
    lines.append(
        f"Git commit           : {result['git_commit']}\r\n"
        f"Git clean            : {result['git_working_tree_clean']}\r\n\r\n"
        f"Fields               : {len(result['fields'])}\r\n"
        f"Total combinations   : {cm['total_combinations']}\r\n"
        f"Distinct hex-codes   : {cm['distinct_hex_codes']}\r\n"
        f"Hex-code range       : {cm['hex_code_range']['first']} -> {cm['hex_code_range']['last']}\r\n"
        f"Lookup key col       : {result['lookup_key_col']}\r\n"
        f"Hex-code col         : {result['hex_code_col']}\r\n\r\n"
    )

    lines.append(_banner("SEAL MANUFACTURER CODES"))
    for m in result["seal_manufacturer_codes"]:
        lines.append(f"  {m['manufacturer']:<25} -> {m['code']}\r\n")
    lines.append("\r\n")

    lines.append(_banner("FIELD OPTIONS"))
    for field, opts in result["field_options"].items():
        lines.append(f"  {field}:\r\n")
        for opt in opts:
            lines.append(f"    - {opt}\r\n")
        lines.append("\r\n")

    lines.append(_banner("DISTINCT VALUES PER FIELD (from full matrix)"))
    for field, vals in cm["distinct_values_per_field"].items():
        lines.append(f"  {field} ({len(vals)}):\r\n")
        for v in vals:
            lines.append(f"    {v}\r\n")
        lines.append("\r\n")

    lines.append(_banner("ALL COMBINATIONS"))
    for row in cm["all_rows"]:
        lines.append(
            f"  ID={row['id']:>3}  hex={row['hex_code']}  "
            f"type_abbrev={row.get('seal_type_abbrev') or '—':<8}  "
            f"mats_abbrev={row.get('seal_materials_abbrev') or '—'}\r\n"
        )
        for _, fname in FIELD_COLS:
            lines.append(f"    {fname:<22} = {row.get(fname) or '—'}\r\n")
        lines.append("\r\n")

    (evidence_dir / "FYBROC_V6_SEAL_ASSEMBLY_HORIZONTAL.txt").write_text(
        "".join(lines), encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="F110.6 Fybroc V6 Seal Assembly Horizontal compiler")
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument("--repo-root", type=Path, default=default_root)
    parser.add_argument("--evidence-dir", type=Path, default=None)
    args = parser.parse_args()

    repo_root: Path = args.repo_root.resolve()
    evidence_dir: Path = (
        args.evidence_dir or (repo_root / "docs" / "evidence" / "F110")
    ).resolve()
    evidence_dir.mkdir(parents=True, exist_ok=True)

    result = build_model(repo_root)
    write_outputs(evidence_dir, result)

    cm = result["combination_matrix"]
    print(json.dumps({
        "step": STEP,
        "output_dir": str(evidence_dir),
        "fields": result["fields"],
        "total_combinations": cm["total_combinations"],
        "distinct_hex_codes": cm["distinct_hex_codes"],
        "hex_code_range": cm["hex_code_range"],
        "seal_manufacturer_codes": result["seal_manufacturer_codes"],
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
