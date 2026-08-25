"""Verify all Fybroc series across all source workbooks."""
import openpyxl
import json

# V6 Attributes - all series defined
attrs = json.load(open("docs/evidence/F110/FYBROC_V6_ATTRIBUTES.json", encoding="utf-8"))
print("=== V6 Nomenclature Attributes - Series Defined ===")
print("Orientations:")
for s in attrs["series_orientation"]:
    print(f"  {s['series']:<8} {s['orientation']}")
print()
print("Series+Flange codes:")
for sf in attrs["series_flange_codes"]:
    print(f"  {sf['series']:<8} + {sf['flange_type']:<8} = {sf['code']}")

# Rev0.3 Selections
print()
print("=== Rev0.3 Selections - Series with configuration data ===")
wb = openpyxl.load_workbook(
    "workbooks/Fybroc/Fybroc Configuration Rev0.3.xlsx",
    read_only=True, data_only=True,
)
ws = wb["Selections"]
series_in_selections = []
for c in range(4, 20):
    v = ws.cell(row=1, column=c).value
    if v:
        series_in_selections.append(str(v).strip())
wb.close()
print(f"Series: {series_in_selections}")

# Pricebook
print()
print("=== Pricebook - Series with pricing blocks ===")
wb2 = openpyxl.load_workbook(
    "workbooks/Fybroc/Price Estimator-Fybroc.xlsm",
    read_only=True, data_only=True,
)
ws2 = wb2["Pricebook"]
for c in range(1, 200):
    v = ws2.cell(row=3, column=c).value
    if v and "Pricing" in str(v):
        print(f"  Col {c}: {v}")
wb2.close()

# Summary
print()
print("=== SUMMARY ===")
v6_series = set(s["series"] for s in attrs["series_orientation"])
rev03_series = set(series_in_selections)
print(f"V6 Attributes series:  {sorted(v6_series)}")
print(f"Rev0.3 Selections:     {sorted(rev03_series)}")
print(f"In V6 but NOT in Rev0.3: {sorted(v6_series - rev03_series)}")
print(f"In Rev0.3 but NOT in V6: {sorted(rev03_series - v6_series)}")
