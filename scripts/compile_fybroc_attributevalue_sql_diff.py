"""F120.5b - Fybroc AttributeValue (SQL) vs V6 Attributes Hex-Code Diff.

Roadmap anchor
--------------
docs/PROJECT_MASTER_ROADMAP.md, Section 6, Milestone F120
("Fybroc Rev0.3 Configuration Model") - "Compare against current SQL
metadata: cfg.AttributeValue". Contributes to FYBROC_CONFIGURATION_MODEL.

SQL's cfg.AttributeValue currently covers exactly 6 fields (confirmed
via the live snapshot): IMPELLER_TRIM, MOTOR_MODIFICATIONS,
PUMP_MATERIAL, SERIES, SIZE, TESTING. Of these:

  - IMPELLER_TRIM, MOTOR_MODIFICATIONS, PUMP_MATERIAL, SIZE: directly
    comparable. V6's Attributes sheet has clean, verified hex codes for
    all four (already proven correct in F110.2a - reusing that exact
    column-reading logic here, not re-deriving it).

  - SERIES: NOT compared. V6 does not have an equivalent single hex
    code for series+flange - see
    docs/milestones/F110_pending_engineering_decisions.md item 1.
    Comparing SQL's SERIES hex codes against "nothing" would produce a
    meaningless 100%-DEPRECATED result that looks like a real finding
    but isn't - reported as a known architectural gap instead.

  - TESTING: NOT compared. V6 uses a completely different 4-axis
    Base-36 permutation encoding (see item 2 in the same tracker doc),
    not a single hex code. Same reasoning as SERIES - not comparable
    without inventing a mapping that doesn't exist yet.

Reuses read_hex_table()/normalize_value() from
compile_fybroc_attributes_hex_diff.py exactly, rather than re-deriving
column positions - those were independently verified against the real
V6 workbook headers in that earlier step.

Inputs:
  workbooks/Fybroc/Nomenclature_V6.xlsm  (read-only, for the 4 comparable fields)
  docs/evidence/F120/FYBROC_SQL_METADATA_SNAPSHOT.json (already extracted, reused)

Outputs:
  docs/evidence/F120/FYBROC_ATTRIBUTEVALUE_SQL_DIFF.{json,txt}
"""

from __future__ import annotations

import argparse
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required (it's already a pinned project dependency - "
        "run this with the project's own .venv interpreter)."
    ) from exc

STEP = "F120.5b"
ROADMAP_VERSION = "1.0"
MILESTONE = "F120"
V6_WORKBOOK_REL = "workbooks/Fybroc/Nomenclature_V6.xlsm"

DATA_START_ROW = 8
DATA_END_ROW = 523

# Verified V6 column positions - identical to compile_fybroc_attributes_hex_diff.py.
COMPARABLE_FIELDS: dict[str, tuple[int, int]] = {
    "SIZE": (10, 11),
    "PUMP_MATERIAL": (13, 14),
    "IMPELLER_TRIM": (16, 17),
    "MOTOR_MODIFICATIONS": (19, 20),
}

NOT_COMPARABLE = {
    "SERIES": (
        "V6 has no single hex code for series+flange - it splits into a bare "
        "series number and a separate, non-hex Flange Type descriptor. "
        "Comparing against 'nothing' would produce a meaningless 100%-"
        "DEPRECATED result. See F110_pending_engineering_decisions.md item 1."
    ),
    "TESTING": (
        "V6 uses a 4-axis Base-36 permutation encoding, not a single hex "
        "code - a different paradigm, not a value remap. See "
        "F110_pending_engineering_decisions.md item 2."
    ),
}


def normalize_value(value: Any) -> str:
    """Identical to compile_fybroc_attributes_hex_diff.py's normalization -
    reused, not re-derived."""
    text = str(value).strip()
    if text.endswith("*"):
        text = text[:-1].strip()
    return text.replace("_", " ").strip().lower()


def read_v6_hex_table(ws, name_col: int, code_col: int) -> dict[str, tuple[Any, Any]]:
    result: dict[str, tuple[Any, Any]] = {}
    for row in range(DATA_START_ROW, DATA_END_ROW + 1):
        v = ws.cell(row=row, column=name_col).value
        if v is None or str(v).strip() == "":
            continue
        code = ws.cell(row=row, column=code_col).value
        result[normalize_value(v)] = (v, code)
    return result


def git_info(repo_root: Path) -> tuple[str, bool]:
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout.strip()
        status = subprocess.run(
            ["git", "status", "--porcelain"], cwd=repo_root, capture_output=True,
            text=True, check=True,
        ).stdout
        return commit, (status.strip() == "")
    except Exception:
        return "unknown", False


def build_diff(repo_root: Path, evidence_dir: Path) -> dict[str, Any]:
    commit, clean = git_info(repo_root)

    sql_path = evidence_dir / "FYBROC_SQL_METADATA_SNAPSHOT.json"
    if not sql_path.exists():
        raise SystemExit(f"Required input missing: {sql_path}")
    sql_snapshot = json.loads(sql_path.read_text(encoding="utf-8"))

    sql_by_field: dict[str, dict[str, tuple[Any, Any]]] = {}
    for r in sql_snapshot["attribute_values"]:
        sql_by_field.setdefault(r["field_code"], {})[normalize_value(r["display_value"])] = (
            r["display_value"], r["identifier_code"]
        )

    wb = openpyxl.load_workbook(str(repo_root / V6_WORKBOOK_REL), read_only=True, data_only=True)
    ws = wb["Attributes"]

    field_results = []
    total_counts = {"UNCHANGED": 0, "CHANGED": 0, "NEW": 0, "DEPRECATED": 0}

    for field_code, (name_col, code_col) in COMPARABLE_FIELDS.items():
        v6_values = read_v6_hex_table(ws, name_col, code_col)
        sql_values = sql_by_field.get(field_code, {})
        all_keys = sorted(set(v6_values) | set(sql_values))
        counts = {"UNCHANGED": 0, "CHANGED": 0, "NEW": 0, "DEPRECATED": 0}
        rows = []
        for key in all_keys:
            in_sql, in_v6 = key in sql_values, key in v6_values
            if in_sql and in_v6:
                _, sql_code = sql_values[key]
                _, v6_code = v6_values[key]
                cls = "UNCHANGED" if str(sql_code) == str(v6_code) else "CHANGED"
            elif in_v6:
                cls = "NEW"
            else:
                cls = "DEPRECATED"
            counts[cls] += 1
            total_counts[cls] += 1
            rows.append({
                "value": (v6_values.get(key) or sql_values.get(key))[0],
                "sql_identifier_code": sql_values.get(key, (None, None))[1],
                "v6_identifier_code": v6_values.get(key, (None, None))[1],
                "classification": cls,
            })
        field_results.append({"field_code": field_code, "row_count": len(rows), "counts": counts, "rows": rows})

    wb.close()

    return {
        "step": STEP, "roadmap_version": ROADMAP_VERSION, "milestone": MILESTONE,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": commit, "git_working_tree_clean": clean,
        "comparable_field_count": len(COMPARABLE_FIELDS),
        "not_comparable_fields": NOT_COMPARABLE,
        "total_classification_counts": total_counts,
        "fields": field_results,
    }


def _banner(title: str) -> str:
    line = "=" * 140
    return f"{line}\r\n{title}\r\n{line}\r\n\r\n"


def write_outputs(evidence_dir: Path, result: dict[str, Any]) -> None:
    payload = {"artifact": "FYBROC_ATTRIBUTEVALUE_SQL_DIFF", **result}
    (evidence_dir / "FYBROC_ATTRIBUTEVALUE_SQL_DIFF.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
    )

    lines = [_banner("F120.5b - FYBROC ATTRIBUTEVALUE (SQL) vs V6 ATTRIBUTES DIFF")]
    lines.append(
        f"Git commit  : {result['git_commit']}\r\nGit clean   : {result['git_working_tree_clean']}\r\n\r\n"
        f"Comparable fields : {result['comparable_field_count']}\r\n"
        f"Totals            : {result['total_classification_counts']}\r\n\r\n"
    )
    lines.append(_banner("NOT COMPARABLE - KNOWN ARCHITECTURAL GAPS, NOT SKIPPED BY OVERSIGHT"))
    for field, reason in result["not_comparable_fields"].items():
        lines.append(f"[{field}] {reason}\r\n\r\n")

    for f in result["fields"]:
        lines.append(_banner(f"{f['field_code']}  ({f['row_count']} values, {f['counts']})"))
        for r in f["rows"]:
            if r["classification"] == "UNCHANGED":
                continue
            lines.append(
                f"  {r['classification']:<12} {r['value']!r:<30} "
                f"sql={r['sql_identifier_code']!r}  v6={r['v6_identifier_code']!r}\r\n"
            )
        lines.append("\r\n")

    (evidence_dir / "FYBROC_ATTRIBUTEVALUE_SQL_DIFF.txt").write_text("".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="F120.5b AttributeValue vs V6 diff")
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument("--repo-root", type=Path, default=default_root)
    parser.add_argument("--evidence-dir", type=Path, default=None)
    args = parser.parse_args()

    repo_root: Path = args.repo_root.resolve()
    evidence_dir: Path = (args.evidence_dir or (repo_root / "docs" / "evidence" / "F120")).resolve()

    result = build_diff(repo_root, evidence_dir)
    write_outputs(evidence_dir, result)

    print(json.dumps({
        "step": STEP, "output_dir": str(evidence_dir),
        "comparable_field_count": result["comparable_field_count"],
        "total_classification_counts": result["total_classification_counts"],
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())