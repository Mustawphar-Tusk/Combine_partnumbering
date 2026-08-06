from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path(
    r"workbooks\Fybroc\Price Estimator-Fybroc.xlsm"
)

TARGETS = (
    ("MainTemplate", "C951"),
    ("MainTemplate", "C952"),
    ("MainTemplate", "C953"),
    ("MainTemplate", "C954"),
    ("MainTemplate", "C955"),
    ("Price Check", "D5"),
    ("Price Check", "D6"),
    ("Price Check", "D7"),
    ("Price Check", "D9"),
    ("Price Check", "D18"),
    ("M-#", "B5"),
)


def value_text(value: Any) -> str:
    if value is None:
        return "<blank>"

    return str(value).replace("\n", "\\n")


def merged_range_for_cell(
    worksheet,
    coordinate: str,
) -> str | None:
    for merged_range in worksheet.merged_cells.ranges:
        if coordinate in merged_range:
            return str(merged_range)

    return None


def print_neighborhood(
    worksheet,
    coordinate: str,
    *,
    row_radius: int = 2,
    column_radius: int = 5,
) -> None:
    anchor = worksheet[coordinate]

    minimum_row = max(1, anchor.row - row_radius)
    maximum_row = min(
        worksheet.max_row,
        anchor.row + row_radius,
    )

    minimum_column = max(
        1,
        anchor.column - column_radius,
    )
    maximum_column = min(
        worksheet.max_column,
        anchor.column + column_radius,
    )

    print()
    print("=" * 100)
    print(
        f"{worksheet.title}!{coordinate} = "
        f"{value_text(anchor.value)}"
    )

    merged = merged_range_for_cell(
        worksheet,
        coordinate,
    )

    if merged:
        print(f"Merged range: {merged}")

    print("-" * 100)

    for row_number in range(
        minimum_row,
        maximum_row + 1,
    ):
        values: list[str] = []

        for column_number in range(
            minimum_column,
            maximum_column + 1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            coordinate_text = (
                f"{get_column_letter(column_number)}"
                f"{row_number}"
            )

            if cell.value is None:
                continue

            merged_cell_range = merged_range_for_cell(
                worksheet,
                coordinate_text,
            )

            suffix = ""

            if merged_cell_range:
                suffix = (
                    f" [merged:{merged_cell_range}]"
                )

            values.append(
                f"{coordinate_text}="
                f"{value_text(cell.value)}"
                f"{suffix}"
            )

        if values:
            print(" | ".join(values))


def formula_references_target(
    formula: str,
    sheet_name: str,
    coordinate: str,
) -> bool:
    escaped_sheet = re.escape(sheet_name)
    escaped_coordinate = re.escape(coordinate)

    patterns = (
        rf"'{escaped_sheet}'!\$?{escaped_coordinate}",
        rf"{escaped_sheet}!\$?{escaped_coordinate}",
    )

    return any(
        re.search(
            pattern,
            formula,
            flags=re.IGNORECASE,
        )
        for pattern in patterns
    )


def find_direct_dependents(
    workbook,
    sheet_name: str,
    coordinate: str,
) -> list[tuple[str, str, str]]:
    matches: list[tuple[str, str, str]] = []

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value

                if not isinstance(value, str):
                    continue

                if not value.startswith("="):
                    continue

                if formula_references_target(
                    value,
                    sheet_name,
                    coordinate,
                ):
                    matches.append(
                        (
                            worksheet.title,
                            cell.coordinate,
                            value,
                        )
                    )

    return matches


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Workbook not found: {WORKBOOK_PATH}"
        )

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=False,
        data_only=False,
        keep_vba=True,
        keep_links=True,
    )

    print("=" * 100)
    print("M020 FYBROC PART-NUMBER TARGET TRACE")
    print("=" * 100)
    print(f"Workbook: {WORKBOOK_PATH}")

    for sheet_name, coordinate in TARGETS:
        if sheet_name not in workbook.sheetnames:
            print(
                f"\nMissing worksheet: {sheet_name}"
            )
            continue

        worksheet = workbook[sheet_name]

        print_neighborhood(
            worksheet,
            coordinate,
        )

        dependents = find_direct_dependents(
            workbook,
            sheet_name,
            coordinate,
        )

        print()
        print(
            f"Direct formula references to "
            f"{sheet_name}!{coordinate}: "
            f"{len(dependents)}"
        )

        for (
            dependent_sheet,
            dependent_coordinate,
            formula,
        ) in dependents[:30]:
            print(
                f"  {dependent_sheet}!"
                f"{dependent_coordinate}: "
                f"{formula}"
            )

        if len(dependents) > 30:
            print(
                f"  ... {len(dependents) - 30} "
                "additional references"
            )


if __name__ == "__main__":
    main()
