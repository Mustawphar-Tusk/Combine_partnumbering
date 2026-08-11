from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

PRICE_FILE = (
    ROOT
    / "workbooks"
    / "Fybroc"
    / "Price Estimator-Fybroc.xlsm"
)

NOM_FILE = (
    ROOT
    / "workbooks"
    / "Fybroc"
    / "Fybroc Nomenclature_V5.xlsm"
)

OUTPUT = (
    ROOT
    / "exports"
    / "fybroc_seal_alignment_proof.json"
)

SEARCH_TERMS = (
    "FKM",
    "Viton",
    "PTFE",
    "Teflon",
    "Carbon vs. Ceramic",
    "Silcar",
    "Customer Supplied",
    "Supplied by others",
    "No Seal",
    "Single Seal Gland",
    "Double Seal Gland",
    "8B2",
    "8-1T",
    "RAC",
    "RXO",
    "CRO",
)


def clean(value):
    if value is None:
        return None
    return str(value)


def search_workbook(
    workbook,
    workbook_name,
):
    results = []

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value

                if not isinstance(value, str):
                    continue

                folded = value.casefold()

                matched = [
                    term
                    for term in SEARCH_TERMS
                    if term.casefold() in folded
                ]

                if not matched:
                    continue

                results.append(
                    {
                        "workbook":
                            workbook_name,
                        "worksheet":
                            ws.title,
                        "cell":
                            cell.coordinate,
                        "value":
                            value,
                        "matched_terms":
                            matched,
                    }
                )

    return results


def show_data_validation(
    ws,
    coordinate,
):
    matches = []

    validations = getattr(
        ws,
        "data_validations",
        None,
    )

    if validations is None:
        return matches

    for dv in validations.dataValidation:
        try:
            contains = coordinate in dv.sqref
        except Exception:
            contains = False

        if not contains:
            continue

        matches.append(
            {
                "cell": coordinate,
                "type": dv.type,
                "formula1": clean(
                    dv.formula1
                ),
                "formula2": clean(
                    dv.formula2
                ),
                "allow_blank":
                    dv.allowBlank,
            }
        )

    return matches


def main():

    price = load_workbook(
        PRICE_FILE,
        data_only=False,
        keep_vba=True,
        read_only=False,
    )

    nomenclature = load_workbook(
        NOM_FILE,
        data_only=False,
        keep_vba=True,
        read_only=False,
    )

    result = {
        "price_check_validation": [],
        "seal_assembly_rows": [],
        "search_hits": [],
    }


    # ============================================================
    # Price Check input validation
    # ============================================================

    print("=" * 100)
    print(
        "PRICE CHECK SEAL INPUT VALIDATION"
    )
    print("=" * 100)

    pc = price["Price Check"]

    for coordinate in (
        "B32",
        "B33",
    ):
        validations = (
            show_data_validation(
                pc,
                coordinate,
            )
        )

        result[
            "price_check_validation"
        ].extend(
            validations
        )

        print()
        print(
            coordinate,
            "value =",
            repr(
                pc[
                    coordinate
                ].value
            ),
        )

        if validations:
            for item in validations:
                print(
                    "    validation type :",
                    item["type"],
                )
                print(
                    "    formula1        :",
                    item["formula1"],
                )
                print(
                    "    formula2        :",
                    item["formula2"],
                )
        else:
            print(
                "    validation       : "
                "<none found>"
            )


    # ============================================================
    # Seal Assembly context
    #
    # Earlier discovery only showed D:G.
    # Now expose A:L so repeated blocks can be explained.
    # ============================================================

    print()
    print("=" * 100)
    print(
        "NOMENCLATURE SEAL ASSEMBLY "
        "FULL ROW CONTEXT"
    )
    print("=" * 100)

    ws = nomenclature[
        "Seal Assembly"
    ]

    columns = (
        "A",
        "B",
        "C",
        "D",
        "E",
        "F",
        "G",
        "H",
        "I",
        "J",
        "K",
        "L",
    )

    for row in range(
        1,
        81,
    ):
        values = {
            col: ws[
                f"{col}{row}"
            ].value
            for col in columns
        }

        if not any(
            value is not None
            for value in values.values()
        ):
            continue

        record = {
            "row": row,
            **{
                col: clean(value)
                for col, value
                in values.items()
            },
        }

        result[
            "seal_assembly_rows"
        ].append(record)

        print()
        print(
            f"ROW {row}"
        )

        for col in columns:
            value = values[col]

            if value is not None:
                print(
                    f"    {col}{row:<3} = "
                    f"{value!r}"
                )


    # ============================================================
    # Search both authoritative workbooks for direct vocabulary
    # evidence.
    # ============================================================

    print()
    print("=" * 100)
    print(
        "VOCABULARY ALIGNMENT SEARCH"
    )
    print("=" * 100)

    hits = []

    hits.extend(
        search_workbook(
            price,
            PRICE_FILE.name,
        )
    )

    hits.extend(
        search_workbook(
            nomenclature,
            NOM_FILE.name,
        )
    )

    result[
        "search_hits"
    ] = hits

    for hit in hits:
        print()
        print(
            f"{hit['workbook']} | "
            f"{hit['worksheet']}!"
            f"{hit['cell']}"
        )
        print(
            "    terms:",
            ", ".join(
                hit[
                    "matched_terms"
                ]
            ),
        )
        print(
            "    value:",
            repr(
                hit[
                    "value"
                ]
            ),
        )


    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT.write_text(
        json.dumps(
            result,
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 100)
    print(
        "Seal Assembly rows:",
        len(
            result[
                "seal_assembly_rows"
            ]
        ),
    )
    print(
        "Vocabulary hits:",
        len(hits),
    )
    print(
        "Output:",
        OUTPUT,
    )
    print("=" * 100)

    price.close()
    nomenclature.close()


if __name__ == "__main__":
    main()
