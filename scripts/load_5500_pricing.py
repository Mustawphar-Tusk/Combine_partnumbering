"""Load 5500 series pricing from Price Estimator Pricebook into SQL."""
import openpyxl
import pyodbc

conn_str = (
    "DRIVER={ODBC Driver 18 for SQL Server};"
    "SERVER=localhost;"
    "DATABASE=PumpConfiguratorDB;"
    "Trusted_Connection=yes;"
    "Encrypt=yes;"
    "TrustServerCertificate=yes;"
)

def main():
    # Read 5500 pricing from Pricebook (cols 87-115)
    wb = openpyxl.load_workbook(
        "workbooks/Fybroc/Price Estimator-Fybroc.xlsm",
        read_only=True, data_only=True,
    )
    ws = wb["Pricebook"]

    # Extract rows: col87=size, col88=setting, col91=VR-1 price
    rows = []
    for r in range(6, 300):
        size = ws.cell(row=r, column=87).value
        setting = ws.cell(row=r, column=88).value
        if size is None and setting is None:
            break
        price_vr1 = ws.cell(row=r, column=91).value
        if size and price_vr1 is not None:
            try:
                amt = float(price_vr1)
                rows.append((str(size).strip(), str(setting).strip() if setting else "1", amt))
            except (ValueError, TypeError):
                pass
    wb.close()
    print(f"Extracted {len(rows)} pricing rows from 5500 block")

    # Load into SQL
    conn = pyodbc.connect(conn_str, autocommit=True)
    cursor = conn.cursor()

    pbv_id = cursor.execute(
        "SELECT PriceBookVersionId FROM price.PriceBookVersion WHERE IsCurrent=1"
    ).fetchone()[0]

    inserted = 0
    for size, setting, amount in rows:
        rule_code = f"5500-BASE-{size}-S{setting}-VR1"
        try:
            cursor.execute(
                "INSERT INTO price.PriceRule "
                "(PriceBookVersionId, RuleCode, RuleName, RuleType, ComponentCode, SeriesCode, "
                " Priority, Amount, PricingStatus, IsActive, SourceSizeValue, SourceOptionValue) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                pbv_id,
                rule_code,
                f"5500 Base {size} Setting {setting} VR-1",
                "base_pump",
                "BASE_PUMP",
                "5500",
                100,
                amount,
                "found",
                1,
                size.upper(),
                "VR-1 (Standard)",
            )
            inserted += 1
        except Exception:
            pass  # skip duplicates

    # Verify
    cnt = cursor.execute(
        "SELECT COUNT(*) FROM price.PriceRule pr "
        "JOIN price.PriceBookVersion pbv ON pbv.PriceBookVersionId=pr.PriceBookVersionId AND pbv.IsCurrent=1 "
        "WHERE pr.SeriesCode='5500' AND pr.IsActive=1"
    ).fetchone()[0]

    print(f"Inserted: {inserted}")
    print(f"Total 5500 price rules: {cnt}")
    conn.close()


if __name__ == "__main__":
    main()
